"""Local regression checks: python -m unittest -v test_pipeline."""

import contextlib
import io
import os
from pathlib import Path
import re
import socket
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parent
APP_NAME = "app"
MAPPING_NAME = "mapping"
GC_NAME = "garbage_check"
with patch.object(socket.socket, "connect", side_effect=AssertionError("Network forbidden")):
    import app as ap
    import mapping as mp
    import garbage_check as gc
MASTER = ROOT / "zuventus_product.xlsx"


class PipelineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.network_guard = patch.object(socket.socket, "connect", side_effect=AssertionError("Network forbidden"))
        cls.network_guard.start()
        wb = load_workbook(MASTER, read_only=True, data_only=True)
        try:
            values = list(wb.active.iter_rows(values_only=True))
            cls.rows = [dict(zip(values[0], row)) for row in values[1:]]
        finally:
            wb.close()
        cls.by_code = {row["PRODUCT_CODE"]: row for row in cls.rows}
        cls.df = mp.load_master_data(str(MASTER))
        cls.brand_map = mp.build_brand_product_map(cls.df)
        cls.variants = mp.extract_all_variants_from_data(cls.brand_map)
        cls.subvariants = mp.extract_all_sub_variants_from_data(cls.brand_map)
        cls.resources = ap.build_garbage_check_resources(MASTER)

    @classmethod
    def tearDownClass(cls):
        cls.network_guard.stop()

    def map_name(self, name, hint=None):
        with contextlib.redirect_stdout(io.StringIO()):
            return mp.process_product(name, self.brand_map, self.variants, self.subvariants, None,
                                      forced_brand=hint, auto_detect=hint is None, verbose=False)

    def assert_maps(self, name, expected):
        classification = gc.classify_single_product(name, *self.resources)
        self.assertEqual(classification[0], "0", (name, classification))
        result = self.map_name(name, classification[1])
        self.assertEqual(result.get("product_code"), expected["PRODUCT_CODE"], (name, result))
        self.assertEqual(result.get("output"), expected["product_name"], (name, result))

    def test_master_integrity(self):
        self.assertEqual(len(self.rows), 805)
        self.assertEqual(len(self.by_code), len(self.rows))
        self.assertEqual(len({re.sub(r"[^A-Z0-9]", "", row["product_name"].upper()) for row in self.rows}), len(self.rows))
        self.assertTrue(all(re.fullmatch(r"\d{9}", row["PRODUCT_CODE"]) for row in self.rows))
        self.assertEqual(set(gc.load_brands(str(MASTER))), set(self.brand_map))

    def test_all_exact_catalog_names(self):
        for row in self.rows:
            with self.subTest(code=row["PRODUCT_CODE"], name=row["product_name"]):
                self.assert_maps(row["product_name"], row)

    def test_all_case_and_spacing_variations(self):
        for row in self.rows:
            name = "  " + "   ".join(row["product_name"].lower().split()) + "  "
            with self.subTest(code=row["PRODUCT_CODE"]):
                self.assert_maps(name, row)

    def test_form_abbreviations(self):
        aliases = {"TABLET": "TAB", "TABLETS": "TAB", "CAPSULE": "CAP", "CAPSULES": "CAP",
                   "INJECTION": "INJ", "SUSPENSION": "SUSP", "SYRUP": "SYP"}
        pattern = re.compile(r"\b(" + "|".join(aliases) + r")\b", re.I)
        for row in self.rows:
            name = pattern.sub(lambda match: aliases[match.group().upper()], row["product_name"])
            if name != row["product_name"]:
                with self.subTest(code=row["PRODUCT_CODE"], input=name):
                    self.assert_maps(name, row)

    def test_repaired_qualifiers_and_strengths(self):
        cases = {
            "BISOTRAK 2.5 TAB": "423331717",
            "BISOTRAK T 2.5 TAB": "423331630",
            "BISOTRAK (T) 2.5 TAB": "423331630",
            "SCAVISTA 12 TAB": "427770269",
            "SCAVISTA 12 A TAB": "423331172",
            "EPORISE 10000 IU PFS": "423330917",
            "EPORISE 40000 IU PFS": "423330850",
        }
        for name, code in cases.items():
            with self.subTest(name=name):
                self.assert_maps(name, self.by_code[code])

    def test_generic_garbage_and_unknowns(self):
        for name in ["TOTAL", "GRAND TOTAL", "PRODUCT NAME", "CLOSING STOCK"]:
            with self.subTest(name=name):
                self.assertEqual(gc.classify_single_product(name, *self.resources)[0], "1")
        self.assertEqual(gc.classify_single_product("QZXWVVUNLISTED", *self.resources)[0], "MAYBE_PRODUCT")
        # Brand exceptions have been removed: these now follow generic brand-first matching.
        self.assertEqual(gc.classify_single_product("ZUVENTUS", *self.resources)[:2], ("0", "ZUVENTUS"))
        self.assertEqual(gc.classify_single_product("NEW", *self.resources)[:2], ("0", "NEW"))

    def test_generic_fuzzy_and_brand_boundaries(self):
        self.assert_maps("BISOTRAKK T 2.5 TAB", self.by_code["423331630"])
        self.assertEqual(gc.classify_single_product("ESLO 2.5 TAB", *self.resources)[:2], ("0", "ESLO"))
        self.assertEqual(gc.classify_single_product("ESLOMET 2.5 TAB", *self.resources)[:2], ("0", "ESLOMET"))
        short_index = [("ZU-C", gc.normalize_text("ZU-C"))]
        self.assertIsNotNone(gc.exact_product_match("ZU-C 500 TAB", short_index))
        self.assertIsNone(gc.exact_product_match("Z U C 500 TAB", short_index))
        self.assertIsNone(gc.fuzzy_product_match("ZUD 500 TAB", short_index))

    def test_generic_numeric_and_form_ocr_normalization(self):
        for source, expected in [("1000K", "1000000"), ("1.25K", "1250"),
                                 ("IN", "INJECTION"), ("GRANULIS", "GRANULES"), ("CHEB", "CHEWABLE")]:
            with self.subTest(source=source):
                self.assertEqual(mp.normalize_ocr_numeric_noise(source), expected)

    def test_only_confirmed_zero_reaches_mapper(self):
        names = [self.rows[0]["product_name"], "TOTAL", "QZXWVVUNLISTED"]
        classified = ap.run_garbage_check_phase(names, *self.resources)
        with patch.object(mp, "process_product", return_value={"product_code": "", "status": "NO_CLEAR_MATCH"}) as call:
            final, diagnostics = ap.run_mapping_phase(classified, MASTER)
        self.assertEqual(call.call_count, 1)
        self.assertEqual(call.call_args.args[0], names[0])
        self.assertEqual([row["PRODUCT_CODE"] for row in final], ["0", "1", "0"])
        self.assertEqual(diagnostics[1:], [None, None])

    def test_complete_workbook_order_duplicates_and_literal_text(self):
        with tempfile.TemporaryDirectory(prefix="zuventus_test_") as directory:
            source, output = Path(directory) / "input.xlsx", Path(directory) / "output.xlsx"
            name = self.by_code["423331630"]["product_name"]
            names = [name, "TOTAL", "", "QZXWVVUNLISTED", name, "=1+1"]
            wb = Workbook()
            wb.active.append(["PRODUCT_NAME"])
            for value in names:
                wb.active.append([value])
                wb.active.cell(wb.active.max_row, 1).data_type = "s"
            wb.save(source)
            wb.close()
            original = source.read_bytes()
            result = ap.run_pipeline(source, MASTER, output, verbose=False)
            self.assertEqual(source.read_bytes(), original)
            self.assertEqual([row["PRODUCT_NAME"] for row in result["mapping"]], names)
            self.assertEqual([result["mapping"][i]["PRODUCT_CODE"] for i in (0, 4)], ["423331630"] * 2)
            self.assertEqual(result["mapping"][1]["MASTER_PRODUCT_NAME"], "GARBAGE")
            self.assertEqual(result["mapping"][3]["MASTER_PRODUCT_NAME"], "NO SUGGESTION")
            wb = load_workbook(output, data_only=False)
            try:
                self.assertEqual(wb.sheetnames, ["garbage_check", "mapping"])
                self.assertEqual([cell.value for cell in wb["mapping"][1]], ap.FINAL_COLS)
                self.assertEqual(wb["mapping"].max_row, len(names) + 1)
                self.assertEqual(wb["mapping"].cell(7, 1).data_type, "s")
                self.assertEqual(wb["mapping"].cell(7, 1).value, "=1+1")
            finally:
                wb.close()
            second = ap.run_pipeline(output, MASTER, verbose=False)
            self.assertIsNotNone(second["backup"])
            self.assertTrue(second["backup"].is_file())
            self.assertEqual(second["mapping"], result["mapping"])

    def test_master_cannot_be_overwritten(self):
        with self.assertRaises(ValueError):
            ap.run_pipeline(MASTER, MASTER, verbose=False)
        with self.assertRaises(ValueError):
            ap.run_pipeline(ROOT / "missing.xlsx", MASTER, MASTER, verbose=False)

    def test_normalized_identity_collision_is_not_arbitrary(self):
        frame = self.df.iloc[:1].copy()
        duplicate = frame.copy()
        duplicate["PRODUCT_CODE"] = "999999999"
        frame = mp.pd.concat([frame, duplicate], ignore_index=True)
        brand_map = mp.build_brand_product_map(frame)
        with contextlib.redirect_stdout(io.StringIO()):
            result = mp.process_product(str(frame.iloc[0]["product_name"]), brand_map, auto_detect=True)
        self.assertFalse(result.get("product_code"), result)
        self.assertEqual(result["status"], "AMBIGUOUS_EXACT_NAME")

    def test_standalone_classification_then_mapping(self):
        with tempfile.TemporaryDirectory(prefix="zuventus_standalone_") as directory:
            path = Path(directory) / "work.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "garbage_check"
            ws.append(["PRODUCT_NAME"])
            for name in [self.rows[0]["product_name"], "TOTAL", "QZXWVVUNLISTED"]:
                ws.append([name])
            wb.create_sheet("notes")["A1"] = "Keep this sheet"
            wb.save(path)
            wb.close()
            with contextlib.redirect_stdout(io.StringIO()):
                gc.main(str(path), str(MASTER))
                frame = mp.process_excel_file(str(path), str(MASTER))
            self.assertEqual(len(frame), 1)
            self.assertEqual(frame.iloc[0][mp.PRODUCT_CODE_COLUMN], self.rows[0]["PRODUCT_CODE"])
            self.assertEqual(str(frame.iloc[0][mp.SOURCE_COLUMN]), "0")
            wb = load_workbook(path, data_only=True)
            try:
                self.assertEqual(wb["notes"]["A1"].value, "Keep this sheet")
                self.assertIn("mapping", wb.sheetnames)
            finally:
                wb.close()

    def test_import_without_api_credentials_or_network(self):
        environment = {key: value for key, value in os.environ.items()
                       if not ("API_KEY" in key or "TOKEN" in key)}
        script = ("import socket; socket.socket.connect=lambda *a,**k: (_ for _ in ()).throw(AssertionError('network')); "
                  f"import {APP_NAME}, {MAPPING_NAME}, {GC_NAME}; print('local imports OK')")
        result = subprocess.run([sys.executable, "-B", "-c", script], cwd=ROOT, env=environment,
                                capture_output=True, text=True, timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("local imports OK", result.stdout)

    def test_cli_entry_point_with_windows_console_encoding(self):
        with tempfile.TemporaryDirectory(prefix="zuventus_cli_") as directory:
            source, output = Path(directory) / "input.xlsx", Path(directory) / "result.xlsx"
            row = self.by_code["423331135"]
            wb = Workbook()
            wb.active.append(["PRODUCT_NAME"])
            wb.active.append([row["product_name"]])
            wb.save(source)
            wb.close()
            environment = dict(os.environ, PYTHONIOENCODING="cp1252")
            environment.pop("NOVITA_API_KEY", None)
            result = subprocess.run(
                [sys.executable, "-B", str(ROOT / f"{APP_NAME}.py"), "--input", str(source),
                 "--master", str(MASTER), "--output", str(output)],
                cwd=directory, env=environment, capture_output=True, timeout=60,
            )
            self.assertEqual(result.returncode, 0, result.stderr.decode("cp1252", errors="replace"))
            wb = load_workbook(output, data_only=True)
            try:
                self.assertEqual(wb["mapping"].cell(2, 2).value, row["PRODUCT_CODE"])
                self.assertEqual(wb["mapping"].cell(2, 3).value, row["product_name"])
            finally:
                wb.close()


if __name__ == "__main__":
    unittest.main()
