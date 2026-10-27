"""Local Zuventus mapping with master-derived brands, variants, and diagnostics."""
import re
import sys
from io import BytesIO
from dataclasses import replace
from functools import lru_cache
from itertools import islice
from pathlib import Path
import pandas as pd
from openpyxl import Workbook, load_workbook
from rapidfuzz import fuzz, process
from rapidfuzz.distance import Levenshtein, JaroWinkler
from excel_autosave import AsyncWorkbookSaver, save_workbook_atomic
from pro_diof import (parse_description as parse_pro_diof_description,
                      select_presentation as select_pro_diof_presentation,
                      normalize_diof_source, match_pro_brand)
from pack_selection import matched_measurement, measurement, parse_master_pack, select_pack_rows, UNIT_PATTERN
from source_brand_aliases import resolve_source_brand_alias, replace_source_brand_alias
from maxtra import (normalize_description as normalize_maxtra_description, tablet_count_evidence,
                    normalize_liquid_quantity, recover_explicit_volume, NUMERIC_OFFER)
MASTER_XLSX_PATH = Path(__file__).resolve().with_name('zuventus_product.xlsx')
INPUT_OUTPUT_XLSX_PATH = Path(__file__).resolve().with_name('mapping.xlsx')
ROW_LIMIT = None
PROCESS_MAYBE_PRODUCT_ONLY = False
PROCESS_CONFIRMED_ZERO_ONLY = True
if PROCESS_CONFIRMED_ZERO_ONLY and PROCESS_MAYBE_PRODUCT_ONLY:
    raise ValueError('PROCESS_CONFIRMED_ZERO_ONLY and PROCESS_MAYBE_PRODUCT_ONLY are mutually exclusive; set at most one of them to True.')
PROCESSING_MODE = 'local'
INPUT_COLUMN = 'Input_Column'
OUTPUT_COLUMN = 'Output_Column'
PRODUCT_CODE_COLUMN = 'Product_Code'
STATUS_COLUMN = 'Match_Status'
CONFIDENCE_COLUMN = 'Confidence'
CANDIDATES_COLUMN = 'Candidates_To_Model'
SUGGESTIONS_COLUMN = 'Suggestions'
SOURCE_COLUMN = 'Source'
PRODUCT_COL = 'product_name'
BRAND_COL = 'BRAND_NAME'
PACK_SIZE_COL = 'PACK_SIZE'
VARIANT_COL = 'VARIANT'
SUB_VARIANT_COL = 'SUB_VARIANT'
PRODUCT_CODE_COL = 'PRODUCT_CODE'
FUZZY_MIN_SCORE = 80
BRAND_NAME_NEXTSTEP_SCORE = 99
JW_MIN_SCORE = 0.85
BRAND_FUZZY_MAX_EDITS = 2
DOSAGE_UNIT_SUFFIXES = {'MG', 'GM', 'ML', 'MCG', 'IU', 'UNIT'}
FORM_HINT_TOKENS = {'BARS', 'SYR', 'TABLETS', 'INJ', 'CAP', 'GRANULES', 'CAPSULE', 'PFS', 'INJECTION', 'CREAM', 'AMPOULE', 'CAPSULES', 'LOTION', 'POWDER', 'TABLET', 'CAPS', 'PEN', 'VIAL', 'AMPS', 'IN', 'GEL', 'GRANUL', 'TAB', 'TONIC', 'AMPOULES', 'BAR', 'SYP', 'OINTMENT', 'SOLUTION', 'SYRUP', 'TABS', 'SACHET', 'SUSP', 'DROPS', 'TA', 'AMP', 'DROP', 'SUSPENSION'}
VARIANT_TOKEN_EQUIVALENTS = {'TA': 'TABLET', 'TAB': 'TABLET', 'TABS': 'TABLET', 'TABLETS': 'TABLET', 'CAP': 'CAPSULE', 'CAPS': 'CAPSULE', 'CAPSULES': 'CAPSULE', 'INJ': 'INJECTION', 'INJECTIONS': 'INJECTION', 'DROP': 'DROPS', 'SYP': 'SYRUP', 'BARS': 'BAR', 'GRANUL': 'GRANULES', 'LOT': 'LOTION', 'LO': 'LOTION', 'CRM': 'CREAM', 'CR': 'CREAM'}
INPUT_TOKEN_EQUIVALENTS = {'ING': 'INJECTION', 'INJ': 'INJECTION', 'INJN': 'INJECTION', 'INJN.': 'INJECTION', 'INJ.': 'INJECTION', 'SYR': 'SYRUP', 'SYP': 'SYRUP', 'SUSP': 'SUSPENSION', 'SUSPN': 'SUSPENSION', 'AMP': 'INJECTION', 'AMPS': 'INJECTION', 'AMPUL': 'INJECTION', 'AMPULE': 'INJECTION', 'AMPULES': 'INJECTION', 'AMPOULE': 'INJECTION', 'AMPOULES': 'INJECTION', 'VIAL': 'VIAL', 'VL': 'VIAL', 'PFS': 'PFS', 'PEN': 'PEN', 'TAB': 'TABLET', 'TABS': 'TABLET', 'TA': 'TABLET', 'TABLETS': 'TABLET', 'CAP': 'CAPSULE', 'CAPS': 'CAPSULE', 'CAPSULE': 'CAPSULE', 'CAPSULES': 'CAPSULE', 'DROP': 'DROPS', 'DROPS': 'DROPS', 'INJECTION': 'INJECTION', 'SYRUP': 'SYRUP', 'SUSPENSION': 'SUSPENSION', 'TABLETSNU': 'TABLETS', 'AMPOULEP': 'INJECTION', 'TABLETP': 'TABLET', 'AMPOUL': 'INJECTION', 'AMPOULS': 'INJECTION', 'INJG': 'INJECTION', 'INJE': 'INJECTION', 'NEDDLES': 'NEEDLES', 'NEEDALES': 'NEEDLES', 'PAKET': 'PACK', 'PACKET': 'PACK', 'VAIL': 'VIAL', 'VALE': 'VIAL', 'INJD': 'INJECTION', 'INGD': 'INJECTION', 'SY': 'SYRUP', 'TABD': 'TABLET', 'CAPD': 'CAPSULE', 'CAPSUL': 'CAPSULE', 'CAPSULS': 'CAPSULE', 'CREA': 'CREAM', 'OINT': 'OINTMENT', 'OINTB': 'OINTMENT', 'POW': 'POWDER', 'GRANUL': 'GRANULES', 'GRANU': 'GRANULES', 'GRANULS': 'GRANULES', 'LOT': 'LOTION', 'LOTION': 'LOTION', 'CRM': 'CREAM', 'CR': 'CREAM', 'CRM.': 'CREAM', 'CR.': 'CREAM', 'OINTM': 'OINTMENT', 'OINTMENT': 'OINTMENT', 'BAR': 'BAR', 'BARS': 'BAR', 'TABLETNU': 'TABLET', 'TABSP': 'TABLET', 'CAPSULP': 'CAPSULE', 'AMPOULP': 'INJECTION', 'DROPSP': 'DROPS', 'IN': 'INJECTION', 'GRANULIS': 'GRANULES', 'CHEB': 'CHEWABLE', 'F/WASH': 'FACE WASH', 'FACEWASH': 'FACE WASH', 'SYNDEN': 'SYNDET', 'MOISTURIS': 'MOISTURIZING', 'HYD': 'HYDRATING', 'CLEN': 'CLEAN', 'FLV': 'FLAVOR', 'FLAV': 'FLAVOR', 'FLAVOUR': 'FLAVOR'}
GENERIC_VARIANT_MATCH_TOKENS = {'TABLET', 'CAPSULE', 'INJECTION', 'SYRUP', 'DROPS', 'CREAM', 'LOTION', 'GEL', 'OINTMENT', 'SOLUTION', 'SUSPENSION', 'POWDER', 'SACHET', 'BAR', 'GRANULES'}

_NASAL_SPRAY_ABBREVIATION = re.compile(r'(?<![A-Z0-9])N\s*/\s*SPRAY(?![A-Z0-9])', re.IGNORECASE)
_VARIANT_ABBREVIATIONS = dict.fromkeys(('FORT', 'FORTS', 'FORET', 'FORTR', 'FOTRE', 'FROTE', 'FORATE', 'FO', 'FRT', 'FT'), 'FORTE')
_FORTE_ALIASES = '|'.join(sorted((*_VARIANT_ABBREVIATIONS, 'FORTE'), key=len, reverse=True))
_FORTE_END = r'(?=$|[^A-Z0-9]|(?:DRY)?(?:SYRUP|SYP|SYR|TABLET|TAB|TA|INJECTION|INJ)(?=$|[^A-Z0-9]|\d)|\d)'
_FORTE_DESCRIPTOR = re.compile(r'(?<![A-Z0-9])(?:' + _FORTE_ALIASES + ')' + _FORTE_END, re.I)
_FORTE_HINT = re.compile(r'F(?:O|RO|RT|T)', re.I)
_CTAX_FORTE_FORMS = re.compile(r'(?<![A-Z0-9])(?:D\s*/\s*(?:SYP|SYR(?:UP)?|S)?|DS|SUP|S)(?![A-Z0-9])', re.I)
_CTAX_FORTE_SPLIT_SYRUP = re.compile(r'(?<![A-Z0-9])SY\s+P(?![A-Z0-9])', re.I)
_CTAX_FORTE_34_6_METADATA = re.compile(r'(?<![A-Z0-9])\(\s*34\s*\+\s*6\s*\)(?![A-Z0-9])', re.I)
_CTAX_FIFTY = re.compile(r'(?<![A-Z0-9.])50(?:MLBOT|\s*M(?:\.L)?)(?![A-Z0-9.])', re.I)
_CTAX_FIFTY_BARE = re.compile(r'(?<![A-Z0-9.])50(?![A-Z0-9.])(?!\s*(?:ML|MG|GM|MCG|G|L|KG|IU|UNIT)\b)', re.I)
_DUPLICATE_50M = re.compile(r'(?<![A-Z0-9.])50M(?![A-Z0-9.])', re.IGNORECASE)
_EXPLICIT_50ML = re.compile(r'(?<![A-Z0-9.])50\s*ML(?![A-Z0-9])', re.IGNORECASE)
_PACK_MEASUREMENT = re.compile(r'(?<![A-Z0-9.])(\d+(?:\.\d+)?|\.\d+)\s*(MCG|MG|GM|ML|KG|G|L|IU|UNIT)\b', re.I)
_EXPLICIT_PACK_SOURCE = re.compile(
    r'(?<![A-Z0-9.])(?:\d+\s*[X*×]\s*)?(?:\d+(?:\.\d+)?|\.\d+)\s*(?:' + UNIT_PATTERN + r')(?![A-Z0-9])', re.I)

def norm(s: str) -> str:
    s = str(s or '').upper().strip()
    s = s.replace('+', ' PLUS ')
    s = s.replace('-', ' ')
    s = s.replace('/', ' ')
    s = s.replace('*', ' ')
    s = re.sub('\\bX\\b', ' ', s)
    s = s.replace('&', ' AND ')
    s = re.sub('\\b([A-Z])\\.([A-Z])\\b', '\\1\\2', s)
    s = re.sub('\\b([A-Z])\\.([A-Z])\\.([A-Z])\\b', '\\1\\2\\3', s)
    s = re.sub('[^A-Z0-9]+', ' ', s)
    s = re.sub('\\s+', ' ', s).strip()
    return s

def normalize_variant_token(s: str) -> str:
    s = str(s or '').upper().replace('+', ' PLUS ').strip()
    s = re.sub('[^A-Z0-9]+', ' ', s)
    return re.sub('\\s+', ' ', s).strip()

def compact_variant_token(s: str) -> str:
    return re.sub('[^A-Z0-9]+', '', str(s or '').upper().replace('+', ' PLUS '))

def compact_brand_token(s: str) -> str:
    return re.sub('[^A-Z0-9]+', '', str(s or '').upper())

def _expand_strength_token(tok: str) -> str:
    m = re.fullmatch('(\\d+(?:\\.\\d+)?)([KM])', str(tok or '').upper())
    if not m:
        return tok
    mult = 1000 if m.group(2) == 'K' else 1000000
    return str(int(float(m.group(1)) * mult))

def _normalize_strength_words(words) -> set:
    return {_expand_strength_token(w) for w in words}

def _input_compact_joins(input_text) -> set:
    toks = re.findall('[A-Z0-9]+', str(input_text or '').upper())
    joins = set()
    for i in range(len(toks) - 1):
        joins.add(toks[i] + toks[i + 1])
    for i in range(len(toks) - 2):
        joins.add(toks[i] + toks[i + 1] + toks[i + 2])
    return joins

def _expand_strength_in_compact(s: str) -> str:
    return re.sub('(\\d+(?:\\.\\d+)?)([KM])(?=\\d|$)', lambda m: str(int(float(m.group(1)) * (1000 if m.group(2) == 'K' else 1000000))), str(s or '').upper())

def canonical_variant_match_token(token: str) -> str:
    token = str(token or '').upper().strip()
    return VARIANT_TOKEN_EQUIVALENTS.get(token, token)

def extract_variant_match_tokens(text: str) -> set:
    normalized = normalize_ocr_numeric_noise(str(text or '').upper())
    tokens = set()
    for token in re.findall('[A-Z0-9]+', normalized):
        if token.isdigit():
            continue
        if re.fullmatch('\\d+(?:\\.\\d+)?', token):
            continue
        if re.fullmatch('\\d+[X*]\\d+(?:MG|GM|ML|MCG|IU|UNIT|G|L|KG)?', token):
            continue
        if re.fullmatch('\\d+(?:MG|GM|ML|MCG|IU|UNIT|G|L|KG)', token):
            continue
        if re.fullmatch('\\d+[STCP]', token):
            continue
        if token in {'X', 'PACK', 'OF'}:
            continue
        token = canonical_variant_match_token(token)
        if token and token not in DOSAGE_UNIT_SUFFIXES:
            tokens.add(token)
    return tokens

@lru_cache(maxsize=1024)
def build_flexible_brand_pattern(brand_name: str) -> str:
    brand_compact = compact_brand_token(brand_name)
    if not brand_compact:
        return ''
    separator = '[\\s\\-]*'
    char_pattern = separator.join((re.escape(ch) for ch in brand_compact))
    return '(?<![A-Z0-9])' + char_pattern + '(?![A-Z0-9])'

def get_brand_removal_name(brand_name: str, available_variants: set=None) -> str:
    tokens = re.findall('[A-Z0-9]+', str(brand_name or '').upper().replace('+', ' PLUS '))
    if len(tokens) <= 1:
        return str(brand_name or '').upper().strip()
    variant_compacts = set()
    if available_variants:
        variant_compacts = {compact_variant_token(variant) for variant in available_variants if variant and str(variant).upper() != 'PLAIN'}
    removable_tokens = [tokens[0]]
    suffix_tokens = tokens[1:]
    i = 0
    while i < len(suffix_tokens):
        token = suffix_tokens[i]
        two_token_compact = ''
        if i + 1 < len(suffix_tokens):
            two_token_compact = compact_variant_token(token + ' ' + suffix_tokens[i + 1])
        preserve_suffix = token in FORM_HINT_TOKENS or compact_variant_token(token) in variant_compacts or (two_token_compact and two_token_compact in variant_compacts)
        if not preserve_suffix:
            removable_tokens.append(token)
        i += 1
    return ' '.join(removable_tokens).strip()

def strip_brand_from_input(input_text: str, brand_name: str, available_variants: set=None) -> str:
    if not brand_name:
        return input_text
    input_upper = str(input_text or '').upper().strip()
    removal_brand = get_brand_removal_name(brand_name, available_variants=available_variants)
    brand_normalized = removal_brand.upper().replace('-', ' ').strip()
    brand_pattern = '\\b' + re.escape(brand_normalized).replace('\\ ', '[-\\s]+') + '\\b'
    stripped = re.sub(brand_pattern, '', input_upper, count=1, flags=re.IGNORECASE).strip()
    stripped = re.sub('\\b' + re.escape(brand_normalized) + '\\b', '', stripped, flags=re.IGNORECASE).strip()
    if stripped != input_upper:
        return re.sub('\\s+', ' ', stripped).strip()
    flexible_pattern = build_flexible_brand_pattern(removal_brand)
    if flexible_pattern:
        stripped = re.sub(flexible_pattern, '', input_upper, count=1, flags=re.IGNORECASE).strip()
        if stripped != input_upper:
            return re.sub('\\s+', ' ', stripped).strip()
    return re.sub('\\s+', ' ', input_upper).strip()


@lru_cache(maxsize=2048)
def _variant_vocabulary(attributes: tuple) -> frozenset:
    return frozenset(token for attribute in attributes
                     for token in re.findall('[A-Z0-9]+', str(attribute).upper()))


def _normalize_variant_abbreviations(text: str, variants, sub_variants=()) -> str:
    """Expand descriptor tokens only when supported by this catalog vocabulary."""
    if not _FORTE_HINT.search(text):
        return text
    vocabulary = _variant_vocabulary(tuple(sorted(set(variants) | set(sub_variants))))
    if 'FORTE' not in vocabulary:
        return text
    for pattern, prefix in _forte_compound_prefixes(tuple(sorted(variants))):
        text = pattern.sub(lambda match: prefix + ' ', text)

    def expand(match):
        token = match.group().upper()
        expanded = _VARIANT_ABBREVIATIONS.get(token, token) if token not in vocabulary else token
        separator = ' ' if match.end() < len(match.string) and match.string[match.end()].isalnum() else ''
        return expanded + separator

    return _FORTE_DESCRIPTOR.sub(expand, text)


@lru_cache(maxsize=512)
def _forte_compound_prefixes(variants: tuple) -> tuple:
    patterns = []
    for variant in variants:
        words = normalize_variant_token(variant).split()
        if 'FORTE' in words and words.index('FORTE'):
            prefix = ''.join(words[:words.index('FORTE')])
            flexible_prefix = r'[\s\-]*'.join(map(re.escape, prefix))
            patterns.append((re.compile(r'(?<![A-Z0-9])' + flexible_prefix + r'(?=[\s\-/]*(?:' + _FORTE_ALIASES + ')' + _FORTE_END + ')', re.I), prefix))
    return tuple(patterns)


@lru_cache(maxsize=128)
def _ctax_forte_identity_pattern(brand: str):
    prefix = build_flexible_brand_pattern(brand).removesuffix('(?![A-Z0-9])')
    return re.compile(prefix + r'[\s\-]*(?:[O0][\s\-]*){0,2}(?=(?:' + _FORTE_ALIASES + ')' + _FORTE_END + ')', re.I)


@lru_cache(maxsize=256)
def _forte_glued_brand_patterns(brand: str, variants: tuple) -> tuple:
    brand_prefix = build_flexible_brand_pattern(brand).removesuffix('(?![A-Z0-9])')
    patterns = []
    for variant in variants:
        words = normalize_variant_token(variant).split()
        if 'FORTE' in words and words.index('FORTE'):
            prefix = r'[\s\-]*'.join(map(re.escape, ''.join(words[:words.index('FORTE')])))
            patterns.append(re.compile(brand_prefix + r'(?=[\s\-]*' + prefix + r'[\s\-]*(?:' + _FORTE_ALIASES + ')' + _FORTE_END + ')', re.I))
    return tuple(patterns)


_FORTE_BRAND_CACHE = (None, ())


def _recover_forte_glued_brand(input_name, brand_map):
    global _FORTE_BRAND_CACHE
    cached_map, patterns_for_map = _FORTE_BRAND_CACHE
    if cached_map is not brand_map:
        patterns_for_map = []
        for brand, items in brand_map.items():
            variants = {item['variant'] for item in items if 'FORTE' in _variant_vocabulary((item['variant'],))}
            patterns = list(_forte_glued_brand_patterns(brand, tuple(sorted(variants))))
            if compact_brand_token(brand) == 'CTAX' and variants == {'O FORTE DRY'}:
                patterns.append(_ctax_forte_identity_pattern(brand))
            if patterns:
                patterns_for_map.append((brand, patterns))
        _FORTE_BRAND_CACHE = (brand_map, tuple(patterns_for_map))
    matches = [brand for brand, patterns in patterns_for_map if any(pattern.search(input_name) for pattern in patterns)]
    return matches[0] if len(matches) == 1 else None


@lru_cache(maxsize=2048)
def _variant_identity_tokens(variant: str) -> frozenset:
    return frozenset(extract_variant_match_tokens(variant) - GENERIC_VARIANT_MATCH_TOKENS - {'PLAIN'})


@lru_cache(maxsize=2048)
def _required_forte_tokens(variant: str, brand_name: str) -> frozenset:
    required = set(_variant_identity_tokens(variant))
    # Only this confirmed compound presentation permits omission of DRY.
    if compact_brand_token(brand_name) == 'CTAX' and normalize_variant_token(variant) == 'O FORTE DRY':
        required.discard('DRY')
    return frozenset(required)


def _prepare_forte_candidates(input_name: str, brand: str, items: list) -> tuple:
    """Keep explicit FORTE identity through all later filters and retries."""
    variants, sub_variants, _ = extract_available_attributes_from_items(items)
    forte_variants = {variant for variant in variants if 'FORTE' in _variant_vocabulary((variant,))}
    for pattern in _forte_glued_brand_patterns(brand, tuple(sorted(forte_variants))):
        input_name = pattern.sub(lambda match: str(brand) + ' ', input_name, count=1)
    if compact_brand_token(brand) == 'CTAX' and forte_variants == {'O FORTE DRY'}:
        # Confirmed C-TAX-only identity recovery: omitted O, OCR 0, or glued O.
        identity = _ctax_forte_identity_pattern(brand)
        input_name = identity.sub(lambda match: str(brand) + ' O ', input_name, count=1)
    brand_match = re.search(build_flexible_brand_pattern(brand), input_name.upper())
    offset = brand_match.end() if brand_match else 0
    input_name = input_name[:offset] + _normalize_variant_abbreviations(input_name[offset:], variants, sub_variants)
    descriptor = strip_brand_from_input(input_name, brand, available_variants=variants)
    descriptor = _normalize_variant_abbreviations(descriptor, variants, sub_variants)
    if not re.search(r'(?<![A-Z0-9])FORTE(?![A-Z0-9])', descriptor):
        return input_name, items, False, ''
    tokens = extract_variant_match_tokens(descriptor)
    forms = detect_dosage_form_in_input(input_name, verbose=False)
    groups = {_FORM_GROUPS[form] for form in forms if form in _FORM_GROUPS}
    if len(groups) > 1:
        return input_name, [], True, 'AMBIGUOUS_FORTE_FORM'
    candidates = []
    for item in items:
        variant = str(item.get('variant', 'PLAIN')).upper()
        sub_variant = str(item.get('sub_variant', 'PLAIN')).upper()
        vocabulary = _variant_vocabulary((variant, sub_variant))
        if 'FORTE' not in vocabulary:
            continue
        required = _required_forte_tokens(variant, brand)
        required |= _required_forte_tokens(sub_variant, '')
        if required.issubset(tokens):
            candidates.append(item)
    candidates = _form_compatible(candidates, forms)
    if not candidates:
        return input_name, [], True, 'VARIANT_NOT_IN_MASTER'
    if compact_brand_token(brand) == 'CTAX' and {'O', 'FORTE'} <= tokens:
        if forte_variants == {'O FORTE DRY'}:
            # Confirmed source meanings, restricted to this unique FORTE identity.
            input_name = _CTAX_FORTE_SPLIT_SYRUP.sub('SYRUP', input_name)
            input_name = _CTAX_FORTE_34_6_METADATA.sub(' ', input_name)
        input_name = re.sub(r'(?<![A-Z0-9])DRY(?=SYP(?:\b|\d))', 'DRY ', input_name)
        input_name = re.sub(r'(?<![A-Z0-9])(?:SYP|SYRUP|SYR)(?=\d)', lambda match: match.group() + ' ', input_name)
        input_name = _CTAX_FORTE_FORMS.sub(' DRY SYRUP ', input_name)
        input_name = re.sub(r'(?<![A-Z0-9])PHI(?![A-Z0-9])', ' ', input_name)
        input_name = _CTAX_FIFTY.sub('50 ML', input_name)
        input_name = _CTAX_FIFTY_BARE.sub('50 ML', input_name)
        input_name = re.sub(r'(?<=ML)32036005(?![A-Z0-9])', ' ', input_name)
        # The approved duplicate is 50M beside an actual 50 ML quantity only.
        if _EXPLICIT_50ML.search(input_name):
            input_name = _DUPLICATE_50M.sub(' ', input_name)
    return input_name, candidates, True, ''

def normalize_master_text(value, default: str='') -> str:
    if pd.isna(value):
        return default
    value = str(value).strip()
    if value in ('nan', 'NaN', 'None'):
        return default
    return value

def normalize_master_code(value) -> str:
    if pd.isna(value):
        return ''
    if isinstance(value, (int, float)) and (not isinstance(value, bool)):
        if float(value).is_integer():
            return str(int(value))
    return normalize_master_text(value, default='')

def normalize_master_attr(value, default: str='PLAIN') -> str:
    normalized = normalize_master_text(value, default='').upper()
    return normalized if normalized else default

def alpha_signature(value: str) -> str:
    letters = re.sub('[^A-Z]', '', str(value or '').upper())
    consonants = re.sub('[AEIOU]', '', letters)
    return consonants or letters

def collapse_spaced_thousands(s: str) -> str:
    prev = None
    out = str(s or '')
    while prev != out:
        prev = out
        out = re.sub('\\b(\\d{1,3})\\s+(\\d{3})\\b', '\\1\\2', out)
    return out

def numeric_strength_variant_value(variant: str) -> str:
    text = collapse_spaced_thousands(str(variant or '').upper().strip())
    text = re.sub('[\\s\\-]*(?:PFS|IU|UNIT|MG|MCG|ML|GM)\\.?$', '', text).strip()
    return text if re.fullmatch('\\d+(?:\\.\\d+)?', text) else ''

def normalize_ocr_numeric_noise(value: str) -> str:
    text = str(value or '').upper().replace('+', ' PLUS ')
    for source_token, target_token in INPUT_TOKEN_EQUIVALENTS.items():
        text = re.sub('\\b' + re.escape(source_token) + '\\b', target_token, text)
    text = re.sub('\\b(\\d+(?:\\.\\d+)?)\\s*K\\b', lambda m: str(int(float(m.group(1)) * 1000)) if float(m.group(1)) * 1000 == int(float(m.group(1)) * 1000) else str(float(m.group(1)) * 1000), text)
    text = re.sub('(?<=\\d)[/\\\\\\-](?=\\d)', '/', text)
    text = re.sub('\\b(\\d+)(10|15|20|30|14|28)S\\b', lambda m: f'{m.group(1)} {m.group(2)}S' if int(m.group(1) + m.group(2)) > 100 else m.group(0), text)
    return re.sub('\\s+', ' ', text).strip()

def extract_numeric_strength(value: str) -> str:
    text = normalize_master_attr(value, default='')
    if not text:
        return ''
    cleaned = re.sub('\\s*(?:MG|GM|ML|MCG|G|M|IU|UNIT|TAB|TABS|TABLET|TABLETS|INJ|INJECTION)\\s*$', '', text).strip()
    if re.fullmatch('\\d+(?:\\.\\d+)?(?:/\\d+(?:\\.\\d+)?)*', cleaned):
        return cleaned
    return ''

def is_numeric_subvariant_like(value: str) -> bool:
    return bool(re.fullmatch('\\d+(?:\\.\\d+)?(?:/\\d+(?:\\.\\d+)?)*', str(value or '').strip().upper()))
_ORAL_FORM_TOKENS = {'TABLET', 'TABLETS', 'TAB', 'TABS', 'CAPSULE', 'CAPSULES', 'CAP', 'CAPS', 'SYRUP', 'SYP', 'DROPS', 'DROP', 'SUSPENSION', 'SUSP', 'CREAM', 'GEL', 'OINTMENT', 'LOTION', 'SACHET', 'POWDER', 'POWD', 'GRANULES', 'GRANULE', 'SPRAY', 'NASALSPRAY', 'SOLUTION', 'SOL', 'TONIC', 'BAR', 'KIT'}

def repair_master_variant_fields(product_name: str, variant: str, sub_variant: str) -> tuple:
    variant = normalize_master_attr(variant)
    sub_variant = normalize_master_attr(sub_variant)
    _vt = [t for t in variant.split() if t.upper() not in _ORAL_FORM_TOKENS]
    variant = ' '.join(_vt).strip() or 'PLAIN'
    numeric_pattern = '\\d+(?:\\.\\d+)?(?:/\\d+(?:\\.\\d+)?)*'
    numeric_with_unit_pattern = numeric_pattern + '(?:\\s*(?:MG|GM|ML|MCG|G|M|IU|UNIT))?'
    noisy_sub_variant_tokens = {'MG', 'GM', 'ML', 'MCG', 'G', 'M', 'IU', 'UNIT', 'TAB', 'TABS', 'TABLET', 'TABLETS', 'INJ', 'INJECTION', 'PFS', 'VIAL', 'AMP', 'AMPS', 'SYR', 'SYRINGE', 'PEN'}
    if sub_variant == 'PLAIN':
        collapsed_variant = collapse_spaced_thousands(variant)
        form_dose = re.fullmatch('(\\d+(?:\\.\\d+)?)\\s*-?\\s*(?:PFS|VIAL|AMP|AMPS|SYR|SYRINGE|PEN|IU|INJ|INJECTION)', collapsed_variant)
        if form_dose:
            return ('PLAIN', form_dose.group(1))
    variant_numeric = extract_numeric_strength(variant)
    if sub_variant in noisy_sub_variant_tokens and variant_numeric:
        return ('PLAIN', variant_numeric)
    if re.fullmatch(numeric_pattern, variant) and re.search('[A-Z]', sub_variant) and (not is_numeric_subvariant_like(sub_variant)) and (sub_variant not in noisy_sub_variant_tokens):
        return (normalize_master_attr(sub_variant), variant)
    if sub_variant != 'PLAIN':
        return (variant, sub_variant)
    if re.fullmatch(numeric_pattern, variant):
        return ('PLAIN', variant)
    pure_numeric_with_unit_match = re.fullmatch('(' + numeric_pattern + ')\\s*(?:MG|GM|ML|MCG|G|M|IU|UNIT)', variant)
    if pure_numeric_with_unit_match:
        return ('PLAIN', pure_numeric_with_unit_match.group(1))
    split_match = re.fullmatch('(.+?)[\\s\\-]+(' + numeric_pattern + ')', variant)
    if split_match:
        repaired_variant = split_match.group(1).strip()
        repaired_sub_variant = split_match.group(2).strip()
        if repaired_variant:
            return (repaired_variant, repaired_sub_variant)
    compact_combo_match = re.fullmatch('([A-Z]+(?:\\s+[A-Z]+)*?)(' + numeric_pattern + ')', variant)
    if compact_combo_match:
        repaired_variant = compact_combo_match.group(1).strip()
        repaired_sub_variant = compact_combo_match.group(2).strip()
        if repaired_variant and repaired_sub_variant and ('/' in repaired_sub_variant or len(repaired_sub_variant.replace('.', '')) >= 2):
            return (repaired_variant, repaired_sub_variant)
    reverse_split_match = re.fullmatch('(' + numeric_pattern + ')\\s+([A-Z]+(?:\\s+[A-Z0-9]+)*)', variant)
    if reverse_split_match:
        repaired_sub_variant = reverse_split_match.group(1).strip()
        repaired_variant = reverse_split_match.group(2).strip()
        if repaired_variant and repaired_variant not in noisy_sub_variant_tokens:
            return (repaired_variant, repaired_sub_variant)
    unit_split_match = re.fullmatch('(.+?)[\\s\\-]+(' + numeric_with_unit_pattern + ')', variant)
    if unit_split_match:
        repaired_variant = unit_split_match.group(1).strip()
        repaired_sub_variant = unit_split_match.group(2).strip()
        repaired_sub_variant = re.sub('\\s*(?:MG|GM|ML|MCG|G|M|IU|UNIT)\\s*$', '', repaired_sub_variant).strip()
        if repaired_variant and repaired_sub_variant:
            return (repaired_variant, repaired_sub_variant)
    return (variant, sub_variant)

def is_reasonable_pack_count(value) -> bool:
    try:
        pack_num = float(str(value).strip())
    except (TypeError, ValueError):
        return False
    return pack_num.is_integer() and 0 < pack_num <= 100

def find_tablike_pack_matches(input_upper: str) -> list:
    matches = []
    pattern = '(?<![\\d.])(\\d+)\\s*(?:TA|TABS?|TABLETS?|STRIPS?|CAPS?|CAPSULES?)\\b'
    for m in re.finditer(pattern, input_upper):
        if is_reasonable_pack_count(m.group(1)):
            matches.append(m)
    return matches

def parse_pack_size_value(pack: str):
    pack_upper = str(pack or '').upper().strip()
    if not pack_upper or pack_upper in ('NAN', 'NONE'):
        return None
    combo_match = re.search('(\\d+)\\s*[X*]\\s*(\\d+)', pack_upper)
    if combo_match:
        return float(int(combo_match.group(1)) * int(combo_match.group(2)))
    numeric_match = re.search('(\\d+(?:\\.\\d+)?)', pack_upper)
    if numeric_match:
        return float(numeric_match.group(1))
    return None

def extract_combination_dosage(input_name: str) -> str:
    input_name = normalize_ocr_numeric_noise(input_name)
    input_name = re.sub('\\s*/\\s*', '/', input_name)
    pattern = '(?<![\\d./])(\\d+(?:\\.\\d+)?(?:/\\d+(?:\\.\\d+)?)+)(?![\\d./])'
    for match in re.finditer(pattern, input_name):
        candidate = match.group(1)
        if re.search('(?<!\\d)\\d+[X*×]', candidate):
            continue
        after = input_name[match.end():match.end() + 2]
        if re.match('[X*×]', after):
            continue
        components = candidate.split('/')
        for comp in components:
            comp_pos = input_name.find(comp, match.start())
            if comp_pos >= 0:
                after_comp = input_name[comp_pos + len(comp):comp_pos + len(comp) + 2]
                if re.match('[X*×]', after_comp):
                    continue
        return candidate
    return ''

def clean_duplicate_words(input_name: str) -> str:
    words = input_name.split()
    cleaned = []
    for i, word in enumerate(words):
        if i == 0 or word.upper() != words[i - 1].upper():
            cleaned.append(word)
    result = cleaned
    for n in (2, 3):
        i = 0
        deduped = []
        while i < len(result):
            chunk = result[i:i + n]
            if len(chunk) == n and result[i + n:i + 2 * n] == chunk:
                deduped.extend(chunk)
                i += 2 * n
            else:
                deduped.append(result[i])
                i += 1
        result = deduped
    return ' '.join(result)

def extract_brand_like_query(input_name: str) -> str:
    text = normalize_ocr_numeric_noise(input_name)
    tokens = re.findall('[A-Z0-9.]+', text)
    if not tokens:
        return ''
    skip_tokens = set(FORM_HINT_TOKENS) | set(DOSAGE_UNIT_SUFFIXES)
    skip_tokens |= {'PACK', 'OF', 'PCS', 'PC', 'NOS', 'NO', 'VIAL', 'AMP', 'AMPS', 'PFS', 'SYR'}
    brand_tokens = []
    for token in tokens:
        if token in skip_tokens:
            break
        if re.search('\\d', token):
            glued = re.match('^([A-Z]{2,})\\d', token)
            if glued and (not brand_tokens):
                brand_tokens.append(glued.group(1))
            break
        brand_tokens.append(token)
    return ' '.join(brand_tokens).strip()

def get_brand_variant_tags(brand_map: dict, brand_hint: str) -> set:
    if not brand_hint or not brand_map:
        return set()
    hint = str(brand_hint).strip().upper()
    key = hint if hint in brand_map else next((k for k in brand_map if k.upper() == hint), None)
    if key is None:
        return set()
    tags = set()
    for item in brand_map[key]:
        variant = str(item.get('variant', '') or '').strip().upper()
        if variant and variant != 'PLAIN' and re.fullmatch('[A-Z]{1,4}', variant):
            tags.add(variant)
    return tags

def strip_parenthetical_noise(s: str, brand_variant_tags: set=None) -> str:
    if not s:
        return s
    text = str(s)
    tag_pattern = '\\(\\s*([A-Z]{1,4})\\s*\\)'
    outside_tokens = None
    if brand_variant_tags:
        outside_tokens = set(re.findall('[A-Z0-9]+', re.sub(tag_pattern, ' ', text, flags=re.IGNORECASE).upper()))

    def _unwrap_or_drop(m):
        tag = m.group(1).strip().upper()
        if brand_variant_tags and tag in brand_variant_tags and (tag not in outside_tokens):
            return ' ' + tag + ' '
        return ' '
    cleaned = re.sub(tag_pattern, _unwrap_or_drop, text, flags=re.IGNORECASE)
    return re.sub('\\s+', ' ', cleaned).strip()

def extract_brand_suffix_tokens(input_name: str, brand_name: str) -> list:
    if not brand_name:
        return []
    input_upper = normalize_ocr_numeric_noise(input_name)
    brand_upper = brand_name.upper().strip()
    tokens = []
    skip_tokens = set(FORM_HINT_TOKENS) | set(DOSAGE_UNIT_SUFFIXES)
    skip_tokens |= {'PACK', 'OF', 'PCS', 'PC', 'NOS', 'NO', 'VIAL', 'AMP', 'AMPS', 'PFS', 'SYR', 'POWDE', 'POWDER', 'GRANUL', 'GRANULE', 'GRANULES', 'GRANULS'}

    def should_keep_token(token: str) -> bool:
        token_letters = re.sub('[^A-Z]', '', token.upper())
        return bool(token_letters) and token_letters not in skip_tokens

    def is_pack_notation(token: str) -> bool:
        return bool(re.match('^\\d+[Xx*×]\\d+', token) or re.match('^\\d+[Xx]\\d+', token))
    pattern1 = re.escape(brand_upper) + '-([A-Z]+\\d+|\\d+[A-Z]+)\\b'
    for m in re.findall(pattern1, input_upper):
        if should_keep_token(m) and (not is_pack_notation(m)):
            tokens.append(m)
    pattern2 = re.escape(brand_upper) + '\\s+([A-Z]+\\d+|\\d+[A-Z]+)\\b'
    for m in re.findall(pattern2, input_upper):
        if should_keep_token(m) and (not is_pack_notation(m)) and (m not in tokens):
            tokens.append(m)
    bm = re.search('\\b' + re.escape(brand_upper) + '\\b', input_upper)
    if bm:
        tail_match = re.match('[\\s\\-]*([A-Z]{2,})\\b', input_upper[bm.end():])
        if tail_match:
            alpha_token = tail_match.group(1)
            if alpha_token not in skip_tokens and alpha_token not in tokens:
                tokens.append(alpha_token)
    return tokens

def filter_items_by_compound_tokens(items: list, compound_tokens: list, verbose=False) -> list:
    if not compound_tokens:
        return items
    if verbose:
        print(f'\n  [COMPOUND TOKEN] Tokens to match: {compound_tokens}')
        print(f'  [COMPOUND TOKEN] Products before filter: {len(items)}')
    normalized_tokens = {compact_variant_token(token) for token in compound_tokens if token}
    matched = []
    for item in items:
        prod_upper = item['product'].upper().replace('-', ' ').replace('_', ' ')
        prod_compact = compact_variant_token(prod_upper)
        product_tokens = re.findall('[A-Z0-9]+', prod_upper)
        product_compact_tokens = {compact_variant_token(tok) for tok in product_tokens if tok}
        for i in range(len(product_tokens) - 1):
            product_compact_tokens.add(compact_variant_token(product_tokens[i] + product_tokens[i + 1]))
        for i in range(len(product_tokens) - 2):
            product_compact_tokens.add(compact_variant_token(product_tokens[i] + product_tokens[i + 1] + product_tokens[i + 2]))
        has_match = False
        for token in compound_tokens:
            normalized = compact_variant_token(token)
            if not token:
                continue
            if re.fullmatch('(?:[A-Z]+\\d+|\\d+[A-Z]+)', token):
                if token in product_tokens or normalized in product_compact_tokens:
                    has_match = True
                    break
                continue
            if token in prod_upper or (normalized and normalized in prod_compact):
                has_match = True
                break
        if has_match:
            matched.append(item)
            if verbose:
                print(f"    [OK] Matched '{item['product']}' on token(s) {compound_tokens}")
    if matched:
        if verbose:
            print(f'  [COMPOUND TOKEN] Products after filter: {len(matched)}')
        return matched
    if verbose:
        print(f'  [COMPOUND TOKEN] No products matched tokens -- keeping all {len(items)} products')
    return items

def get_compound_token_digits(compound_tokens: list) -> set:
    digits = set()
    for token in compound_tokens:
        for d in re.findall('\\d+', token):
            digits.add(d)
    return digits

def get_all_pack_positions(input_upper: str) -> list:
    positions = []
    for m in re.finditer('\\b(\\d+)\\s*[X*×]\\s*(\\d+(?:\\.\\d+)?)(?:\\s*(?:ML|GM|MG|MCG|IU|G|L|KG|TABLETS?|TABS?|TA|CAPSULES?|CAPS?|STRIPS?)(?![A-Z]))?', input_upper):
        positions.append((m.start(), m.end()))
    for m in re.finditer('\\bX\\s*(\\d+)\\b', input_upper):
        positions.append((m.start(), m.end()))
    for m in re.finditer('\\b(\\d+)([STCP])\\b', input_upper):
        if m.group(2) in ['S', 'T', 'C', 'P']:
            positions.append((m.start(), m.end()))
    for m in re.finditer('\\b(\\d+)\\s+(?:S|T|PCS)\\b', input_upper):
        positions.append((m.start(), m.end()))
    tablike_matches = find_tablike_pack_matches(input_upper)
    if tablike_matches and (not positions):
        last_match = tablike_matches[-1]
        positions.append((last_match.start(), last_match.end()))
    return positions

def is_position_in_pack_pattern(match_start: int, match_end: int, pack_positions: list) -> bool:
    for pack_start, pack_end in pack_positions:
        if pack_start <= match_start < pack_end or pack_start < match_end <= pack_end:
            return True
    return False

def split_strength_and_variant(input_name: str, available_variants: set, verbose=False) -> tuple:
    detected_variants = set()
    detected_sub_variants = set()
    input_upper = input_name.upper()
    pack_positions = get_all_pack_positions(input_upper)
    pattern = '\\b(\\d+(?:\\.\\d+)?)([A-Z]+)\\b'
    for match in re.finditer(pattern, input_upper):
        number_part = match.group(1)
        letter_part = match.group(2)
        found_letter_variant = False
        found_full_letter_variant = False
        is_dosage_suffix = letter_part in DOSAGE_UNIT_SUFFIXES
        if is_position_in_pack_pattern(match.start(), match.end(), pack_positions):
            continue
        if not is_dosage_suffix:
            if all((letter in available_variants for letter in letter_part)):
                for i, letter in enumerate(letter_part):
                    if letter in available_variants:
                        found_letter_variant = True
                        detected_variants.add(letter)
                        if i == 0:
                            detected_sub_variants.add(number_part)
            if letter_part in available_variants:
                found_full_letter_variant = True
                detected_variants.add(letter_part)
                detected_sub_variants.add(number_part)
        if number_part in available_variants and (not found_letter_variant) and (not found_full_letter_variant):
            detected_variants.add(number_part)
    return (detected_variants, detected_sub_variants)

def detect_compact_variant_strength_combo(input_name: str, available_variants: set) -> tuple:
    detected_variants = set()
    detected_sub_variants = set()
    tokens = re.findall('[A-Z0-9]+', normalize_ocr_numeric_noise(input_name))
    skip_tokens = _pack_form_skip_tokens() | {'PACK', 'OF'}
    for i in range(len(tokens) - 1):
        compact_token = tokens[i]
        next_token = tokens[i + 1]
        m = re.fullmatch('([A-Z]+)(\\d+(?:\\.\\d+)?)', compact_token)
        if not m:
            continue
        letter_part, numeric_part = m.groups()
        if letter_part in available_variants:
            detected_variants.add(letter_part)
            detected_sub_variants.add(numeric_part)
        if next_token in skip_tokens:
            continue
        candidate_variant = f'{letter_part} {next_token}'
        if candidate_variant in available_variants:
            detected_variants.add(candidate_variant)
            detected_sub_variants.add(numeric_part)
    return (detected_variants, detected_sub_variants)

def detect_variant_numeric_pairs(input_name: str, available_variants: set) -> tuple:
    detected_variants = set()
    detected_sub_variants = set()
    normalized = normalize_ocr_numeric_noise(input_name)
    seen = set()
    patterns = ['\\b([A-Z]+)(\\d+(?:\\.\\d+)?)\\b', '\\b([A-Z]+)[\\-\\s]+(\\d+(?:\\.\\d+)?)\\b']
    for pattern in patterns:
        for m in re.finditer(pattern, normalized):
            letter_part, numeric_part = m.groups()
            if letter_part in DOSAGE_UNIT_SUFFIXES:
                continue
            key = (letter_part, numeric_part)
            if key in seen:
                continue
            seen.add(key)
            if letter_part in available_variants:
                detected_variants.add(letter_part)
                detected_sub_variants.add(numeric_part)
    return (detected_variants, detected_sub_variants)

def product_contains_combination(product_name: str, combination: str) -> bool:
    product_norm = product_name.upper().replace(' ', '').replace('-', '')
    combo_norm = combination.upper().replace(' ', '').replace('-', '')
    if not combo_norm:
        return False
    pat = '(?<![\\d.])' + re.escape(combo_norm) + '(?![\\d.])'
    return bool(re.search(pat, product_norm))

def extract_available_attributes_from_items(items: list) -> tuple:
    available_variants = {'PLAIN'}
    available_sub_variants = {'PLAIN'}
    available_pack_sizes = set()
    for item in items:
        variant = str(item.get('variant', 'PLAIN')).strip().upper()
        sub_variant = str(item.get('sub_variant', 'PLAIN')).strip().upper()
        pack_size = str(item.get('pack_size', '')).strip()
        if variant and variant not in ('', 'NAN', 'NONE'):
            available_variants.add(variant)
        if sub_variant and sub_variant not in ('', 'NAN', 'NONE'):
            available_sub_variants.add(sub_variant)
        if pack_size and pack_size not in ('', 'nan', 'NaN', 'None'):
            available_pack_sizes.add(pack_size)
    return (available_variants, available_sub_variants, available_pack_sizes)

def extract_pack_size_from_input(s: str) -> str:
    return _extract_pack_size(s)


def _extract_pack_size(s: str, with_measurement=False):
    s_upper = normalize_ocr_numeric_noise(s)

    def result(value, match):
        if with_measurement and match.start() and s_upper[match.start() - 1] == '.':
            value = '.' + str(value)
        return matched_measurement(value, match, s_upper) if with_measurement else value

    triple_nxmxp = re.search('(?<![\\.\\d])(\\d+(?:\\.\\d+)?)\\s*[X*×]\\s*(\\d+)\\s*[X*×]\\s*(15)\\b', s_upper)
    if triple_nxmxp:
        return result('15', triple_nxmxp)
    nxm_matches = list(re.finditer('\\b(\\d+)\\s*[X*×]\\s*(\\d+(?:\\.\\d+)?)\\s*([STCP])?\\b', s_upper))
    for m in nxm_matches:
        n1 = float(m.group(1))
        n2 = float(m.group(2))
        total = n1 * n2
        if n1 <= 10 and n2 <= 100 and (total <= 500):
            return result(str(int(total) if total.is_integer() else total), m)
    full_nx1xm = re.search('\\b(\\d+)\\s*[X*×]\\s*1\\s*[X*×]\\s*(\\d+(?:\\.\\d+)?)\\s*([STCP])?\\b', s_upper)
    if full_nx1xm:
        n1 = float(full_nx1xm.group(1))
        n2 = float(full_nx1xm.group(2))
        total = 1 * n2
        return result(str(int(total) if total.is_integer() else total), full_nx1xm)
    if len(nxm_matches) > 1:
        pass
    m = re.search('\\b(\\d+)\\s*[X*×]\\s*(\\d+(?:\\.\\d+)?)(?:\\s*(?:ML|GM|MG|MCG|IU|G|L|KG|TA|TABS?|TABLETS?|CAPS?|CAPSULES?|T|C))?\\b', s_upper)
    if m:
        n1 = float(m.group(1))
        n2 = float(m.group(2))
        total = n1 * n2
        if n1 <= 10 and n2 <= 100 and (total <= 500):
            return result(str(int(total) if total.is_integer() else total), m)
    m = re.search('\\bPACK\\s*(?:OF\\s*)?(\\d+)\\b', s_upper)
    if m:
        return result(m.group(1), m)
    m = re.search('(?<!\\d)(\\d+)(?=[STCP])', s_upper)
    if m:
        return result(m.group(1), m)
    m = re.search('\\b(\\d+)\\s+(?:S|T|PCS)\\b', s_upper)
    if m:
        return result(m.group(1), m)
    adjacent_form_matches = []
    for m in re.finditer('(?<![\\d.])(\\d+)\\s*(?:TA|TABS?|TABLETS?|STRIPS?|CAPS?|CAPSULES?)\\b', s_upper):
        if is_reasonable_pack_count(m.group(1)):
            adjacent_form_matches.append((m.start(), m.group(1), m))
    for m in re.finditer('\\b(?:TA|TABS?|TABLETS?|STRIPS?|CAPS?|CAPSULES?)\\s*(\\d+)\\b', s_upper):
        if is_reasonable_pack_count(m.group(1)):
            adjacent_form_matches.append((m.start(), m.group(1), m))
    if adjacent_form_matches:
        adjacent_form_matches.sort(key=lambda x: x[0], reverse=True)
        return result(adjacent_form_matches[0][1], adjacent_form_matches[0][2])
    m = re.search('\\b(\\d+(?:\\.\\d+)?)\\s*(?:ML|GM|G|L|KG)\\b', s_upper)
    if m:
        return result(m.group(1), m)
    return None if with_measurement else ''

def load_master(path: str) -> pd.DataFrame:
    df = pd.read_excel(path, engine='openpyxl', dtype=object)
    for required in (PRODUCT_COL, BRAND_COL):
        if required not in df.columns:
            raise ValueError(f"Missing column '{required}'. Found: {list(df.columns)}")
    cols = [PRODUCT_COL, BRAND_COL]
    if PACK_SIZE_COL in df.columns:
        cols.append(PACK_SIZE_COL)
        print(f"[OK] Found '{PACK_SIZE_COL}' column in master - pack-size matching ENABLED.")
    else:
        print(f"[WARN] '{PACK_SIZE_COL}' column NOT found in master - pack sizes will be ignored.")
    if VARIANT_COL in df.columns:
        cols.append(VARIANT_COL)
        print(f"[OK] Found '{VARIANT_COL}' column in master - variant-based filtering ENABLED.")
    else:
        print(f"[WARN] '{VARIANT_COL}' column NOT found in master - all products will be shown.")
    if SUB_VARIANT_COL in df.columns:
        cols.append(SUB_VARIANT_COL)
        print(f"[OK] Found '{SUB_VARIANT_COL}' column in master - sub-variant filtering ENABLED.")
    else:
        print(f"[WARN] '{SUB_VARIANT_COL}' column NOT found in master - sub-variant filtering DISABLED.")
    if PRODUCT_CODE_COL in df.columns:
        cols.append(PRODUCT_CODE_COL)
        print(f"[OK] Found '{PRODUCT_CODE_COL}' column in master - product codes will be returned.")
    else:
        print(f"[WARN] '{PRODUCT_CODE_COL}' column NOT found in master - product codes will not be available.")
    df = df[cols].copy()
    df[PRODUCT_COL] = df[PRODUCT_COL].map(lambda value: '' if pd.isna(value) else str(value))
    df[BRAND_COL] = df[BRAND_COL].map(normalize_master_text)
    if PACK_SIZE_COL in df.columns:
        df[PACK_SIZE_COL] = df[PACK_SIZE_COL].map(normalize_master_text)
    else:
        df[PACK_SIZE_COL] = ''
    if VARIANT_COL in df.columns:
        df[VARIANT_COL] = df[VARIANT_COL].map(normalize_master_attr)
    else:
        df[VARIANT_COL] = 'PLAIN'
    if SUB_VARIANT_COL in df.columns:
        df[SUB_VARIANT_COL] = df[SUB_VARIANT_COL].map(normalize_master_attr)
    else:
        df[SUB_VARIANT_COL] = 'PLAIN'
    if PRODUCT_CODE_COL in df.columns:
        df[PRODUCT_CODE_COL] = df[PRODUCT_CODE_COL].map(normalize_master_code)
    else:
        df[PRODUCT_CODE_COL] = ''
    df = df[(df[PRODUCT_COL] != '') & (df[BRAND_COL] != '')]
    df['product_norm'] = df[PRODUCT_COL].map(norm)
    df['brand_norm'] = df[BRAND_COL].map(norm)
    return df

def build_brand_product_map(df: pd.DataFrame) -> dict:
    brand_map = {}
    for brand_norm, group in df.groupby('brand_norm'):
        original_brand = group.iloc[0][BRAND_COL]
        seen = set()
        items = []
        for _, row in group.iterrows():
            prod = str(row[PRODUCT_COL])
            pack = str(row.get(PACK_SIZE_COL, '')).strip()
            pack = '' if pack in ('nan', 'NaN', 'None') else pack
            variant = str(row.get(VARIANT_COL, 'PLAIN')).strip().upper()
            variant = 'PLAIN' if variant in ('', 'NAN', 'NONE') else variant
            sub_variant = str(row.get(SUB_VARIANT_COL, 'PLAIN')).strip().upper()
            sub_variant = 'PLAIN' if sub_variant in ('', 'NAN', 'NONE') else sub_variant
            variant, sub_variant = repair_master_variant_fields(prod, variant, sub_variant)
            product_code = str(row.get(PRODUCT_CODE_COL, '')).strip()
            product_code = '' if product_code in ('nan', 'NaN', 'None') else product_code
            key = (prod, pack, product_code, variant, sub_variant)
            if prod and key not in seen:
                seen.add(key)
                items.append({'product': prod, 'pack_size': pack, 'variant': variant, 'sub_variant': sub_variant, 'product_code': product_code, 'product_norm': norm(prod)})
        brand_map[original_brand] = items
    return brand_map

def extract_all_variants_from_data(brand_map: dict) -> set:
    all_variants = set()
    for brand, items in brand_map.items():
        for item in items:
            variant = item.get('variant', 'PLAIN')
            if variant != 'PLAIN':
                all_variants.add(variant.upper())
    all_variants.add('PLAIN')
    return all_variants

def extract_all_sub_variants_from_data(brand_map: dict) -> set:
    all_sub_variants = set()
    for brand, items in brand_map.items():
        for item in items:
            sub_variant = item.get('sub_variant', 'PLAIN')
            if sub_variant != 'PLAIN':
                all_sub_variants.add(sub_variant.upper())
    all_sub_variants.add('PLAIN')
    return all_sub_variants

def find_potential_brands(input_name: str, brand_map: dict) -> list:
    input_upper = input_name.upper()
    input_words = _normalize_strength_words(re.findall('\\b[A-Z0-9]+\\b', input_upper))
    input_compact = _expand_strength_in_compact(compact_brand_token(input_upper))
    potential = []
    for brand, items in brand_map.items():
        brand_upper = brand.upper().replace('-', ' ').strip()
        brand_words = _normalize_strength_words(brand_upper.split())
        brand_compact = _expand_strength_in_compact(compact_brand_token(brand))
        brand_pattern = '\\b' + re.escape(brand_upper) + '\\b'
        if re.search(brand_pattern, input_upper):
            extra_words = brand_words - input_words
            if not extra_words:
                potential.append((brand, items, 3))
            else:
                potential.append((brand, items, 1))
            continue
        flexible_brand_pattern = build_flexible_brand_pattern(brand)
        if flexible_brand_pattern and re.search(flexible_brand_pattern, input_upper):
            potential.append((brand, items, 4))
            continue
        if brand_compact and input_compact.startswith(brand_compact):
            remainder = input_compact[len(brand_compact):]
            if not remainder or remainder[0].isdigit():
                potential.append((brand, items, 3))
                continue
        if brand_words.issubset(input_words):
            potential.append((brand, items, 2))
            continue
        brand_first_word = brand_upper.split()[0] if brand_upper.split() else ''
        if brand_first_word and brand_first_word in input_words:
            potential.append((brand, items, 0))
    potential.sort(key=lambda x: (-x[2], len(x[0])))
    return [(brand, items) for brand, items, _ in potential]

def group_related_brands(primary_brand: str, brand_map: dict, input_name: str='', verbose=False) -> list:
    primary_upper = primary_brand.upper().replace('-', ' ').strip()
    input_words = _normalize_strength_words(re.findall('\\b[A-Z0-9]+\\b', input_name.upper())) if input_name else set()
    grouped_items = []
    related_brands = []
    primary_parts = primary_upper.split()
    primary_root = primary_parts[0] if primary_parts else ''
    for brand, items in brand_map.items():
        brand_upper = brand.upper().replace('-', ' ').strip()
        if brand_upper == primary_upper:
            grouped_items.extend(items)
            related_brands.append(brand)
        elif brand_upper.startswith(primary_upper + ' '):
            primary_words = _normalize_strength_words(primary_upper.split())
            brand_words = _normalize_strength_words(brand_upper.split())
            extra_words = {word for word in brand_words - primary_words if word not in _pack_form_skip_tokens()}
            if not extra_words or (input_words and extra_words.issubset(input_words)):
                grouped_items.extend(items)
                related_brands.append(brand)
            elif verbose:
                print(f"    [WARN] Skipping sub-brand '{brand}' -- extra words {extra_words} not in input")
        elif primary_root:
            brand_parts = brand_upper.split()
            brand_root = brand_parts[0] if brand_parts else ''
            if brand_root == primary_root:
                extra_words = {word for word in _normalize_strength_words(brand_parts[1:]) if word not in _pack_form_skip_tokens()}
                if extra_words and input_words and extra_words.issubset(input_words):
                    grouped_items.extend(items)
                    related_brands.append(brand)
    if verbose and len(related_brands) > 1:
        print(f'\n  Found {len(related_brands)} related brands:')
        for rb in sorted(related_brands):
            print(f"    - '{rb}' ({len(brand_map[rb])} products)")
    return grouped_items

def _pack_form_skip_tokens():
    tokens = set(FORM_HINT_TOKENS) | set(DOSAGE_UNIT_SUFFIXES)
    tokens |= {'PCS', 'PC', 'NOS', 'NO', 'VIAL', 'AMP', 'AMPS', 'PFS', 'SYR', 'POWDE', 'POWDER', 'GRANUL', 'GRANULE', 'GRANULES', 'GRANULS'}
    return tokens

def _brand_letters(s: str) -> str:
    return re.sub('[^A-Z]', '', str(s or '').upper())

def _is_strong_brand_token(token: str, brand_compacts: set) -> bool:
    c = compact_brand_token(token)
    if not c:
        return False
    if c in brand_compacts:
        return True
    return any((c.startswith(bc) and len(c) > len(bc) and c[len(bc)].isdigit() for bc in brand_compacts))
_DESEG_UNITS = 'MG|MCG|ML|GM|IU'
_DESEG_FORMS = 'TABLETS|TABLET|TABS|TAB|CAPSULES|CAPSULE|CAPS|CAP|INJECTION|INJ|SYRUP|SYP|DROPS|GRANULES|SUSPENSION'

def _split_glued_brand_token(tok: str, brand_letters: set) -> str:
    m = re.match('^[A-Za-z]+', tok)
    if not m:
        return tok
    lead = m.group(0).upper()
    best = ''
    for b in brand_letters:
        if len(best) < len(b) <= len(lead) and lead.startswith(b):
            best = b
    if best and best != lead:
        return tok[:len(best)] + ' ' + tok[len(best):].lstrip('-')
    if best == lead and len(tok) > len(lead):
        return lead + ' ' + tok[len(lead):].lstrip('-')
    return tok
_VST_CACHE_MAP = None
_VST_CACHE_VAL = None

def _variant_subvariant_tokens(brand_map: dict) -> tuple:
    global _VST_CACHE_MAP, _VST_CACHE_VAL
    if _VST_CACHE_MAP is brand_map:
        return _VST_CACHE_VAL
    var, subvar = (set(), set())
    for items in brand_map.values():
        for it in items:
            for t in re.findall('[A-Z]{2,}', str(it.get('variant', '')).upper()):
                var.add(t)
            for t in re.findall('[A-Z]{2,}', str(it.get('sub_variant', '')).upper()):
                subvar.add(t)
    _VST_CACHE_MAP, _VST_CACHE_VAL = (brand_map, (var, subvar))
    return (var, subvar)

def _split_glued_variant(tok: str, var_tokens: set, subvar_tokens: set) -> str:
    u = tok.upper()
    if not tok.isalpha() or len(u) < 4 or u in var_tokens or (u in subvar_tokens):
        return tok
    for i in range(2, len(u) - 1):
        if u[:i] in var_tokens and u[i:] in subvar_tokens:
            return tok[:i] + ' ' + tok[i:]
    return tok

def desegment_ocr_input(input_name: str, brand_map: dict) -> str:
    s = str(input_name or '')
    s = re.sub('[Xx]{2,}(?=\\s*-{2,})', '', s)
    s = re.sub('(?<!\\S)[Xx]{2,}(?!\\S)', ' ', s)
    s = re.sub('\\s*-{1,}\\s*$', '', s)
    s = re.sub('(\\d)\\.\\s+(\\d)', '\\1.\\2', s)
    s = re.sub('\\s+', ' ', s).strip()
    tokens = s.split()
    if not tokens:
        return s
    brand_compacts = {compact_brand_token(b) for b in brand_map}
    brand_letters = {_brand_letters(b) for b in brand_map if len(_brand_letters(b)) >= 4}
    if len(tokens) >= 2 and len(tokens[0]) == 1 and tokens[0].isalpha() and _is_strong_brand_token(tokens[0] + tokens[1], brand_compacts):
        tokens = [tokens[0] + tokens[1]] + tokens[2:]
    tokens[0] = _split_glued_brand_token(tokens[0], brand_letters)
    s = ' '.join(tokens)
    var_tokens, subvar_tokens = _variant_subvariant_tokens(brand_map)
    known_subbrands = var_tokens | subvar_tokens
    parts = s.split()
    new_parts = []
    for p in parts:
        split_done = False
        for b in brand_compacts:
            if p.startswith(b) and len(p) > len(b):
                suffix = p[len(b):]
                if suffix in known_subbrands:
                    new_parts.append(b)
                    new_parts.append(suffix)
                    split_done = True
                    break
        if not split_done:
            new_parts.append(p)
    s = ' '.join(new_parts)
    s = re.sub('\\bIN[- ](\\d+)', 'INJ \\1', s)
    s = re.sub('\\bSY[- ](\\d+)', 'SYP \\1', s)
    s = re.sub('\\bSUSP[- ](\\d+)', 'SUSPENSION \\1', s)
    var_tokens, subvar_tokens = _variant_subvariant_tokens(brand_map)
    parts = s.split()
    s = ' '.join(parts[:1] + [_split_glued_variant(p, var_tokens, subvar_tokens) for p in parts[1:]])
    s = re.sub('\\b(AMPOULEP)(\\d)', '\\1 \\2', s)
    s = re.sub('\\b(AMPOULE)(\\d)', '\\1 \\2', s)
    s = re.sub('\\b(AMP)(\\d)', '\\1 \\2', s)
    s = re.sub('\\b(AMPS)(\\d)', '\\1 \\2', s)
    s = re.sub('(?i)(\\d)\\s*(' + _DESEG_UNITS + ')(' + _DESEG_FORMS + ')', '\\1 \\2 \\3', s)
    s = re.sub('(?i)(?<=[A-WYZa-wyz])(?=\\d)', ' ', s)
    s = re.sub('(?i)(\\d)(' + _DESEG_FORMS + ')\\b', '\\1 \\2', s)
    return re.sub('\\s+', ' ', s).strip()

def is_substring_brand_reject(query_brand: str, candidate_brand: str) -> bool:
    q = _brand_letters(query_brand)
    c = _brand_letters(candidate_brand)
    if not q or not c:
        return False
    return c != q and c in q

def find_unmatched_qualifier_tokens(input_name: str, brand: str, available_variants, available_sub_variants, brand_items=None) -> list:
    text = str(input_name or '').upper().replace('-', ' ')
    brand_up = str(brand or '').upper().replace('-', ' ').strip()

    def _find_boundary(hay, needle):
        if not needle:
            return -1
        start = 0
        while True:
            i = hay.find(needle, start)
            if i < 0:
                return -1
            end = i + len(needle)
            before_ok = i == 0 or not hay[i - 1].isalnum()
            after_ok = end >= len(hay) or not hay[end].isalnum()
            if before_ok and after_ok:
                return i
            start = i + 1
    idx = _find_boundary(text, brand_up)
    brand_match = None
    if idx < 0 and brand:
        possible = re.search(build_flexible_brand_pattern(brand), text)
        if possible:
            remainder = text[possible.end():]
            strength = re.match(r'\s*(\d+(?:\.\d+)?)(?![\d.])', remainder)
            known_strengths = {str(it.get('sub_variant', '')) for it in brand_items or []
                               if is_numeric_subvariant_like(it.get('sub_variant', ''))}
            printed_doses = re.findall(r'(?<![\d.])(\d+(?:\.\d+)?)\s*(MG|MCG|GM|IU|UNIT)\b', remainder)
            compound_dose = re.match(r'\s*\d+(?:\.\d+)?\s*(?:MG|MCG|GM|IU|UNIT)?\s*(?:PLUS|/)\s*\d', remainder)
            if (strength and strength.group(1) in known_strengths and not compound_dose
                    and len({unit for _, unit in printed_doses}) <= 1
                    and all(dose == strength.group(1) for dose, _ in printed_doses)):
                catalog_units = {unit for it in brand_items or []
                                 if str(it.get('sub_variant', '')) == strength.group(1)
                                 for unit in re.findall(r'\d\s*(MG|MCG|GM|IU|UNIT)\b', str(it.get('product', '')).upper())}
                if not catalog_units or all(unit in catalog_units for _, unit in printed_doses):
                    brand_match = possible
    if idx >= 0:
        tail = text[idx + len(brand_up):]
    elif brand_match:
        tail = text[brand_match.end():]
    else:
        first = brand_up.split()[0] if brand_up.split() else ''
        pos = _find_boundary(text, first)
        if pos >= 0:
            tail = text[pos + len(first):]
        else:
            bq = extract_brand_like_query(input_name).upper()
            bq_first = bq.split()[0] if bq.split() else bq
            bpos = _find_boundary(text, bq_first)
            tail = text[bpos + len(bq_first):] if bpos >= 0 else ''
    tail = _normalize_variant_abbreviations(tail, available_variants or (), available_sub_variants or ())
    known = set(FORM_HINT_TOKENS) | set(DOSAGE_UNIT_SUFFIXES) | _pack_form_skip_tokens()
    known |= {'PLAIN'}
    known |= {str(v).upper() for v in available_variants or set()}
    known |= {str(v).upper() for v in available_sub_variants or set()}
    brand_vocab = set()
    for it in brand_items or []:
        for tok in re.findall('[A-Z]+', str(it.get('product', '')).upper()):
            brand_vocab.add(tok)
    prefix_targets = (brand_vocab | known) - ({'PLAIN'} | set(FORM_HINT_TOKENS))
    brand_compact = _brand_letters(brand_up)
    orphans = []
    tail_tokens = re.findall('[A-Z0-9]+', tail)
    for ti, tok in enumerate(tail_tokens):
        if any((ch.isdigit() for ch in tok)):
            break
        if tok in FORM_HINT_TOKENS:
            break
        canon = canonical_variant_match_token(tok)
        if tok in known or canon in known or tok in brand_vocab:
            continue
        if brand_compact + tok in brand_vocab:
            continue
        prev_join = tail_tokens[ti - 1] + tok if ti > 0 else ''
        next_join = tok + tail_tokens[ti + 1] if ti + 1 < len(tail_tokens) else ''
        if prev_join and (prev_join in known or prev_join in brand_vocab) or (next_join and (next_join in known or next_join in brand_vocab)):
            continue
        if len(tok) >= 3 and any((t != tok and t.startswith(tok) for t in prefix_targets)):
            continue
        orphans.append(tok)
    if idx > 0:
        for tok in re.findall('[A-Z]+', text[:idx]):
            if tok in {'SEMI', 'HALF'} and tok not in brand_vocab and (tok not in orphans):
                orphans.append(tok)
    return orphans

def find_abbrev_promoted_brand(input_name: str, matched_brand: str, brand_map: dict) -> str:
    if not matched_brand or not brand_map:
        return None
    input_upper = str(input_name or '').upper()
    matched_upper = str(matched_brand).upper().replace('-', ' ').strip()
    if not matched_upper:
        return None
    bm = re.search('\\b' + re.escape(matched_upper) + '\\b', input_upper)
    if not bm:
        return None
    tail_match = re.match('[\\s\\-]*([A-Z]{2,})\\b', input_upper[bm.end():])
    if not tail_match:
        return None
    abbrev_token = tail_match.group(1)
    if abbrev_token in _pack_form_skip_tokens():
        return None
    prefix = matched_upper + ' '
    candidates = []
    for brand in brand_map.keys():
        b_upper = str(brand).upper().replace('-', ' ').strip()
        if b_upper == matched_upper or not b_upper.startswith(prefix):
            continue
        next_word = b_upper[len(prefix):].split()[0] if b_upper[len(prefix):] else ''
        if len(next_word) < 2:
            continue
        if next_word.startswith(abbrev_token) or abbrev_token.startswith(next_word):
            candidates.append(brand)
    return candidates[0] if len(candidates) == 1 else None

def filter_items_by_product_name_abbrev(input_name: str, brand_name: str, items: list, verbose=False) -> list:
    if not brand_name or not items:
        return items
    input_upper = str(input_name or '').upper()
    matched_upper = str(brand_name).upper().replace('-', ' ').strip()
    if not matched_upper:
        return items
    bm = re.search('\\b' + re.escape(matched_upper) + '\\b', input_upper)
    if not bm and _FORTE_HINT.search(input_upper):
        variants, sub_variants, _ = extract_available_attributes_from_items(items)
        vocabulary = _variant_vocabulary(tuple(sorted(variants | sub_variants)))
        if 'FORTE' in vocabulary:
            descriptor = strip_brand_from_input(input_upper, brand_name, variants)
            descriptor = _normalize_variant_abbreviations(descriptor, variants, sub_variants)
            if 'FORTE' in extract_variant_match_tokens(descriptor):
                pattern = build_flexible_brand_pattern(brand_name)
                bm = re.search(pattern, input_upper) if pattern else None
    if not bm:
        return items
    tail_tokens = re.findall('[A-Z]{2,}', input_upper[bm.end():])
    if not tail_tokens:
        return items
    brand_tokens = set(matched_upper.split())
    filtered_items = list(items)
    for abbrev in tail_tokens:
        if abbrev in _pack_form_skip_tokens():
            continue
        expansions = set()
        abbrev_as_exact_match = False
        for item in filtered_items:
            prod = str(item['product']).upper().replace('-', ' ')
            for tok in re.findall('\\b[A-Z]{2,}\\b', prod):
                if tok in brand_tokens or tok in _pack_form_skip_tokens():
                    continue
                if tok == abbrev:
                    abbrev_as_exact_match = True
                elif len(tok) > len(abbrev) and tok.startswith(abbrev):
                    expansions.add(tok)
        if abbrev_as_exact_match:
            continue
        if len(expansions) != 1:
            continue
        expansion = next(iter(expansions))
        narrowed = [i for i in filtered_items if re.search('\\b' + re.escape(expansion) + '\\b', str(i['product']).upper().replace('-', ' '))]
        if narrowed and len(narrowed) < len(filtered_items):
            if verbose:
                print(f"  [PRODUCT-NAME ABBREV] '{abbrev}' -> '{expansion}': filtered {len(filtered_items)} -> {len(narrowed)} items")
            filtered_items = narrowed
    return filtered_items

def find_best_brand_for_input(input_name: str, brand_map: dict, verbose=False) -> tuple:
    if verbose:
        print('\n' + '=' * 80)
        print('STEP 1: BRAND DETECTION')
        print('=' * 80)
        print(f"Input name: '{input_name}'")
    potential = find_potential_brands(input_name, brand_map)
    if verbose:
        print(f'\nPattern matching found {len(potential)} potential brand(s):')
        for i, (brand, items) in enumerate(potential[:5], 1):
            print(f"  {i}. '{brand}' ({len(items)} products)")
    input_upper = input_name.upper().replace('-', ' ')
    input_words = _normalize_strength_words(re.findall('\\b[A-Z0-9]+\\b', input_upper))
    input_compact = _expand_strength_in_compact(compact_brand_token(input_name))

    def _brand_position(brand_name: str) -> int:
        brand_upper = str(brand_name or '').upper().replace('-', ' ').strip()
        if not brand_upper:
            return 10 ** 6
        m = re.search('\\b' + re.escape(brand_upper) + '\\b', input_upper)
        if m:
            return m.start()
        brand_compact = _expand_strength_in_compact(compact_brand_token(brand_name))
        if brand_compact:
            compact_pos = input_compact.find(brand_compact)
            if compact_pos >= 0:
                return compact_pos
        return 10 ** 6
    input_joins = _input_compact_joins(input_name)
    filtered_potential = []
    for brand, items in potential:
        brand_upper = brand.upper().replace('-', ' ')
        brand_words = _normalize_strength_words(brand_upper.split())
        brand_compact = _expand_strength_in_compact(compact_brand_token(brand))
        extra_words = {w for w in brand_words - input_words if compact_brand_token(w) not in input_joins}
        compact_match = False
        if brand_compact and input_compact.startswith(brand_compact):
            remainder = input_compact[len(brand_compact):]
            if not remainder or remainder[0].isdigit():
                compact_match = True
        if extra_words and (not compact_match):
            if verbose:
                print(f"  [FAIL] Rejecting '{brand}' - extra words: {extra_words}")
            continue
        filtered_potential.append((brand, items))
    if filtered_potential:
        filtered_potential.sort(key=lambda x: (_brand_position(x[0]), -len(x[0])))
        selected_brand, _ = filtered_potential[0]
        promoted = find_abbrev_promoted_brand(input_name, selected_brand, brand_map)
        if promoted and promoted != selected_brand:
            if verbose:
                print(f"\n[OK] Promoted '{selected_brand}' -> '{promoted}' (input contains abbreviation of sub-brand)")
            selected_brand = promoted
        if verbose:
            print(f"\n[OK] Selected brand: '{selected_brand}'")
        grouped_items = group_related_brands(selected_brand, brand_map, input_name=input_name, verbose=verbose)
        return (selected_brand, grouped_items)
    if not potential:
        input_norm = norm(input_name)
        brand_query_norm = norm(extract_brand_like_query(input_name))
        if verbose:
            print(f'\nFuzzy matching against {len(brand_map)} brands...')
        all_brands = list(brand_map.keys())
        candidate_brands = []
        for brand in all_brands:
            brand_upper = brand.upper().replace('-', ' ')
            brand_first_word = brand_upper.split()[0] if brand_upper.split() else ''
            if not brand_first_word or brand_first_word in input_upper or len(brand.split()) == 1:
                candidate_brands.append(brand)
        if not candidate_brands:
            candidate_brands = all_brands
        best = None
        best_source = 'fuzzy'
        brand_root_norm = brand_query_norm.split()[0] if brand_query_norm else ''
        brand_tokens_norm = [t for t in brand_query_norm.split() if len(t) >= 4]
        if brand_query_norm:
            best = process.extractOne(brand_query_norm, candidate_brands, scorer=fuzz.ratio, score_cutoff=BRAND_NAME_NEXTSTEP_SCORE)
            if best:
                best_source = 'brand-only'
        for fuzzy_query in [brand_query_norm, brand_root_norm, *brand_tokens_norm, input_norm]:
            if best:
                break
            if not fuzzy_query:
                continue
            best = process.extractOne(fuzzy_query, candidate_brands, scorer=fuzz.token_sort_ratio, score_cutoff=FUZZY_MIN_SCORE)
            if best:
                break
        if not best:
            for fuzzy_query in [brand_query_norm, brand_root_norm, *brand_tokens_norm]:
                if not fuzzy_query:
                    continue
                best = process.extractOne(fuzzy_query, candidate_brands, scorer=JaroWinkler.normalized_similarity, score_cutoff=JW_MIN_SCORE)
                if best:
                    best_source = 'jaro-winkler'
                    break
        if best:
            best_brand, score, _ = best
            if is_substring_brand_reject(extract_brand_like_query(input_name), best_brand):
                if verbose:
                    print(f"\n[FAIL] Rejecting fuzzy brand '{best_brand}' -- it is a sub-token of the input brand (different brand, not a typo)")
                return (None, None)
            cand_compact = compact_brand_token(best_brand)
            edit_dists = [Levenshtein.distance(compact_brand_token(q), cand_compact) for q in (extract_brand_like_query(input_name), brand_root_norm, *brand_tokens_norm) if compact_brand_token(q)]
            if edit_dists and min(edit_dists) > BRAND_FUZZY_MAX_EDITS:
                if verbose:
                    print(f"\n[FAIL] Rejecting fuzzy brand '{best_brand}' -- {min(edit_dists)} edits from input brand (> {BRAND_FUZZY_MAX_EDITS}); different brand, not a typo")
                return (None, None)
            promoted = find_abbrev_promoted_brand(input_name, best_brand, brand_map)
            if promoted and promoted != best_brand:
                if verbose:
                    print(f"\n[OK] Promoted fuzzy match '{best_brand}' -> '{promoted}'")
                best_brand = promoted
            if verbose:
                if best_source == 'brand-only':
                    print(f"\n[OK] Brand-only high-confidence match: '{best_brand}' (score: {score})")
                else:
                    print(f"\n[OK] Fuzzy match: '{best_brand}' (score: {score})")
            grouped_items = group_related_brands(best_brand, brand_map, input_name=input_name, verbose=verbose)
            return (best_brand, grouped_items)
    return (None, None)

def detect_variants_in_input(input_name: str, brand_name: str, available_variants: set, compound_token_letters: set=None, verbose=False, brand_items: list=None, _exclude_inferred_forte=False) -> set:
    input_upper = normalize_ocr_numeric_noise(input_name)
    pack_positions = get_all_pack_positions(input_upper)

    def _single_letter_variant_only_in_pack(token: str) -> bool:
        token = str(token or '').strip().upper()
        if token not in {'S', 'T', 'C', 'P'}:
            return False
        token_matches = list(re.finditer('\\b' + re.escape(token) + '\\b', input_upper))
        if not token_matches:
            return False
        return all((is_position_in_pack_pattern(m.start(), m.end(), pack_positions) for m in token_matches))
    if brand_name:
        input_without_brand = strip_brand_from_input(input_upper, brand_name, available_variants=available_variants)
    else:
        input_without_brand = input_upper
    input_without_brand = _normalize_variant_abbreviations(input_without_brand, available_variants)
    input_without_brand = re.sub('\\s+', ' ', input_without_brand).strip()
    normalized_input = normalize_variant_token(input_without_brand)
    compact_input_token_list = [compact_variant_token(token) for token in re.findall('[A-Z0-9.\\-]+', input_without_brand) if compact_variant_token(token)]
    compact_input_tokens = set(compact_input_token_list)
    for i in range(len(compact_input_token_list) - 1):
        compact_input_tokens.add(compact_input_token_list[i] + compact_input_token_list[i + 1])
    for i in range(len(compact_input_token_list) - 2):
        compact_input_tokens.add(compact_input_token_list[i] + compact_input_token_list[i + 1] + compact_input_token_list[i + 2])
    detected_variants = set()
    split_variants, _ = split_strength_and_variant(input_without_brand, available_variants, verbose=verbose)
    if compound_token_letters:
        split_variants -= compound_token_letters
    detected_variants.update(split_variants)
    combo_variants, combo_sub_variants = detect_compact_variant_strength_combo(input_without_brand, available_variants)
    if compound_token_letters:
        combo_variants -= compound_token_letters
    detected_variants.update(combo_variants)
    pair_variants, _ = detect_variant_numeric_pairs(input_without_brand, available_variants)
    detected_variants.update(pair_variants)
    words = re.findall('[A-Z0-9]+', input_without_brand)
    input_match_tokens = extract_variant_match_tokens(input_without_brand)
    input_descriptor_tokens = {token for token in input_match_tokens if token not in GENERIC_VARIANT_MATCH_TOKENS}
    input_forms_detected = detect_dosage_form_in_input(input_name, verbose=False)
    for variant in available_variants:
        if variant == 'PLAIN':
            continue
        variant_upper = variant.upper()
        if len(variant_upper) == 1 and _single_letter_variant_only_in_pack(variant_upper):
            continue
        variant_compact = compact_variant_token(variant_upper)
        if compound_token_letters and variant_upper in compound_token_letters:
            continue
        if variant_compact and variant_compact in compact_input_tokens:
            detected_variants.add(variant)
            continue
        if variant_upper in words:
            detected_variants.add(variant)
            continue
        patterns = ['[\\-\\s]' + re.escape(variant_upper) + '[\\-\\s]', '[\\-\\s]' + re.escape(variant_upper) + '\\b', '\\b' + re.escape(variant_upper) + '[\\-\\s]']
        for pattern in patterns:
            if re.search(pattern, input_without_brand):
                detected_variants.add(variant)
                break
    for variant in available_variants:
        if variant == 'PLAIN' or len(variant) <= 1:
            continue
        if compound_token_letters and variant.upper() in compound_token_letters:
            continue
        variant_upper = variant.upper()
        variant_normalized = normalize_variant_token(variant_upper)
        if variant_normalized and re.search('\\b' + re.escape(variant_normalized).replace('\\ ', '\\s+') + '\\b', normalized_input):
            detected_variants.add(variant)
            continue
        if re.search('\\b' + re.escape(variant_upper) + '\\b', input_without_brand):
            detected_variants.add(variant)
    for variant in available_variants:
        if variant == 'PLAIN':
            continue
        if len(str(variant).strip()) == 1 and _single_letter_variant_only_in_pack(variant):
            continue
        variant_match_tokens = extract_variant_match_tokens(variant)
        if not variant_match_tokens:
            continue
        if variant_match_tokens.issubset(input_match_tokens):
            if variant not in detected_variants:
                detected_variants.add(variant)
            continue
        if len(variant_match_tokens) >= 2 and input_descriptor_tokens and (len(input_descriptor_tokens) >= 2) and input_descriptor_tokens.issubset(variant_match_tokens):
            detected_variants.add(variant)
        variant_form_tokens = variant_match_tokens & GENERIC_VARIANT_MATCH_TOKENS
        variant_non_form_tokens = variant_match_tokens - GENERIC_VARIANT_MATCH_TOKENS
        if variant not in detected_variants and variant_non_form_tokens and variant_non_form_tokens.issubset(input_match_tokens) and input_forms_detected and variant_form_tokens and forms_share_group(variant_form_tokens, input_forms_detected):
            detected_variants.add(variant)
    input_form_tokens = input_match_tokens & GENERIC_VARIANT_MATCH_TOKENS
    for form_token in sorted(input_form_tokens):
        form_variants = []
        for variant in available_variants:
            if variant == 'PLAIN':
                continue
            variant_match_tokens = extract_variant_match_tokens(variant)
            if form_token in variant_match_tokens:
                form_variants.append(variant)
        if len(form_variants) == 1:
            variant = form_variants[0]
            non_form = extract_variant_match_tokens(variant) - GENERIC_VARIANT_MATCH_TOKENS
            if non_form and (not non_form.issubset(input_match_tokens)):
                continue
            if variant not in detected_variants:
                detected_variants.add(variant)
    input_collapsed = collapse_spaced_thousands(input_without_brand)
    input_strength_numbers = set(re.findall('\\d+(?:\\.\\d+)?', input_collapsed))
    if input_strength_numbers:
        for variant in available_variants:
            if variant in detected_variants:
                continue
            vnum = numeric_strength_variant_value(variant)
            if vnum and vnum in input_strength_numbers:
                detected_variants.add(variant)
    _INJ_FORM_VAR = {'IV', 'INJ', 'INJECTION', 'VIAL', 'PFS', 'AMP', 'AMPS', 'SYR', 'SYRINGE'}
    if input_forms_detected and forms_share_group({'INJECTION'}, input_forms_detected):
        for variant in available_variants:
            if variant == 'PLAIN' or variant in detected_variants:
                continue
            vt = {re.sub('[^A-Z]', '', t) for t in str(variant).upper().split()}
            vt = {t for t in vt if t}
            if vt and vt <= _INJ_FORM_VAR:
                if input_strength_numbers and brand_items is not None:
                    has_dose = any((str(it.get('sub_variant', '')).strip() in input_strength_numbers for it in brand_items if str(it.get('variant', '')).upper() == str(variant).upper()))
                    if not has_dose:
                        continue
                detected_variants.add(variant)
    if brand_items:
        input_combos = set(re.findall('(?<![\\d./])\\d+(?:\\.\\d+)?(?:/\\d+(?:\\.\\d+)?)+(?![\\d./])', input_without_brand))
        if input_combos:
            for it in brand_items:
                ivar = str(it.get('variant', 'PLAIN')).strip().upper()
                isub = str(it.get('sub_variant', '')).strip().upper()
                if ivar and ivar != 'PLAIN' and ('/' in isub) and (isub in input_combos):
                    detected_variants.add(ivar)
    variant_suffix_patterns = {v: '\\b' + re.escape(v) + '\\b' for v in available_variants if v != 'PLAIN'}
    input_upper_for_var = input_without_brand.upper()
    pack_positions = get_all_pack_positions(input_upper_for_var)

    def is_in_pack_pattern(start, end):
        for ps, pe in pack_positions:
            if ps <= start < pe or ps < end <= pe:
                return True
        return False
    for var_name, var_pattern in variant_suffix_patterns.items():
        if var_name not in available_variants:
            continue
        if len(var_name) == 1:
            for m in re.finditer(var_pattern, input_upper_for_var):
                if not is_in_pack_pattern(m.start(), m.end()):
                    detected_variants.add(var_name)
                    break
        elif re.search(var_pattern, input_upper_for_var):
            detected_variants.add(var_name)
    single_letter_variants = {v for v in available_variants if len(v) == 1 and v.isalpha()}
    for var_name in single_letter_variants:
        if var_name in available_variants:
            pattern_glued = '\\d(?:\\.\\d+)?' + re.escape(var_name) + '(?![A-Z])'
            for m in re.finditer(pattern_glued, input_upper_for_var):
                if not is_in_pack_pattern(m.start(), m.end()):
                    detected_variants.add(var_name)
                    break
            pattern_token_start = '(?<![A-Z0-9])' + re.escape(var_name) + '(?=\\d)'
            for m in re.finditer(pattern_token_start, input_upper_for_var):
                if not is_in_pack_pattern(m.start(), m.end()):
                    detected_variants.add(var_name)
                    break
            pattern_spaced = '\\d(?:\\.\\d+)?\\s+' + re.escape(var_name) + '(?![A-Z])'
            for m in re.finditer(pattern_spaced, input_upper_for_var):
                if not is_in_pack_pattern(m.start(), m.end()):
                    detected_variants.add(var_name)
                    break
    if 'FORTE' in input_match_tokens:
        input_tokens = input_match_tokens
        detected_variants = {variant for variant in detected_variants
                             if 'FORTE' not in _variant_vocabulary((variant,))
                             or _required_forte_tokens(variant, brand_name).issubset(input_tokens)}
    if _exclude_inferred_forte:
        detected_variants = {variant for variant in detected_variants if 'FORTE' not in variant.split()}
    if not any((variant != 'PLAIN' for variant in detected_variants)):
        detected_variants.add('PLAIN')
    if verbose:
        print(f'  Final detected variants: {sorted(detected_variants)}')
    return detected_variants

def detect_sub_variants_in_input(input_name: str, brand_name: str, available_sub_variants: set, available_variants: set, compound_token_digits: set=None, verbose=False) -> set:
    input_upper = normalize_ocr_numeric_noise(input_name)
    detected_sub_variants = set()
    if brand_name:
        input_without_brand = strip_brand_from_input(input_upper, brand_name, available_variants=available_variants)
    else:
        input_without_brand = input_upper
    combo_dosage = extract_combination_dosage(input_upper)
    combo_components: set = set()
    if combo_dosage:
        detected_sub_variants.add(combo_dosage)
        combo_components = set(re.findall('\\d+(?:\\.\\d+)?', combo_dosage))
    _, split_sub_variants = split_strength_and_variant(input_without_brand, available_variants, verbose=verbose)
    if compound_token_digits:
        split_sub_variants -= compound_token_digits
    if combo_components:
        split_sub_variants -= combo_components
    detected_sub_variants.update(split_sub_variants)
    _, combo_sub_variants = detect_compact_variant_strength_combo(input_without_brand, available_variants)
    if combo_components:
        combo_sub_variants -= combo_components
    detected_sub_variants.update(combo_sub_variants)
    _, pair_sub_variants = detect_variant_numeric_pairs(input_without_brand, available_variants)
    if combo_components:
        pair_sub_variants -= combo_components
    detected_sub_variants.update(pair_sub_variants)
    pack_positions = get_all_pack_positions(input_upper)
    sorted_sub_variants = sorted([sv for sv in available_sub_variants if sv != 'PLAIN'], key=lambda x: (len(str(x).replace('.', '')), len(str(x))), reverse=True)
    matched_positions = []

    def is_position_matched(start, end):
        for ms, me in matched_positions:
            if ms <= start < me or ms < end <= me:
                return True
        return False
    for sub_var in sorted_sub_variants:
        if sub_var in detected_sub_variants:
            continue
        if compound_token_digits and str(sub_var) in compound_token_digits:
            continue
        if combo_components and str(sub_var) in combo_components:
            continue
        sub_str = extract_numeric_strength(sub_var)
        if not sub_str:
            continue
        unit_pattern = '(?<![\\d.])' + re.escape(sub_str) + '\\s*(?:MG|GM|MCG|G|M|IU|UNIT)(?:\\b|/)'
        for m in re.finditer(unit_pattern, input_upper):
            if not is_position_in_pack_pattern(m.start(), m.end(), pack_positions) and (not is_position_matched(m.start(), m.end())):
                detected_sub_variants.add(sub_var)
                matched_positions.append((m.start(), m.end()))
                break
    for sub_var in sorted_sub_variants:
        if sub_var in detected_sub_variants:
            continue
        if compound_token_digits and str(sub_var) in compound_token_digits:
            continue
        if combo_components and str(sub_var) in combo_components:
            continue
        sub_str = str(sub_var)
        pattern = '(?<![\\d.])' + re.escape(sub_str) + '(?![\\d.])'
        for m in re.finditer(pattern, input_upper):
            context_start = max(0, m.start() - 5)
            context_end = min(len(input_upper), m.end() + 5)
            context = input_upper[context_start:context_end]
            if '/' in context:
                continue
            if re.match('\\s*M?L\\b', input_upper[m.end():]):
                continue
            if m.start() > 0 and input_upper[m.start() - 1] == '.' or re.match('\\.\\d', input_upper[m.end():]):
                continue
            if is_position_in_pack_pattern(m.start(), m.end(), pack_positions):
                continue
            if not is_position_in_pack_pattern(m.start(), m.end(), pack_positions) and (not is_position_matched(m.start(), m.end())):
                detected_sub_variants.add(sub_var)
                matched_positions.append((m.start(), m.end()))
                break
    for sub_var in sorted_sub_variants:
        if sub_var in detected_sub_variants:
            continue
        if compound_token_digits and str(sub_var) in compound_token_digits:
            continue
        if combo_components and str(sub_var) in combo_components:
            continue
        sub_str = str(sub_var)
        pattern = '[-\\s/]' + re.escape(sub_str) + '(?![\\d.])'
        for m in re.finditer(pattern, input_upper):
            if re.match('\\s*M?L\\b', input_upper[m.end():]):
                continue
            if re.match('\\.\\d', input_upper[m.end():]):
                continue
            if not is_position_in_pack_pattern(m.start(), m.end(), pack_positions) and (not is_position_matched(m.start(), m.end())):
                detected_sub_variants.add(sub_var)
                matched_positions.append((m.start(), m.end()))
                break
    input_alpha_tokens = {t for t in re.findall('[A-Z]+', input_upper) if len(t) >= 3}
    input_alpha_signatures = {alpha_signature(t) for t in input_alpha_tokens}
    input_tokens_set_upper = set(re.findall('[A-Z]+', input_upper))
    for sub_var in sorted_sub_variants:
        if sub_var in detected_sub_variants or is_numeric_subvariant_like(sub_var):
            continue
        sub_var_upper = str(sub_var).strip().upper()
        if sub_var_upper not in input_tokens_set_upper:
            continue
        sub_signature = alpha_signature(sub_var)
        if len(sub_signature) >= 3 and sub_signature in input_alpha_signatures:
            detected_sub_variants.add(sub_var)
    normalized_descriptor = _normalize_variant_abbreviations(input_without_brand, available_variants, available_sub_variants)
    if 'FORTE' in extract_variant_match_tokens(normalized_descriptor) and 'FORTE' in available_sub_variants:
        detected_sub_variants.add('FORTE')
    if not any((is_numeric_subvariant_like(s) for s in detected_sub_variants if s != 'PLAIN')):
        rescue_text = normalize_ocr_numeric_noise(input_without_brand)
        input_numbers = set()
        for m in re.finditer('(\\d+(?:\\.\\d+)?)\\s*([A-Z]*)', rescue_text):
            unit = m.group(2)
            if unit == 'L' or unit.startswith('ML'):
                continue
            input_numbers.add(m.group(1))
        strength_subs = {s for s in available_sub_variants if is_numeric_subvariant_like(s) and '/' not in str(s)}
        hits = {s for s in strength_subs if str(s) in input_numbers}
        if len(hits) == 1:
            detected_sub_variants.add(next(iter(hits)))
    detected_sub_variants.add('PLAIN')
    if verbose:
        print(f'  Final detected sub-variants: {sorted(detected_sub_variants)}')
    return detected_sub_variants

def _specific_detected_variants(detected_variants: set) -> set:
    specific = {v for v in detected_variants if v and v != 'PLAIN'}
    if len(specific) <= 1:
        return detected_variants
    word_sets = {v: set(str(v).upper().split()) for v in specific}
    kept = set()
    for v, vset in word_sets.items():
        is_subset_of_other = any((v != other and vset < word_sets[other] for other in word_sets))
        if not is_subset_of_other:
            kept.add(v)
    if 'PLAIN' in detected_variants:
        kept.add('PLAIN')
    return kept if kept else detected_variants

def filter_items_by_variant(items: list, detected_variants: set, reference_items: list=None, verbose=False) -> list:
    if verbose:
        print('\n' + '=' * 80)
        print('STEP 2: VARIANT-BASED FILTERING')
        print('=' * 80)
        print(f'Detected variants in input: {sorted(detected_variants)}')
        print(f'Total products before filtering: {len(items)}')
    specific_variants = _specific_detected_variants(detected_variants)
    if verbose and specific_variants != detected_variants:
        print(f'Narrowed to most-specific: {sorted(specific_variants)}')

    def _partition(variant_set):
        vmatches, plains = ([], [])
        for item in items:
            iv = item.get('variant', 'PLAIN').upper()
            if iv == 'PLAIN':
                plains.append(item)
            elif iv in variant_set:
                vmatches.append(item)
        return (vmatches, plains)
    variant_matches, plain_products = _partition(specific_variants)
    if not variant_matches and specific_variants != detected_variants:
        if verbose:
            print(f'[WARN] No items for specific variants -- falling back to full detected set')
        variant_matches, plain_products = _partition(detected_variants)
    if variant_matches:
        filtered = variant_matches
        if verbose:
            print(f'\n[OK] Found {len(variant_matches)} specific variant matches')
            print(f'  Excluding {len(plain_products)} PLAIN products')
    elif reference_items and reference_items is not items:
        ref_matches, ref_plains = ([], [])
        for item in reference_items:
            iv = item.get('variant', 'PLAIN').upper()
            if iv == 'PLAIN':
                ref_plains.append(item)
            elif iv in specific_variants or iv in detected_variants:
                ref_matches.append(item)
        if ref_matches:
            filtered = ref_matches
            if verbose:
                print(f'\n[OK] Recovered {len(ref_matches)} variant matches from full brand list')
        else:
            filtered = plain_products if plain_products else ref_plains
            if verbose:
                print(f'\n[WARN] No variant matches; using {len(filtered)} PLAIN as fallback')
    else:
        filtered = plain_products
        if verbose:
            print(f'\n[WARN] No variant matches; using {len(plain_products)} PLAIN as fallback')
    if verbose:
        print(f'\nProducts after variant filtering: {len(filtered)}')
        if variant_matches:
            variants_kept = {item.get('variant', 'PLAIN') for item in filtered}
            for var in sorted(variants_kept):
                count = sum((1 for i in filtered if i.get('variant', 'PLAIN') == var))
                print(f'  - {var}: {count} products')
    if not filtered:
        return items
    return filtered

def filter_by_subvariant_in_product_name(items: list, detected_sub_variants: set, verbose=False) -> list:
    numeric_sub_variants = {sv for sv in detected_sub_variants if sv != 'PLAIN' and '/' not in str(sv)}
    if not numeric_sub_variants:
        return items
    matches_with_scores = []
    for item in items:
        product_name = item['product'].upper()
        best_score = 0
        matched_svs = set()
        for sub_var in numeric_sub_variants:
            sub_str = str(sub_var)
            score = 0
            if re.search('(?<![\\d.])' + re.escape(sub_str) + '\\s*(?:MG|GM|ML|MCG|IU)\\b', product_name):
                score = 100
            elif re.search('(?<![\\d.])' + re.escape(sub_str) + '(?![\\d.])', product_name):
                if not re.search('(?<![\\d.])' + re.escape(sub_str) + '\\s*[X*×]', product_name):
                    score = 90
            elif re.search('[A-Z]' + re.escape(sub_str) + '\\b', product_name):
                score = 75
            elif sub_str in product_name:
                is_false_match = False
                for m in re.finditer('(?<![\\d.])' + re.escape(sub_str) + '(?![\\d.])', product_name):
                    after = product_name[m.end():m.end() + 3]
                    if after and after[0].isdigit() and (not re.match('\\s*(?:MG|GM|ML|MCG|IU|G|L|KG)\\b', after)):
                        is_false_match = True
                        break
                    before = product_name[max(0, m.start() - 3):m.start()]
                    if before and before[-1].isdigit() and (not re.match('(?<![\\d.])\\d+$', before)):
                        is_false_match = True
                        break
                if not is_false_match:
                    score = 50
            if score > best_score:
                best_score = score
                matched_svs = {sub_var}
            elif score == best_score and score > 0:
                matched_svs.add(sub_var)
        if best_score >= 75:
            matches_with_scores.append((item, best_score, matched_svs))
    if matches_with_scores:
        matches_with_scores.sort(key=lambda x: x[1], reverse=True)
        best_score = matches_with_scores[0][1]
        return [item for item, score, _ in matches_with_scores if score == best_score]
    return items

def calculate_product_name_priority_score(product_name: str, input_name: str, detected_variants: set=None, detected_sub_variants: set=None, _form_order=None) -> int:
    _TOKEN_SPLIT = '[\\s\\-+/]+'
    input_upper = input_name.upper().strip()
    pack_positions = get_all_pack_positions(input_upper)
    input_words_raw = {w for w in re.split(_TOKEN_SPLIT, input_upper) if w}
    input_words = set()
    for w in input_words_raw:
        is_pack_word = False
        for m in re.finditer('\\b' + re.escape(w) + '\\b', input_upper):
            if is_position_in_pack_pattern(m.start(), m.end(), pack_positions):
                is_pack_word = True
                break
        if not is_pack_word:
            input_words.add(w)
    input_raw_numbers = set(re.findall('\\d+(?:\\.\\d+)?', input_upper))
    input_numbers = set()
    for num in input_raw_numbers:
        is_pack_num = False
        for m in re.finditer('(?<![\\d.])' + re.escape(num) + '(?![\\d.])', input_upper):
            if is_position_in_pack_pattern(m.start(), m.end(), pack_positions):
                is_pack_num = True
                break
        if not is_pack_num:
            input_numbers.add(num)
    product_upper = product_name.upper().strip()
    form_precedence = _form_precedence(input_name) if _form_order is None else _form_order
    product_comparison = _form_neutral_text(product_upper, form_precedence, product=True)
    product_words = {w for w in re.split(_TOKEN_SPLIT, product_comparison) if w}
    _PACK_SHAPE = re.compile('^\\d+[Xx*]\\d+[A-Z]?$|^\\d+[STCP]$|^\\d+PCS?$|^\\d+NOS?$', re.IGNORECASE)
    _UNIT_SUFFIX = re.compile('^(?:MG|ML|GM|MCG|IU|KG)\\.?$', re.IGNORECASE)

    def _is_structural(token: str) -> bool:
        if _PACK_SHAPE.match(token) or _UNIT_SUFFIX.match(token):
            return True
        bare = token.rstrip('.')
        if bare in FORM_HINT_TOKENS:
            return True
        glued = re.match('^\\d+([A-Z]+)$', token)
        if glued and glued.group(1) in FORM_HINT_TOKENS:
            return True
        return False
    score = 0
    extra_words = product_words - input_words
    if extra_words:
        extra_words = {w for w in extra_words if not _is_structural(w)}
        numeric_re = re.compile('^\\d+(?:\\.\\d+)?(?:/\\d+(?:\\.\\d+)?)*$')
        extra_numeric = {w for w in extra_words if numeric_re.match(w)}
        extra_numeric_penalized = {w for w in extra_numeric if w not in input_numbers}
        extra_text = extra_words - extra_numeric
        score -= len(extra_text) * 40
        score -= len(extra_numeric_penalized) * 25
    if detected_variants:
        for var in detected_variants:
            if var != 'PLAIN' and _form_neutral_text(var, form_precedence) in product_comparison:
                score += 100
    if detected_sub_variants:
        for sv in detected_sub_variants:
            if sv != 'PLAIN' and _form_neutral_text(sv, form_precedence) in product_comparison:
                score += 50
    product_sub_var = 'PLAIN'
    strength_match = re.search('(\\d+(?:\\.\\d+)?)\\s*(?:MG|GM|ML|MCG|IU|G|KG)\\b', product_comparison)
    if strength_match:
        product_sub_var = strength_match.group(1)
    else:
        form_words = ['TABLET', 'TABLETS', 'CAPSULE', 'CAPSULES', 'TAB', 'TABS', 'CAP', 'CAPS', 'INJECTION', 'INJ', 'SYRUP', 'SYP', 'SUSPENSION', 'SUSP', 'DROP', 'DROPS', 'CREAM', 'OINTMENT', 'GEL', 'LOTION', 'POWDER', 'GRANULES', 'GRANUL']
        for form in form_words:
            pattern = '(\\d+(?:\\.\\d+)?)\\s+' + re.escape(form) + '\\b'
            m = re.search(pattern, product_comparison)
            if m:
                product_sub_var = m.group(1)
                break
    if product_sub_var != 'PLAIN' and product_sub_var not in ('', 'PLAIN'):
        strength_num = product_sub_var
        for m in re.finditer('(?<![\\d.])' + re.escape(strength_num) + '(?![\\d.])', input_upper):
            if is_position_in_pack_pattern(m.start(), m.end(), pack_positions):
                continue
            after = input_upper[m.end():m.end() + 3]
            if re.match('\\s*[STCP]', after) or re.match('\\s*[X*×]', after):
                continue
            score += 75
            break
    if product_sub_var != 'PLAIN' and product_sub_var not in ('', 'PLAIN'):
        strength_num = product_sub_var
        typical_tablet_counts = {'10', '15', '20', '30', '14', '21', '28', '7', '5', '30'}
        for m in re.finditer('(?<![\\d.])' + re.escape(strength_num) + '\\s*[X*×]\\s*(?:' + '|'.join(typical_tablet_counts) + ')(?![\\d.])', input_upper):
            before = input_upper[max(0, m.start() - 5):m.start()]
            after = input_upper[m.end():m.end() + 5]
            if not re.search('\\d\\s*$', before):
                if not re.match('\\s*[STCP]', after):
                    score += 100
                    break
    if 'SEMI' not in input_words and 'SEMI' in product_upper:
        score -= 20
    return score

def prioritize_by_product_name_match(items: list, input_name: str, verbose=False, _form_order=None) -> list:
    if len(items) <= 1:
        return items
    form_order = _form_precedence(input_name) if _form_order is None else _form_order
    scored_items = [(item, calculate_product_name_priority_score(item['product'], input_name, _form_order=form_order)) for item in items]
    scored_items.sort(key=lambda x: x[1], reverse=True)
    if verbose:
        print(f'\n  Product name prioritization scores:')
        for item, score in scored_items:
            print(f"    {score}: {item['product'][:50]}")
    return [item for item, _ in scored_items]

def filter_items_by_sub_variant(items: list, detected_sub_variants: set, input_name: str='', reference_items: list=None, verbose=False, _catalog_form_recovery=False) -> list:
    if verbose:
        print('\n' + '=' * 80)
        print('STEP 3: SUB-VARIANT-BASED FILTERING')
        print('=' * 80)
        print(f'Detected sub-variants: {sorted(detected_sub_variants)}')
    specific_detected = {sv for sv in detected_sub_variants if sv != 'PLAIN'}
    match_targets = specific_detected if specific_detected else set(detected_sub_variants)
    detected_forms = detect_dosage_form_in_input(input_name, verbose=False) if input_name else set()
    alpha_detected = {sv for sv in specific_detected if re.search('[A-Z]', str(sv)) and (not is_numeric_subvariant_like(str(sv)))}
    combo_dosages = {sv for sv in detected_sub_variants if '/' in str(sv)}
    if combo_dosages:
        combo_matches = []
        for item in items:
            for combo in combo_dosages:
                if product_contains_combination(item['product'], combo):
                    combo_matches.append(item)
                    break
        if combo_matches:
            if verbose:
                print(f'\n[OK] {len(combo_matches)} literal combo matches')
            return combo_matches
        ordered_components = []
        for combo in combo_dosages:
            parts = [p.strip() for p in str(combo).split('/') if p.strip()]
            for idx, part in enumerate(parts):
                if idx >= len(ordered_components):
                    ordered_components.append(set())
                ordered_components[idx].add(part.upper())
        input_combo_len = len(ordered_components)
        master_combo_lengths = set()
        for item in items:
            subv = str(item.get('sub_variant', 'PLAIN')).strip().upper()
            if '/' in subv:
                master_combo_lengths.add(len(subv.split('/')))
        is_2to3_combo = input_combo_len == 2 and any((l >= 3 for l in master_combo_lengths))
        if is_2to3_combo:
            if verbose:
                print(f'\n[2->3 COMBO] Input has 2 components, master has 3+ -- will match middle component')
            middle_components = set()
            for item in items:
                subv = str(item.get('sub_variant', 'PLAIN')).strip().upper()
                if '/' in subv:
                    parts = subv.split('/')
                    if len(parts) >= 3:
                        middle_components.add(parts[1].upper())
            if middle_components:
                comp_col_matches = [item for item in items if len(str(item.get('sub_variant', 'PLAIN')).strip().upper().split('/')) >= 3 if str(item.get('sub_variant', 'PLAIN')).strip().upper().split('/')[1].upper() in middle_components]
                if comp_col_matches:
                    if verbose:
                        print(f'\n[OK] 2->3 combo -- using middle component {sorted(middle_components)} (sub_variant column)')
                    middle_values = []
                    for item in comp_col_matches:
                        subv = str(item.get('sub_variant', 'PLAIN')).strip().upper()
                        parts = subv.split('/')
                        if len(parts) >= 3:
                            try:
                                middle_values.append((float(parts[1]), item))
                            except ValueError:
                                pass
                    if middle_values:
                        middle_values.sort(key=lambda x: x[0])
                        lowest_middle = middle_values[0][0]
                        lowest_items = [item for val, item in middle_values if val == lowest_middle]
                        if verbose:
                            print(f'\n[DEFAULT] Picking lowest middle component: {lowest_middle}')
                        if input_name:
                            return prioritize_by_product_name_match(lowest_items, input_name, verbose=verbose)
                        return lowest_items
        for position, components in enumerate(ordered_components):
            if not components:
                continue
            comp_col_matches = [item for item in items if str(item.get('sub_variant', 'PLAIN')).strip().upper() in components]
            if comp_col_matches:
                if verbose:
                    print(f'\n[OK] No literal combo -- using component #{position + 1} {sorted(components)} (sub_variant column)')
                if input_name:
                    return prioritize_by_product_name_match(comp_col_matches, input_name, verbose=verbose)
                return comp_col_matches
            comp_name_matches = []
            for item in items:
                pname = item['product'].upper()
                for comp in components:
                    if re.search('\\b' + re.escape(str(comp)) + '\\b', pname):
                        comp_name_matches.append(item)
                        break
            if comp_name_matches:
                if verbose:
                    print(f'\n[OK] No literal combo -- using component #{position + 1} {sorted(components)} (product name)')
                if input_name:
                    return prioritize_by_product_name_match(comp_name_matches, input_name, verbose=verbose)
                return comp_name_matches
        if len(ordered_components) >= 2:
            master_combo_lengths = set()
            for item in items:
                subv = str(item.get('sub_variant', 'PLAIN')).strip().upper()
                if '/' in subv:
                    master_combo_lengths.add(len(subv.split('/')))
            if any((l > len(ordered_components) for l in master_combo_lengths)):
                last_components = ordered_components[-1]
                if last_components:
                    comp_col_matches = [item for item in items if str(item.get('sub_variant', 'PLAIN')).strip().upper() in last_components]
                    if comp_col_matches:
                        if verbose:
                            print(f'\n[OK] Master has longer combo -- using last component {sorted(last_components)} (sub_variant column)')
                        if input_name:
                            return prioritize_by_product_name_match(comp_col_matches, input_name, verbose=verbose)
                        return comp_col_matches
                    comp_name_matches = []
                    for item in items:
                        pname = item['product'].upper()
                        for comp in last_components:
                            if re.search('\\b' + re.escape(str(comp)) + '\\b', pname):
                                comp_name_matches.append(item)
                                break
                    if comp_name_matches:
                        if verbose:
                            print(f'\n[OK] Master has longer combo -- using last component {sorted(last_components)} (product name)')
                        if input_name:
                            return prioritize_by_product_name_match(comp_name_matches, input_name, verbose=verbose)
                        return comp_name_matches
                if len(ordered_components) == 2 and len(comp_name_matches) > 1:
                    middle_values = []
                    for item in comp_name_matches:
                        subv = str(item.get('sub_variant', 'PLAIN')).strip().upper()
                        parts = subv.split('/')
                        if len(parts) == 3 and parts[0] == list(ordered_components[0])[0] and (parts[2] == list(ordered_components[-1])[0]):
                            try:
                                middle_values.append((float(parts[1]), item))
                            except ValueError:
                                pass
                    if middle_values:
                        middle_values.sort(key=lambda x: x[0])
                        lowest_middle = middle_values[0][0]
                        lowest_items = [item for val, item in middle_values if val == lowest_middle]
                        if verbose:
                            print(f'\n[DEFAULT] 2-component input matches multiple 3-component combos -- picking lowest middle component: {lowest_middle}')
                        if input_name:
                            return prioritize_by_product_name_match(lowest_items, input_name, verbose=verbose)
                        return lowest_items
    if alpha_detected:
        alpha_matches = [i for i in items if i.get('sub_variant', 'PLAIN').upper() in alpha_detected]
        if alpha_matches:
            narrowed = filter_by_subvariant_in_product_name(alpha_matches, detected_sub_variants, verbose=verbose)
            if input_name:
                return prioritize_by_product_name_match(narrowed, input_name, verbose=verbose)
            return narrowed
    exact_matches = []
    plain_products = []
    numeric_detected = {sv for sv in detected_sub_variants if sv != 'PLAIN' and is_numeric_subvariant_like(str(sv))}
    for item in items:
        item_sv = item.get('sub_variant', 'PLAIN').upper()
        if item_sv in match_targets:
            exact_matches.append(item)
        elif item_sv == 'PLAIN':
            plain_products.append(item)
    if exact_matches:
        if not specific_detected and detected_forms:
            form_matches = [item for item in exact_matches if detect_dosage_form_in_product(item['product'], _catalog_form_recovery) & detected_forms or forms_share_group(detect_dosage_form_in_product(item['product'], _catalog_form_recovery), detected_forms)]
            if form_matches:
                if input_name:
                    return prioritize_by_product_name_match(form_matches, input_name, verbose=verbose)
                return form_matches
        numeric_exact = {str(i.get('sub_variant', '')).strip() for i in exact_matches if is_numeric_subvariant_like(str(i.get('sub_variant', '')))}
        if input_name and len(numeric_exact) > 1:
            up = normalize_ocr_numeric_noise(input_name)

            def _pos(sub):
                m = re.search('(?<![\\d.])' + re.escape(str(sub)) + '(?![\\d.])', up)
                return m.start() if m else 10 ** 6
            best_pos = min((_pos(s) for s in numeric_exact))
            earliest = [i for i in exact_matches if _pos(str(i.get('sub_variant', '')).strip()) == best_pos]
            if earliest and len(earliest) < len(exact_matches):
                exact_matches = earliest
        if input_name:
            return prioritize_by_product_name_match(exact_matches, input_name, verbose=verbose)
        return exact_matches
    if numeric_detected:
        first_comp = [item for item in items if '/' in str(item.get('sub_variant', '')) and str(item.get('sub_variant', '')).split('/')[0].strip().upper() in numeric_detected]
        if first_comp and len(first_comp) < len(items):
            first_comp_sub_vars = {str(item.get('sub_variant', '')).strip().upper() for item in first_comp}
            if len(first_comp_sub_vars) > 1:
                if verbose:
                    print(f'\n[WARN] Multiple combo sub-variants share first component {sorted(numeric_detected)} -- not narrowing')
            else:
                if verbose:
                    print(f"\n[OK] Single number -> combo first-component match: {[i['product'] for i in first_comp]}")
                if input_name:
                    return prioritize_by_product_name_match(first_comp, input_name, verbose=verbose)
                return first_comp
    if numeric_detected:
        name_matches = filter_by_subvariant_in_product_name(items, detected_sub_variants, verbose=verbose)
        if name_matches and len(name_matches) < len(items):
            if input_name:
                return prioritize_by_product_name_match(name_matches, input_name, verbose=verbose)
            return name_matches
    if numeric_detected and plain_products:
        name_matches = filter_by_subvariant_in_product_name(plain_products, detected_sub_variants, verbose=verbose)
        if name_matches and len(name_matches) < len(plain_products):
            return name_matches
    if plain_products:
        return plain_products
    return items

def _dosage_form_mentions(text: str, product=False) -> list:
    """Bounded aliases in printed order; retain injection presentation identity."""
    upper = str(text or '').upper()
    mentions = []
    for match in _DOSAGE_FORM_PATTERN.finditer(upper):
        alias = match.group().rstrip('.')
        # This catalog typo remains restricted to the existing scoped recovery.
        if product and alias == 'CAPSULS':
            continue
        if not product and alias == 'SY' and upper[match.end():match.end() + 1].isdigit():
            continue  # Legacy SY1X1 pack notation is not an explicit form token.
        form = _DOSAGE_FORM_ALIASES.get(alias, 'INJECTION')
        if not product and alias == 'SYR' and match.start() and upper[match.start() - 1].isdigit():
            form = 'INJECTION'  # Preserve the legacy counted-syringe interpretation.
        mentions.append((match.start(), match.end(), form))
    return mentions


def detect_dosage_form_in_input(input_name: str, verbose=False) -> set:
    detected_forms = {form for _, _, form in _dosage_form_mentions(input_name)}
    input_upper = normalize_ocr_numeric_noise(input_name).upper()
    # Keep unrelated topical/bar detection unchanged.
    for pattern, form in [(r'\bCREAMS?\b|\bCRM\.?\b', 'CREAM'),
                          (r'\bOINTMENTS?\b|\bOI\.?\b|\bOINT\.?\b', 'OINTMENT'),
                          (r'\bGELS?\b', 'GEL'), (r'\bBARS?\b|\bSYNDET\b', 'BAR'),
                          (r'\bLOTIONS?\b|\bLOT\.?\b|\bLO\b', 'LOTION')]:
        if re.search(pattern, input_upper):
            detected_forms.add(form)
    if detected_forms & {'VIAL', 'AMPULE'}:
        detected_forms.add('INJECTION')
    return detected_forms


def detect_dosage_form_in_product(product_name: str, _catalog_form_recovery=False) -> set:
    product_upper = product_name.upper()
    if _catalog_form_recovery:
        product_upper = re.sub(r'\bCAPSULS\b', 'CAPSULES', product_upper)
        product_upper = re.sub(r'\bMGINJECTION\b', 'MG INJECTION', product_upper)
    detected_forms = {form for _, _, form in _dosage_form_mentions(product_upper, product=True)}
    for pattern, form in [(r'\bCREAMS?\b', 'CREAM'), (r'\bOINTMENTS?\b', 'OINTMENT'),
                          (r'\bGELS?\b', 'GEL'), (r'\bBARS?\b|\bSYNDET\b', 'BAR'),
                          (r'\bLOTIONS?\b', 'LOTION')]:
        if re.search(pattern, product_upper):
            detected_forms.add(form)
    if detected_forms & {'VIAL', 'AMPULE'}:
        detected_forms.add('INJECTION')
    return detected_forms


_FORM_GROUPS = {'TABLET': 'solid_oral', 'CAPSULE': 'solid_oral', 'SUSPENSION': 'liquid_oral', 'SYRUP': 'liquid_oral', 'LIQUID': 'liquid_oral', 'SOLUTION': 'liquid_oral', 'DROPS': 'liquid_oral', 'TONIC': 'liquid_oral', 'INJECTION': 'parenteral', 'VIAL': 'parenteral', 'AMPULE': 'parenteral', 'CREAM': 'topical', 'GEL': 'topical', 'OINTMENT': 'topical', 'LOTION': 'topical', 'POWDER': 'powder', 'GRANULES': 'powder', 'SACHET': 'powder', 'BAR': 'bar'}
_EQUIVALENT_FORM_GROUPS = {'solid_oral', 'liquid_oral', 'parenteral', 'powder'}
_NEUTRAL_GROUP_FORM = {'solid_oral': 'TABLET', 'liquid_oral': 'SUSPENSION', 'parenteral': 'INJECTION', 'powder': 'POWDER'}


def _form_precedence(input_name: str) -> tuple:
    forms = detect_dosage_form_in_input(input_name)
    groups = {_FORM_GROUPS[form] for form in forms if form in _FORM_GROUPS}
    # First-mention precedence is defined only within one equivalence group.
    if len(groups) != 1 or not groups.issubset(_EQUIVALENT_FORM_GROUPS):
        return ()
    return tuple(dict.fromkeys(form for _, _, form in _dosage_form_mentions(input_name)))


def _form_neutral_text(text: str, precedence: tuple, product=False) -> str:
    if not precedence:
        return text
    group = _FORM_GROUPS[precedence[0]]
    neutral = _NEUTRAL_GROUP_FORM[group]
    upper = str(text or '').upper()
    parts, end = [], 0
    for start, stop, form in _dosage_form_mentions(upper, product=product):
        if _FORM_GROUPS[form] == group:
            parts.extend((upper[end:start], neutral))
            end = stop
    parts.append(upper[end:])
    return ''.join(parts)


def _final_form_preference(product_name: str, precedence: tuple) -> int:
    if not precedence:
        return 0
    product_forms = {form for _, _, form in _dosage_form_mentions(product_name, product=True)}
    return next((len(precedence) - i for i, form in enumerate(precedence) if form in product_forms), 0)

def forms_share_group(product_forms: set, detected_forms: set) -> bool:
    product_groups = {_FORM_GROUPS[f] for f in product_forms if f in _FORM_GROUPS}
    detected_groups = {_FORM_GROUPS[f] for f in detected_forms if f in _FORM_GROUPS}
    return bool(product_groups & detected_groups)

def prioritize_by_dosage_form(items: list, detected_forms: set, verbose=False, _catalog_form_recovery=False) -> list:
    if not detected_forms:
        return items
    if verbose:
        print('\n' + '=' * 80)
        print('STEP 4.5: DOSAGE FORM PRIORITIZATION')
        print('=' * 80)
        print(f'Detected forms in input: {sorted(detected_forms)}')
    detected_groups = {_FORM_GROUPS[f] for f in detected_forms if f in _FORM_GROUPS}
    if detected_groups & _EQUIVALENT_FORM_GROUPS:
        # Keep every recognized equivalent, without promoting an unknown spelling
        # ahead of recognized forms. Scoped catalog-typo recovery remains intact.
        equivalent = []
        for item in items:
            product_forms = detect_dosage_form_in_product(item['product'], _catalog_form_recovery)
            product_groups = {_FORM_GROUPS[form] for form in product_forms}
            if product_forms & detected_forms or product_groups & detected_groups & _EQUIVALENT_FORM_GROUPS:
                equivalent.append(item)
        if equivalent:
            return equivalent
        compatible = _form_compatible(items, detected_forms, _catalog_form_recovery)
        if compatible:
            return compatible
        return items  # Retain the existing no-compatible-form fallback.
    exact_form_matches = []
    other_products = []
    for item in items:
        product_forms = detect_dosage_form_in_product(item['product'], _catalog_form_recovery)
        if product_forms & detected_forms:
            exact_form_matches.append(item)
        else:
            other_products.append(item)
    if exact_form_matches:
        if verbose:
            print(f'\n[OK] Found {len(exact_form_matches)} products matching dosage form')
        return exact_form_matches
    if detected_groups:
        non_conflicting = []
        excluded = []
        for item in items:
            product_forms = detect_dosage_form_in_product(item['product'], _catalog_form_recovery)
            product_groups = {_FORM_GROUPS[f] for f in product_forms if f in _FORM_GROUPS}
            if product_groups and product_groups.isdisjoint(detected_groups):
                excluded.append(item)
            else:
                non_conflicting.append(item)
        if non_conflicting and len(non_conflicting) < len(items):
            if verbose:
                print(f'\n[WARN] No exact form match -- excluded {len(excluded)} conflicting-form products')
            return non_conflicting
    return items

def _form_compatible(items: list, detected_forms: set, _catalog_form_recovery=False) -> list:
    detected_groups = {_FORM_GROUPS[f] for f in detected_forms if f in _FORM_GROUPS}
    if not detected_groups:
        return list(items)
    keep = []
    for item in items:
        product_groups = {_FORM_GROUPS[f] for f in detect_dosage_form_in_product(item['product'], _catalog_form_recovery) if f in _FORM_GROUPS}
        if not product_groups or product_groups & detected_groups:
            keep.append(item)
    return keep

def _pack_selection_input(input_name, items):
    """Keep the lowest explicit compatible pack, without discarding strengths."""
    matches = list(_EXPLICIT_PACK_SOURCE.finditer(input_name))
    if len(matches) < 2 or len({(item.get('variant'), item.get('sub_variant')) for item in items}) != 1:
        return input_name
    dimensions = {pack[1] for item in items if (pack := parse_master_pack(item.get('pack_size', '')))}
    count_dimensions = {dimension for dimension in dimensions if dimension.startswith('COUNT')}
    quantities = []
    for match in matches:
        pack = parse_master_pack(match[0].replace('×', '*'))
        if any(separator in match[0].upper() for separator in ('X', '*', '×')):
            # Input multipliers retain the existing extraction interpretation;
            # the master-pack parser must not substitute a stock aggregate.
            pack = _extract_pack_size(match[0], with_measurement=True) or pack
        if pack and pack[1] == 'COUNT' and len(count_dimensions) == 1:
            pack = (pack[0], next(iter(count_dimensions)))
        if pack and pack[1] in dimensions:
            quantities.append((match, pack))
    # Different dimensions cannot compete for the lowest pack.
    if len(quantities) < 2 or len({pack[1] for _, pack in quantities}) != 1:
        return input_name
    lowest = min(pack[0] for _, pack in quantities)
    for match, pack in reversed(quantities):
        if pack[0] != lowest:
            input_name = input_name[:match.start()] + ' ' * (match.end() - match.start()) + input_name[match.end():]
    return input_name


def _input_pack_measurement(input_name: str, items: list, *, _prefer_volume=False):
    input_name = _pack_selection_input(input_name, items)
    requested = _extract_pack_size(input_name, with_measurement=True)
    if _prefer_volume:
        requested = recover_explicit_volume(input_name, requested)
    if requested and requested[1] != 'COUNT':
        dimensions = {pack[1] for item in items if (pack := parse_master_pack(item.get('pack_size', '')))}
        if dimensions and requested[1] not in dimensions:
            # A content mass (20G/50ML) must not hide the catalog's liquid pack.
            compatible_text = _PACK_MEASUREMENT.sub(
                lambda match: match.group() if (value := measurement(*match.groups())) and value[1] in dimensions else ' ',
                input_name.upper())
            compatible = _extract_pack_size(compatible_text, with_measurement=True)
            if compatible and compatible[1] in dimensions:
                return compatible
        return requested
    if not requested and not any(unit in input_name.upper() for unit in ('MG', 'MCG', 'IU', 'UNIT')):
        return None
    dimensions = {pack[1] for item in items if (pack := parse_master_pack(item.get('pack_size', '')))}
    if requested and requested[1] == 'COUNT' and not any(d.startswith('COUNT') for d in dimensions):
        return None  # Unitless multipliers do not establish a liquid/mass quantity.
    if requested:
        return requested
    # MG/MCG/IU can be packaging only when the catalog establishes that dimension
    # and the printed number has not already been used as a numeric sub-variant.
    if dimensions & {'MASS', 'ACTIVITY'}:
        for match in re.finditer(r'(?<![A-Z0-9.])(\d+(?:\.\d+)?|\.\d+)\s*(MCG|MG|IU|UNIT)\b', input_name.upper()):
            if any(is_numeric_subvariant_like(item.get('sub_variant', '')) for item in items):
                continue
            candidate = measurement(match[1], match[2])
            if candidate and candidate[1] in dimensions:
                return candidate
    return None


_PACK_BRAND_CACHE = (None, {})


def _pack_source_brands(brand_map):
    global _PACK_BRAND_CACHE
    cached_map, sources = _PACK_BRAND_CACHE
    if cached_map is not brand_map:
        sources = {id(item): brand for brand, items in brand_map.items() for item in items}
        _PACK_BRAND_CACHE = (brand_map, sources)
    return sources


def _input_variant_tokens_outside_packs(input_name, brand, variants):
    """Do not let letters printed only inside packs establish another variant."""
    descriptor = strip_brand_from_input(normalize_ocr_numeric_noise(input_name), brand, variants)
    for start, end in reversed(get_all_pack_positions(descriptor)):
        if end and descriptor[end - 1].isdigit():
            suffix = re.match(r'\s*[STCP]\b', descriptor[end:])
            if suffix:
                end += suffix.end()
        descriptor = descriptor[:start] + ' ' * (end - start) + descriptor[end:]
    return extract_variant_match_tokens(descriptor)


def _explicit_variant_items_before_pack(items, input_name, brand):
    variants = {str(item.get('variant', 'PLAIN')).upper() for item in items}
    tokens = _input_variant_tokens_outside_packs(input_name, brand, variants)
    explicit = [item for item in items if (required := _variant_identity_tokens(str(item.get('variant', 'PLAIN')).upper()))
                and required.issubset(tokens)]
    return explicit if explicit else items


def _prepare_diof_candidates(text, items):
    """Fence the ordinary family before retries can discard DS or descriptors."""
    variants = {item['variant'] for item in items}
    tokens = _input_variant_tokens_outside_packs(text, 'DIOF', variants)
    known = _pack_form_skip_tokens() | {'DS', 'PLAIN', 'LIQUID', 'PACK', 'OF', 'X', 'L', 'G', 'KG'}
    unknown = {word for token in tokens for word in re.findall('[A-Z]+', token)} - known
    if unknown:
        return [], 'VARIANT_NOT_IN_MASTER'
    identity = 'DS' if 'DS' in tokens else 'PLAIN'
    selected = [item for item in items if item['variant'] == identity]
    forms = detect_dosage_form_in_input(text, verbose=False)
    if forms:
        selected = [item for item in selected
                    if all(forms_share_group(detect_dosage_form_in_product(item['product']), {form}) for form in forms)]
    if forms & {'SUSPENSION', 'SYRUP', 'SOLUTION'} and any(
            measurement(*match.groups())[1] != 'VOLUME'
            for match in _PACK_MEASUREMENT.finditer(text) if measurement(*match.groups())):
        return [], 'NO_COMPATIBLE_PACK'
    return (selected, '') if selected else ([], 'VARIANT_NOT_IN_MASTER')


@lru_cache(maxsize=32)
def _maxtra_identity_vocabulary(attributes):
    return frozenset(token for value in attributes for token in extract_variant_match_tokens(value)) - {'PLAIN'}


@lru_cache(maxsize=256)
def _maxtra_variant_identity(value):
    """Catalog identity is immutable; do not retokenize it for every source row."""
    return _variant_identity_tokens(value)


_MAXTRA_BRAND_CACHE = (None, None)


def _prepare_maxtra_source(text, brand_map, forced_brand, auto_detect):
    """Preserve approved source evidence before generic destructive cleaning."""
    upper = text.upper()
    if str(forced_brand or '').upper() != 'MAXTRA' and 'MAXTRA' not in upper and not upper.startswith('MAXT'):
        return text
    items = brand_map.get('MAXTRA', ())
    if not items or forced_brand and str(forced_brand).upper() != 'MAXTRA':
        return text
    variants = tuple(sorted({item['variant'] for item in items}))
    leading = re.match(r'^\s*([A-Z]+)(?=$|[\s-])', text, re.I)
    if leading:
        token = leading[1].upper()
        # MAXTRAP is a supported glued identity, not a misspelled brand.
        glued = token.startswith('MAXTRA') and token[6:] in _maxtra_identity_vocabulary(variants)
        if token != 'MAXTRA' and token not in brand_map and not glued:
            from garbage_check import edit_distance_leq_one, build_brand_index, exact_product_match, fuzzy_product_match
            if edit_distance_leq_one(token, 'MAXTRA'):
                confirmed = str(forced_brand or '').upper() == 'MAXTRA'
                if not confirmed and auto_detect:
                    global _MAXTRA_BRAND_CACHE
                    cached_map, index = _MAXTRA_BRAND_CACHE
                    if cached_map is not brand_map:
                        index = build_brand_index(brand_map)
                        _MAXTRA_BRAND_CACHE = (brand_map, index)
                    match = exact_product_match(text, index) or fuzzy_product_match(text, index)
                    confirmed = bool(match and match['brand'] == 'MAXTRA' and match['matched_text'].upper() == token)
                if confirmed:
                    text = text[:leading.start(1)] + 'MAXTRA' + text[leading.end(1):]
    return normalize_maxtra_description(text, variants)


def _prepare_maxtra_candidates(input_name, items):
    """Establish approved MAXTRA identity/form before the normal filters run.

    Return original master rows. Quantity interpretation and SKU selection stay
    in the existing parser/universal selector; PS is an explicit family rule.
    """
    variants, subvariants, _ = extract_available_attributes_from_items(items)
    text = input_name
    if NUMERIC_OFFER.search(text):
        return text, [], True, 'VARIANT_NOT_IN_MASTER'
    if tablet_count_evidence(text):
        text += ' TABLET'
    tokens = _input_variant_tokens_outside_packs(text, 'MAXTRA', variants)
    identity_vocab = _maxtra_identity_vocabulary(tuple(sorted(variants)))
    identity = tokens & identity_vocab
    forms = detect_dosage_form_in_input(text)
    # Complete O/S SPRAY only when the loaded catalog has no other such family.
    for token in ('O', 'S'):
        if identity == {token} and re.search(r'\b' + token + r'[\s-]+SPRAY\b', text, re.I):
            family = [item for item in items if token in _maxtra_variant_identity(item['variant'])]
            if family and all(re.search(r'\bNASAL\s+SPRAY\b', item['product'], re.I) for item in family):
                identity = identity | {'NASAL'}
                text = re.sub(r'\b' + token + r'[\s-]+SPRAY\b', token + ' NASAL SPRAY', text, flags=re.I)
    liquids = {'SYRUP', 'SUSPENSION', 'SOLUTION', 'LIQUID', 'DROPS', 'TONIC'}
    liquid_identity = (identity in ({'GARGLE'}, {'P', 'DS'}, {'S', 'NASAL'}, {'S', 'NASAL', 'PS'}, {'O', 'NASAL'})
                       or identity in ({'P'}, {'ORAL'}, {'P', 'ORAL'}, {'HS'}, {'HS', 'ORAL'}, {'C'}, {'N'}, {'DC'}, {'LM'})
                       and bool(forms & liquids) and not forms - liquids
                       or not identity and forms == {'DROPS'})
    if liquid_identity:
        family = [item for item in items if _maxtra_variant_identity(item['variant']) - {'PS'} == identity - {'PS'}]
        sizes = {pack[0] for item in family if (pack := parse_master_pack(item.get('pack_size', ''))) and pack[1] == 'VOLUME'}
        normalized = normalize_liquid_quantity(text, sizes, nasal_spray='NASAL' in identity and bool(re.search(r'\bSPRAY\b', text, re.I)))
        if normalized != text:
            text = normalized
            # The generic 15 ML default must see the interpreted unit, not an
            # obsolete M qualifier left over from the pre-normalization scan.
            tokens = _input_variant_tokens_outside_packs(text, 'MAXTRA', variants)
    requested = _input_pack_measurement(text, items, _prefer_volume=True)
    if requested and requested[1] in {'MASS', 'ACTIVITY'} and not any(
            (pack := parse_master_pack(item.get('pack_size', ''))) and pack[1] == requested[1] for item in items):
        return text, [], True, 'NO_COMPATIBLE_PACK'
    if requested is None and any(measurement(*match.groups()) for match in _PACK_MEASUREMENT.finditer(text.upper())):
        # An unsupported physical measurement must not masquerade as a missing
        # quantity and reach a blank-pack liquid/tablet default.
        return text, [], True, 'NO_COMPATIBLE_PACK'
    # Preserve the old bare-brand path when no descriptor or form is supplied.
    if not (tokens - {'PLAIN'} or forms):
        return text, items, False, ''
    if identity == {'S'} and requested == (20, 'VOLUME'):
        s_items = [item for item in items if 'S' in _maxtra_variant_identity(item['variant'])]
        s_identities = {_maxtra_variant_identity(item['variant']) - {'PS'} for item in s_items}
        s_twenty = [item for item in s_items if parse_master_pack(item.get('pack_size', '')) == (20, 'VOLUME')]
        # The approved omission is limited to the unique S 20 ML presentation.
        # Do not infer a nasal identity when the catalog has another S family.
        if (len(s_twenty) == 1 and _maxtra_variant_identity(s_twenty[0]['variant']) == {'S', 'NASAL', 'PS'}
                and s_identities == {frozenset({'S', 'NASAL'})} and all(
                re.search(r'\bNASAL\s+SPRAY\b', item['product'], re.I) for item in s_items)):
            identity = identity | {'NASAL'}
            text += ' NASAL SPRAY'
    inferred_ps = False
    if ({'S', 'NASAL'} <= identity and not identity & {'O', 'P', 'HS', 'N', 'ORAL'}
            or identity == {'O', 'NASAL'}):
        # Approved O/S NASAL drops terminology stays in its own spray family.
        text = re.sub(r'\b(?:DROPS?|DR)\b', 'SPRAY', text, flags=re.I)
        if 'S' in identity and requested == (20, 'VOLUME') and 'PS' not in identity:
            identity = identity | {'PS'}
            text += ' PS'
            inferred_ps = True
        forms = detect_dosage_form_in_input(text)
    elif (identity == {'HS'} and forms & {'DROPS', 'SYRUP', 'SUSPENSION', 'LIQUID', 'SOLUTION', 'TONIC'}
          and {variant for variant in variants if 'HS' in _maxtra_variant_identity(variant)} == {'HS ORAL'}):
        identity = identity | {'ORAL'}
        text += ' ORAL'
    elif (not identity and 'DROPS' in forms and requested == (15, 'VOLUME')
          and not (tokens - _pack_form_skip_tokens() - {'PLAIN'})):
        identity = {'ORAL'}
        text += ' ORAL'

    # Unlike the legacy fail-open compound filter, a printed family token may
    # not disappear into a different family during a low-confidence retry.
    selected = [item for item in items
                if _maxtra_variant_identity(item['variant']) == identity]
    if inferred_ps and not any(parse_master_pack(item.get('pack_size', '')) == (20, 'VOLUME') for item in selected):
        return text, [], True, 'MISSING_MAXTRA_TARGET'
    sub_identity = tokens & _maxtra_identity_vocabulary(tuple(sorted(subvariants)))
    selected = [item for item in selected
                if _maxtra_variant_identity(str(item['sub_variant'])) == sub_identity]
    upper = normalize_ocr_numeric_noise(text)
    if re.search(r'\bGARGLE\b', upper):
        selected = [item for item in selected if re.search(r'\bGARGLE\b', item['product'], re.I)]
        if forms - {'SYRUP', 'SUSPENSION', 'LIQUID', 'SOLUTION'}:
            selected = []
    elif re.search(r'\bNASAL\s+SPRAY\b', upper):
        selected = [item for item in selected if re.search(r'\bNASAL\s+SPRAY\b', item['product'], re.I)]
        if forms:
            selected = []  # A second, incompatible printed form is not erased.
    elif forms:
        # Keep the existing syrup/liquid and solid-oral aliases. Explicit drops
        # remain distinct here: P syrup must not win over a missing P drops row.
        non_drop_liquids = {'SYRUP', 'SUSPENSION', 'LIQUID', 'SOLUTION', 'TONIC'}

        def compatible(product):
            product_forms = detect_dosage_form_in_product(product)
            return all(bool(product_forms & non_drop_liquids) if form in non_drop_liquids
                       else bool(product_forms & {'TABLET', 'CAPSULE'}) if form in {'TABLET', 'CAPSULE'}
                       else form in product_forms for form in forms)

        family_items = selected
        selected = [item for item in family_items if compatible(item['product'])]
        if not selected and 'ORAL' in identity:
            # Confirmed ORAL/HS identity retains the existing liquid-form
            # equivalence when its catalog offers only the drops presentation.
            selected = [item for item in family_items
                        if all(forms_share_group(detect_dosage_form_in_product(item['product']), {form}) for form in forms)]
    # Scan beyond the first form as well: later unknown/conflicting descriptors
    # cannot be dropped by the older qualifier retry.
    printed = _input_variant_tokens_outside_packs(text, 'MAXTRA', variants)
    known = _pack_form_skip_tokens() | {'PLAIN', 'SPRAY', 'NASAL', 'ORAL', 'GARGLE', 'LIQUID', 'BOTTLE', 'BOTTLES', 'PACK', 'OF', 'X', 'UNIT', 'UNITS', 'L', 'G', 'KG'}
    known |= identity_vocab | _maxtra_identity_vocabulary(tuple(sorted(subvariants)))
    unknown = {word for token in printed for word in re.findall('[A-Z]+', token)} - known
    if unknown or not selected:
        return text, [], True, 'VARIANT_NOT_IN_MASTER'
    return text, selected, True, ''


def filter_items_by_pack_size(items: list, input_pack_size: str, reference_items: list=None, verbose=False,
                              _keep_form_ties=False, _measurement=None, _legacy=False) -> list:
    if _legacy or not input_pack_size:
        return _filter_items_by_pack_size_legacy(items, input_pack_size, reference_items, verbose, _keep_form_ties)
    numbered, blank = [], []
    units = set()
    for item in items:
        pack = str(item.get('pack_size', '')).strip()
        if not pack or pack.upper() in {'NAN', 'NONE'}:
            blank.append(item)
        elif parsed := parse_master_pack(pack):
            numbered.append((item, parsed))
            units.add(re.sub(r'[^A-Z]', '', pack.upper()) or 'COUNT')
        else:
            return []  # A malformed master pack is not a valid base presentation.
    requested = _measurement or parse_master_pack(str(input_pack_size))
    if requested is None and re.fullmatch(r'\d+(?:\.\d+)?', str(input_pack_size)):
        if len(units) == 1:
            requested = measurement(input_pack_size, next(iter(units)))
        elif not numbered:
            requested = measurement(input_pack_size, 'COUNT')
    if requested is None:
        return []
    identities = {(str(item.get('variant', 'PLAIN')).upper(), str(item.get('sub_variant', 'PLAIN')).upper()) for item in items}
    if len(identities) != 1:
        return []
    if blank:
        form_groups = {
            'VOLUME': {_FORM_GROUPS['SYRUP'], _FORM_GROUPS['INJECTION']},
            'COUNT_ORAL': {_FORM_GROUPS['TABLET']},
            'COUNT_INJECTION': {_FORM_GROUPS['INJECTION']},
        }.get(requested[1])
        if form_groups:
            blank = [item for item in blank
                     if not (groups := {_FORM_GROUPS[form] for form in detect_dosage_form_in_product(item['product'])})
                     or groups & form_groups]
    selected, refusal = select_pack_rows(numbered, blank, requested)
    if verbose:
        print(f'  [PACK] Exact / lower / base / lowest: {len(items)} -> {len(selected)} ({refusal or requested})')
    return selected


def _filter_items_by_pack_size_legacy(items: list, input_pack_size: str, reference_items: list=None, verbose=False, _keep_form_ties=False) -> list:
    if verbose:
        print('\n' + '=' * 80)
        print('STEP 5: PACK SIZE FILTERING')
        print('=' * 80)
        print(f"Input pack size: '{input_pack_size or 'NONE'}'")
    products_with_pack = []
    products_without_pack = []
    for item in items:
        pack = item['pack_size'].strip()
        if pack and pack not in ('', 'nan', 'NaN', 'None'):
            pack_num = parse_pack_size_value(pack)
            if pack_num is not None:
                products_with_pack.append((item, pack_num))
            else:
                products_without_pack.append(item)
        else:
            products_without_pack.append(item)
    if verbose:
        print(f'  Products with pack size:    {len(products_with_pack)}')
        print(f'  Products without pack size: {len(products_without_pack)}')
    if not input_pack_size:
        if products_without_pack:
            if verbose:
                print(f'\n[OK] Input has no pack -- returning {len(products_without_pack)} BASE products')
            return products_without_pack
        if products_with_pack:
            products_with_pack.sort(key=lambda x: x[1])
            if _keep_form_ties:
                return [item for item, pn in products_with_pack if pn == products_with_pack[0][1]]
            return [item for item, _ in products_with_pack[:1]]
        return items
    try:
        input_pack_num = float(input_pack_size)
    except ValueError:
        return products_without_pack if products_without_pack else items
    exact_matches = [item for item, pn in products_with_pack if pn == input_pack_num]
    if exact_matches:
        return exact_matches
    if products_without_pack:
        if verbose:
            print(f'\n[OK] No exact pack match -- falling back to BASE products')
        return products_without_pack
    if products_with_pack:
        closest = min(products_with_pack, key=lambda x: abs(x[1] - input_pack_num))
        closest_diff = abs(closest[1] - input_pack_num)
        if closest_diff <= input_pack_num * 0.5 or closest_diff <= 20:
            if verbose:
                print(f'\n[OK] No exact pack match and no base products -- using closest pack {closest[1]} (diff={closest_diff})')
            if _keep_form_ties:
                return [item for item, pn in products_with_pack if pn == closest[1]]
            return [closest[0]]
    return items

def _clean_input_for_fuzzy(input_name: str) -> str:
    s_upper = normalize_ocr_numeric_noise(input_name).upper()
    pack_positions = get_all_pack_positions(s_upper)
    _TOKEN_SPLIT = '[\\s\\-+/]+'
    tokens = []
    for m in re.finditer('\\S+', s_upper):
        token = m.group()
        is_pack = any((ps <= m.start() < pe or ps < m.end() <= pe for ps, pe in pack_positions))
        if not is_pack:
            tokens.append(token)
    return ' '.join(tokens)

def rank_local_candidates(items: list, input_name: str, detected_variants: set=None, detected_sub_variants: set=None) -> list:
    if not items:
        return []
    form_precedence = _form_precedence(input_name)
    ranked_items = prioritize_by_product_name_match(items, input_name, verbose=False, _form_order=form_precedence)
    if len(ranked_items) == 1:
        return [(100.0, ranked_items[0])]
    input_clean = _clean_input_for_fuzzy(_form_neutral_text(input_name, form_precedence))
    input_tokens_upper = set(re.findall('[A-Z0-9]+', input_clean.upper())) | set(re.findall('[A-Z]+', input_clean.upper()))
    KNOWN_GHOST_TOKENS = {token for candidate in ranked_items for field in ('variant', 'sub_variant') for token in re.findall('[A-Z]+', _form_neutral_text(str(candidate.get(field, '')).upper(), form_precedence, product=True)) if token != 'PLAIN' and token not in FORM_HINT_TOKENS and (token not in DOSAGE_UNIT_SUFFIXES)}
    scored_items = []
    for item in ranked_items:
        product_text = _form_neutral_text(item['product'], form_precedence, product=True)
        product_norm = norm(product_text)
        fuzzy_score = max(fuzz.token_sort_ratio(input_clean, product_norm), fuzz.token_set_ratio(input_clean, product_norm), fuzz.ratio(input_clean, product_norm))
        priority_score = calculate_product_name_priority_score(item['product'], input_name, detected_variants, detected_sub_variants, _form_order=form_precedence)
        alpha_ghost_penalty = 0
        item_sv = str(item.get('sub_variant', 'PLAIN')).strip().upper() or 'PLAIN'
        item_sv = _form_neutral_text(item_sv, form_precedence, product=True)
        if item_sv != 'PLAIN' and re.match('^[A-Z]+$', item_sv) and (item_sv not in input_tokens_upper):
            alpha_ghost_penalty = -200
        item_v = str(item.get('variant', 'PLAIN')).strip().upper() or 'PLAIN'
        item_v = _form_neutral_text(item_v, form_precedence, product=True)
        if item_v != 'PLAIN' and re.match('^[A-Z ]+$', item_v):
            v_words = set(item_v.split())
            if not v_words.issubset(input_tokens_upper):
                alpha_ghost_penalty -= 150
        product_tokens = set(re.findall('[A-Z]+', product_text.upper()))
        ghost_in_product = (product_tokens & KNOWN_GHOST_TOKENS) - input_tokens_upper
        if ghost_in_product:
            alpha_ghost_penalty -= 120 * len(ghost_in_product)
        score = fuzzy_score + priority_score + alpha_ghost_penalty
        scored_items.append((score, item))
    # A form preference never outweighs identity/strength/pack suitability.
    scored_items.sort(key=lambda x: (x[0], _final_form_preference(x[1]['product'], form_precedence)), reverse=True)
    return scored_items

def choose_best_local_candidate(items: list, input_name: str, detected_variants: set=None, detected_sub_variants: set=None) -> dict:
    ranked = rank_local_candidates(items, input_name, detected_variants, detected_sub_variants)
    return ranked[0][1] if ranked else None

def has_ambiguous_missing_strength(items: list, detected_sub_variants: set) -> bool:
    specific_detected = {sv for sv in detected_sub_variants if sv != 'PLAIN'}
    item_sub_variants = {str(item.get('sub_variant', 'PLAIN')).strip().upper() for item in items if str(item.get('sub_variant', 'PLAIN')).strip().upper() not in ('', 'PLAIN')}
    if len(item_sub_variants) <= 1:
        return False
    if not all((is_numeric_subvariant_like(sub_variant) for sub_variant in item_sub_variants)):
        return False
    item_variants = {str(item.get('variant', 'PLAIN')).strip().upper() for item in items}
    if len(item_variants) == 1:
        combo_subs = [sv for sv in item_sub_variants if '/' in sv]
        if len(combo_subs) >= 2:
            first_comps = {sv.split('/')[0].strip().upper() for sv in combo_subs}
            if len(first_comps) == 1:
                detected_combos = {sv for sv in detected_sub_variants if '/' in str(sv)}
                if not detected_combos:
                    return True
    if specific_detected and len(item_variants) == 1:
        combo_subs = [sv for sv in item_sub_variants if '/' in sv]
        if len(combo_subs) >= 2:
            first_comps = {sv.split('/')[0].strip().upper() for sv in combo_subs}
            if specific_detected.issubset(first_comps):
                detected_combos = {sv for sv in detected_sub_variants if '/' in str(sv)}
                if not detected_combos:
                    return True
    return False
TOP_N_SUGGESTIONS = 3
NO_BRAND_RECOVERY_MAX_EDITS = 1

def nearest_brand_within_edits(input_name, suggestions, max_edits=NO_BRAND_RECOVERY_MAX_EDITS):
    query = compact_brand_token(extract_brand_like_query(input_name))
    if not query:
        return (None, None)
    best, best_dist = (None, None)
    for s in suggestions or []:
        cand = compact_brand_token(s.get('brand', ''))
        if not cand:
            continue
        dist = Levenshtein.distance(query, cand)
        if best_dist is None or dist < best_dist:
            best, best_dist = (s, dist)
    if best is not None and best_dist is not None and (best_dist <= max_edits):
        return (best, best_dist)
    return (None, best_dist)

def extract_pack_notation(input_name: str) -> str:
    s_upper = normalize_ocr_numeric_noise(input_name).upper()
    m = re.search('\\b(\\d+)\\s*[X*×]\\s*(\\d+(?:\\.\\d+)?)(?=[A-Z]?\\b)', s_upper)
    if m:
        n1, n2 = m.groups()
        after = s_upper[m.end():m.end() + 5]
        form_suffix = ''
        if after:
            form_match = re.match('^([STCP]|PCS?|AB)', after)
            if form_match:
                form_suffix = form_match.group(1)
        return f'{n1}X{n2}{form_suffix}'
    m = re.search('\\b(\\d+)\\s*[X*×]\\s*(\\d+(?:\\.\\d+)?)\\s*(PCS?)\\b', s_upper)
    if m:
        n1, n2, suffix = m.groups()
        return f'{n1}X{n2}{suffix}'
    m = re.search('\\b(\\d+)\\s*([STCP])\\b', s_upper)
    if m:
        return f'{m.group(1)}{m.group(2)}'
    m = re.search('\\b(\\d+)\\s*(?:TA|TABS?|TABLETS?|STRIPS?|CAPS?|CAPSULES?)\\b', s_upper)
    if m:
        return f'{m.group(1)}T'
    m = re.search('\\b(?:TA|TABS?|TABLETS?|STRIPS?|CAPS?|CAPSULES?)\\s*(\\d+)\\b', s_upper)
    if m:
        return f'{m.group(1)}T'
    return ''

def format_output_with_pack_size(product_name: str, input_name: str='', master_pack_size: str='') -> str:
    return product_name

def make_result(output, product_code='', status='MATCHED', confidence='HIGH', candidate_count=0, suggestions=None):
    return {'output': output, 'product_code': product_code, 'status': status, 'confidence': confidence, 'candidate_count': int(candidate_count or 0), 'suggestions': suggestions or []}

def _suggestions_from_ranked(ranked, top_n=TOP_N_SUGGESTIONS):
    out = []
    for score, item in ranked[:top_n]:
        out.append({'product': item.get('product', ''), 'product_code': item.get('product_code', ''), 'score': round(float(score), 1)})
    return out

def build_item_suggestions(items, input_name, top_n=TOP_N_SUGGESTIONS):
    if not items:
        return []
    return _suggestions_from_ranked(rank_local_candidates(items, input_name), top_n)

def suggest_nearest_brands(input_name, brand_map, top_n=TOP_N_SUGGESTIONS):
    if not brand_map:
        return []
    query = norm(extract_brand_like_query(input_name)) or norm(input_name)
    if not query:
        return []
    norm_to_brand = {}
    for brand in brand_map.keys():
        norm_to_brand.setdefault(norm(brand), brand)
    matches = process.extract(query, list(norm_to_brand.keys()), scorer=fuzz.token_sort_ratio, limit=top_n)
    brand_query = extract_brand_like_query(input_name)
    out = []
    for norm_brand, score, _ in matches:
        brand = norm_to_brand[norm_brand]
        if is_substring_brand_reject(brand_query, brand):
            continue
        brand_items = brand_map.get(brand, [])
        best = choose_best_local_candidate(brand_items, input_name) if brand_items else None
        out.append({'brand': brand, 'product': best['product'] if best else brand, 'product_code': best.get('product_code', '') if best else '', 'score': round(float(score), 1)})
    return out

def format_suggestions(suggestions):
    if not suggestions:
        return ''
    parts = []
    for s in suggestions:
        code = f" ({s['product_code']})" if s.get('product_code') else ''
        parts.append(f"{s.get('product', '')}{code} [{s.get('score', '')}]")
    return ' | '.join(parts)

def filter_items_by_name_strength(items: list, input_name: str, verbose=False) -> list:
    if len(items) <= 1:
        return items
    txt = collapse_spaced_thousands(normalize_ocr_numeric_noise(input_name))
    dose_pairs = re.findall('(\\d+(?:\\.\\d+)?)\\s*(MG|MCG|GM|IU|ML|UNIT)\\b', txt)
    if not dose_pairs:
        return items
    matched = []
    for it in items:
        name = collapse_spaced_thousands(str(it.get('product', '')).upper())
        if any((re.search('(?<![\\d.])' + re.escape(num) + '\\s*' + unit, name) for num, unit in dose_pairs)):
            matched.append(it)
    if matched and len(matched) < len(items):
        if verbose:
            print(f'  [NAME-STRENGTH] dose {dose_pairs} -> {len(items)}->{len(matched)} items')
        return matched
    return items

def filter_items_prefer_exact_brand_token(items: list, brand: str, input_name: str, verbose=False) -> list:
    if len(items) <= 1 or not brand:
        return items
    brand_first = brand.upper().replace('-', ' ').split()[0] if brand.split() else ''
    if not brand_first:
        return items
    input_tokens = set(re.findall('[A-Z]+', normalize_ocr_numeric_noise(input_name)))
    exact, glued = ([], [])
    for it in items:
        toks = re.findall('[A-Z]+', str(it.get('product', '')).upper())
        first = toks[0] if toks else ''
        if first != brand_first and first.startswith(brand_first) and (len(first) > len(brand_first)) and (first not in input_tokens) and (first[len(brand_first):] not in input_tokens):
            glued.append(it)
        else:
            exact.append(it)
    if exact and glued:
        if verbose:
            print(f"  [EXACT-BRAND] dropping glued sub-brand items ({[i['product'] for i in glued]})")
        return exact
    return items

def _explicit_form_recovery_strength(input_name: str, brand: str) -> str:
    """Limit catalog-form recovery to the verified explicit-strength families."""
    upper = input_name.upper()
    if brand == 'ENZELO':
        strengths = set(re.findall(r'(?<!\d)(10000|25000)(?!\d)', upper))
        if (len(strengths) == 1 and not re.search(r'(?<!\d)(150|300)(?!\d)', upper)
                and detect_dosage_form_in_input(input_name) & {'TABLET', 'CAPSULE'}):
            return next(iter(strengths))
    elif brand == 'SCUD':
        if (re.search(r'(?<!\d)250(?!\d)', upper)
                and not re.search(r'\bXP\b', upper)
                and not re.search(r'(?<!\d)(?:500|750|1000|1(?:\.5)?)\s*(?:MG|GM)\b', upper)
                and 'INJECTION' in detect_dosage_form_in_input(input_name)):
            return '250'
    return ''


def _expand_supported_glued_dt(input_name: str, items: list) -> str:
    matches = list(re.finditer(r'(?<![\d.])(\d+(?:\.\d+)?)DT\b', input_name.upper()))
    if not matches:
        return input_name
    tokens = set(re.findall(r'[A-Z]+', input_name.upper()))
    strengths = {str(item['sub_variant']) for item in items
                 if 'DT' in item['variant'].split()
                 and set(item['variant'].split()).issubset(tokens)}
    if all(match[1] in strengths for match in matches):
        return re.sub(r'(?<![\d.])(\d+(?:\.\d+)?)DT\b', r'\1 DT', input_name, flags=re.I)
    return input_name


def process_product(input_name: str, brand_map: dict, all_variants: set=None, all_sub_variants: set=None, client=None, verbose=False, _lower_variant_retry=False, forced_brand: str=None, auto_detect: bool=False) -> dict:
    raw_input = str(input_name or '').strip()
    hint = str(forced_brand or '').strip().upper()
    if hint == 'PRO-DIOF' or not hint and auto_detect and 'DIOF' in raw_input.upper():
        description = parse_pro_diof_description(raw_input)
        if description is not None:
            # Keep the established matching pass and its complete diagnostics.
            # New descriptor recovery never promotes previously rejected unknown
            # text; existing accepted descriptions still receive the pack policy.
            existing = _process_product_core(raw_input, brand_map, all_variants, all_sub_variants, client,
                                             verbose=verbose, _lower_variant_retry=_lower_variant_retry,
                                             forced_brand=forced_brand, auto_detect=auto_detect)
            if existing.get('product_code') and (description.error == 'VARIANT_NOT_IN_MASTER'
                                                  or description.five_ml and description.error == 'UNSUPPORTED_PRO_DIOF_QUANTITY'):
                description = replace(description, error='')
            item, refusal = select_pro_diof_presentation(description, brand_map)
            if item is None:
                return make_result('NO_CLEAR_MATCH', status=refusal, confidence='NONE')
            if existing.get('product_code') == item['product_code'] and existing.get('output') == item['product']:
                return existing
            return make_result(item['product'], item['product_code'], status='MATCHED', confidence='HIGH',
                               candidate_count=1, suggestions=[{'product': item['product'],
                                                               'product_code': item['product_code'], 'score': 100.0}])
    if hint == 'BEVON' and 'SOFT' in raw_input.upper():
        targets = [item for item in brand_map.get('BEVON', [])
                   if str(item.get('product', '')).strip().upper() == 'BEVON SOFTULES']
        if len(targets) != 1:
            raise ValueError('BEVON SOFT exception requires exactly one BEVON SOFTULES master product.')
        item = targets[0]
        return make_result(item['product'], item['product_code'], status='MATCHED', confidence='HIGH',
                           candidate_count=1, suggestions=[{'product': item['product'],
                                                           'product_code': item['product_code'], 'score': 100.0}])
    # These corrections require the confirmed catalog hint used by app.py.
    scoped_items = brand_map.get(forced_brand, [])
    strength = _explicit_form_recovery_strength(raw_input, forced_brand) if scoped_items else ''
    exclude_forte = bool(scoped_items and forced_brand == 'C - TAX'
                         and re.search(r'\bO\b.*\bDRY\b', raw_input.upper())
                         and not re.search(r'\bFORT(?:E|S)?\b', raw_input.upper()))
    prepared = _expand_supported_glued_dt(raw_input, scoped_items) if scoped_items else raw_input
    options = dict(verbose=verbose, _lower_variant_retry=_lower_variant_retry,
                   forced_brand=forced_brand, auto_detect=auto_detect,
                   _exclude_inferred_forte=exclude_forte)
    result = _process_product_core(prepared, brand_map, all_variants, all_sub_variants, client, **options)
    if strength:
        recovered = _process_product_core(prepared, brand_map, all_variants, all_sub_variants, client,
                                          _catalog_form_recovery=True, **options)
        # Keep original catalog identities and reject recovery to another variant/strength.
        if any(item['variant'] == 'PLAIN' and str(item['sub_variant']) == strength
               and item['product_code'] == recovered.get('product_code')
               and item['product'] == recovered.get('output') for item in scoped_items):
            return recovered
    return result


def _process_product_core(input_name: str, brand_map: dict, all_variants: set=None, all_sub_variants: set=None, client=None, verbose=False, _lower_variant_retry=False, forced_brand: str=None, auto_detect: bool=False, _catalog_form_recovery=False, _exclude_inferred_forte=False) -> dict:
    input_name = str(input_name or '').strip()
    if 'CORLION' in input_name.upper():
        alias = resolve_source_brand_alias(input_name, brand_map)
        if alias and (not forced_brand or compact_brand_token(forced_brand) == compact_brand_token(alias)):
            input_name = replace_source_brand_alias(input_name, alias)
    exact_items = _get_exact_catalog_index(brand_map).get(norm(input_name), [])
    if len(exact_items) == 1:
        item = exact_items[0]
        return make_result(item['product'], item.get('product_code', ''), status='MATCHED', confidence='HIGH', candidate_count=1, suggestions=[{'product': item['product'], 'product_code': item.get('product_code', ''), 'score': 100.0}])
    if len(exact_items) > 1:
        return make_result('NO_CLEAR_MATCH', '', status='AMBIGUOUS_EXACT_NAME', confidence='NONE', candidate_count=len(exact_items), suggestions=[{'product': item['product'], 'product_code': item.get('product_code', ''), 'score': 100.0} for item in exact_items[:TOP_N_SUGGESTIONS]])
    form_items = _get_form_catalog_index(brand_map).get(_canonical_form_name(input_name), [])
    if len(form_items) == 1:
        item = form_items[0]
        return make_result(item['product'], item.get('product_code', ''), status='MATCHED', confidence='HIGH', candidate_count=1, suggestions=[{'product': item['product'], 'product_code': item.get('product_code', ''), 'score': 100.0}])
    if len(form_items) > 1:
        return make_result('NO_CLEAR_MATCH', '', status='AMBIGUOUS_FORM_NAME', confidence='NONE', candidate_count=len(form_items), suggestions=[{'product': item['product'], 'product_code': item.get('product_code', ''), 'score': 100.0} for item in form_items[:TOP_N_SUGGESTIONS]])
    diof_source = None
    if 'DIOF' in input_name.upper() and (not forced_brand or compact_brand_token(forced_brand) == 'DIOF'):
        leading = re.match(r'[A-Z0-9]+', input_name.upper())
        literal_other_brand = (not forced_brand and leading and leading[0] in brand_map
                               and compact_brand_token(leading[0]) != 'DIOF')
        diof_source = None if literal_other_brand else normalize_diof_source(
            input_name, confirmed=bool(forced_brand) or auto_detect)
        if diof_source is not None:
            input_name = diof_source
    input_name = _prepare_maxtra_source(input_name, brand_map, forced_brand, auto_detect)
    input_name = clean_duplicate_words(input_name)
    forced_brand_variant_tags = get_brand_variant_tags(brand_map, forced_brand) if forced_brand else None
    input_name = strip_parenthetical_noise(input_name, forced_brand_variant_tags)
    input_name = desegment_ocr_input(input_name, brand_map)
    if verbose:
        print('\n' + '=' * 80)
        print('PRODUCT MATCHING')
        print('=' * 80)
        print(f"Input: '{input_name}'")
    brand = None
    items = []
    if forced_brand:
        forced_upper = forced_brand.strip().upper()
        if forced_upper in brand_map:
            brand = forced_upper
            items = list(brand_map[brand])
            if verbose:
                print(f"\n  [forced_brand] Using hint '{brand}' -- {len(items)} candidates")
        else:
            matched_key = next((k for k in brand_map if k.upper() == forced_upper), None)
            if matched_key:
                brand = matched_key
                items = list(brand_map[brand])
                if verbose:
                    print(f"\n  [forced_brand] Using hint '{brand}' (normalised) -- {len(items)} candidates")
            elif auto_detect:
                if verbose:
                    print(f"  [forced_brand] Hint '{forced_brand}' not in brand_map -> auto-detecting brand")
                brand, items = find_best_brand_for_input(input_name, brand_map, verbose=verbose)
            else:
                if verbose:
                    print(f"  [forced_brand] Hint '{forced_brand}' not in brand_map -> NO_CLEAR_MATCH")
                return make_result('NO_CLEAR_MATCH', '', status='NO_BRAND', confidence='NONE', candidate_count=0, suggestions=[])
    elif auto_detect:
        if verbose:
            print(f"  [auto_detect] No brand hint -- running auto brand detection for '{input_name[:50]}'")
        brand, items = find_best_brand_for_input(input_name, brand_map, verbose=verbose)
    else:
        if verbose:
            print(f"  [forced_brand] No brand hint for '{input_name[:50]}' -> NO_CLEAR_MATCH")
        return make_result('NO_CLEAR_MATCH', '', status='NO_BRAND', confidence='NONE', candidate_count=0, suggestions=[])
    if not brand and auto_detect and _FORTE_HINT.search(input_name):
        brand = _recover_forte_glued_brand(input_name, brand_map)
        items = list(brand_map[brand]) if brand else []
    if not brand:
        suggestions = suggest_nearest_brands(input_name, brand_map)
        top, dist = nearest_brand_within_edits(input_name, suggestions)
        if top and top.get('product'):
            if verbose:
                print(f"\nRECOVERED (no brand) -> '{top['product']}' (brand '{top.get('brand', '')}' within {dist} edit(s))")
            return make_result(top['product'], top.get('product_code', ''), status='RECOVERED_NO_BRAND', confidence='LOW', candidate_count=0, suggestions=suggestions)
        if verbose:
            print(f'\nNO_CLEAR_MATCH: no brand within {NO_BRAND_RECOVERY_MAX_EDITS} edits (nearest gap {dist})')
        return make_result('NO_CLEAR_MATCH', '', status='NO_BRAND', confidence='NONE', candidate_count=0, suggestions=suggestions)
    maxtra_input = False
    diof_input = compact_brand_token(brand) == 'DIOF'
    if diof_input:
        if match_pro_brand(input_name):
            return make_result('NO_CLEAR_MATCH', status='CONFLICTING_BRAND', confidence='NONE')
        prepared = diof_source or normalize_diof_source(input_name, confirmed=True)
        if prepared is None:
            return make_result('NO_CLEAR_MATCH', status='NO_BRAND', confidence='NONE')
        input_name = desegment_ocr_input(prepared, brand_map)
        items, refusal = _prepare_diof_candidates(input_name, items)
        if refusal:
            return make_result('NO_CLEAR_MATCH', status=refusal, confidence='NONE')
    if brand == 'MAXTRA':
        input_name, items, maxtra_input, refusal = _prepare_maxtra_candidates(input_name, items)
        if refusal:
            return make_result('NO_CLEAR_MATCH', '', status=refusal, confidence='NONE')
    forte_input = False
    if _FORTE_HINT.search(input_name) and compact_brand_token(brand) != 'PRODIOF':
        input_name, items, forte_input, refusal = _prepare_forte_candidates(input_name, brand, items)
        if refusal:
            return make_result('NO_CLEAR_MATCH', '', status=refusal, confidence='NONE', candidate_count=0, suggestions=[])
    if forte_input:
        _exclude_inferred_forte = False
    elif compact_brand_token(brand) == 'CTAX' and re.search(r'\bO\b.*\bDRY\b', input_name.upper()):
        _exclude_inferred_forte = True
    _input_forms = detect_dosage_form_in_input(input_name, verbose=False)
    if _catalog_form_recovery and _input_forms & {'TABLET', 'CAPSULE'}:
        _input_forms = _input_forms | {'TABLET', 'CAPSULE'}
    items = prioritize_by_dosage_form(items, _input_forms, verbose=verbose,
                                     _catalog_form_recovery=_catalog_form_recovery)
    compound_tokens = extract_brand_suffix_tokens(input_name, brand)
    compound_token_digits = get_compound_token_digits(compound_tokens)
    compound_token_letters = set()
    for token in compound_tokens:
        letters = re.sub('\\d', '', token)
        if letters:
            compound_token_letters.add(letters)
    if compound_tokens:
        if verbose:
            print(f'\n  Compound tokens: {compound_tokens}')
        items = filter_items_by_compound_tokens(items, compound_tokens, verbose=verbose)
    items = filter_items_by_product_name_abbrev(input_name, brand, items, verbose=verbose)
    items = filter_items_prefer_exact_brand_token(items, brand, input_name, verbose=verbose)

    def _is_semi_item(it):
        return bool(re.search('\\b(?:SEMI|HALF)\\b', str(it.get('product', '')).upper()))
    _input_has_semi = bool(re.search('\\b(?:SEMI|HALF)\\b', normalize_ocr_numeric_noise(input_name)))
    _semi_items = [it for it in items if _is_semi_item(it)]
    _plain_items = [it for it in items if not _is_semi_item(it)]
    if _semi_items and _plain_items:
        items = _semi_items if _input_has_semi else _plain_items
        if verbose:
            print(f'\n  SEMI/HALF prefix discriminator: input_has_semi={_input_has_semi} -> kept {len(items)} of {len(_semi_items) + len(_plain_items)} SKUs')
    brand_scope_items = list(items)
    available_variants, available_sub_variants, _ = extract_available_attributes_from_items(brand_scope_items)
    normalized_for_qual = normalize_ocr_numeric_noise(input_name)
    orphan_quals = find_unmatched_qualifier_tokens(normalized_for_qual, brand, available_variants, available_sub_variants, brand_items=brand_scope_items)
    if compound_token_letters:
        orphan_quals = [q for q in orphan_quals if q not in compound_token_letters]
    forte_conflicts = set()
    if forte_input:
        # Qualifier scanning stops at the first form; explicit conflicting identity
        # tokens remain disqualifying even when printed after that form.
        variant_tokens, subvariant_tokens = _variant_subvariant_tokens(brand_map)
        source_items = brand_map.get(brand, ())
        identity_tokens = _variant_vocabulary(tuple(str(item.get(field, 'PLAIN')).upper()
                                                   for item in source_items for field in ('variant', 'sub_variant')))
        protected = {token for token in variant_tokens | subvariant_tokens | identity_tokens
                     if not token.isdigit()} - _pack_form_skip_tokens() - {'PLAIN'}
        active = _variant_vocabulary(tuple(sorted(available_variants | available_sub_variants)))
        printed = _input_variant_tokens_outside_packs(input_name, brand, available_variants)
        forte_conflicts = ((printed | set(orphan_quals)) & protected) - active
        if 'OF' in forte_conflicts and re.search(r'\b(?:BOTTLES?|PACK)\s+OF\s+\d', normalized_for_qual):
            forte_conflicts.discard('OF')
        orphan_quals += sorted(forte_conflicts - set(orphan_quals))
    if orphan_quals:
        def _retry_nasal_abbreviation(result):
            # Preserve successful existing matches and expand only a complete phrase.
            if result.get('product_code') or result.get('status') != 'VARIANT_NOT_IN_MASTER' or '/' not in input_name:
                return result
            expanded = _NASAL_SPRAY_ABBREVIATION.sub('NASAL SPRAY', input_name)
            if expanded == input_name:
                return result
            recovered = _process_product_core(expanded, brand_map, all_variants, all_sub_variants, client,
                                        verbose=verbose, _lower_variant_retry=_lower_variant_retry,
                                        forced_brand=forced_brand, auto_detect=auto_detect,
                                        _catalog_form_recovery=_catalog_form_recovery,
                                        _exclude_inferred_forte=_exclude_inferred_forte)
            if recovered.get('status') != 'MATCHED' or recovered.get('confidence') != 'HIGH':
                return result
            descriptor_text = strip_brand_from_input(expanded, brand, available_variants)
            if descriptor_text == expanded.upper().strip():
                return result
            descriptors = extract_variant_match_tokens(descriptor_text) - _pack_form_skip_tokens()
            if not descriptors.issubset(extract_variant_match_tokens(recovered['output'])):
                return result
            quantity_pattern = r'(?<![\d.])(\d+(?:\.\d+)?)\s*(MCG|MG|GM|ML|KG|MD|G|L)\b'
            catalog_quantities = {(float(number), 'GM' if unit == 'G' else unit)
                                  for number, unit in re.findall(quantity_pattern, recovered['output'].upper())}
            if catalog_quantities and any(
                    (float(number), 'GM' if unit == 'G' else unit) not in catalog_quantities
                    for number, unit in re.findall(quantity_pattern, expanded.upper())):
                return result
            non_identity_suffixes = DOSAGE_UNIT_SUFFIXES | {'G', 'L', 'KG', 'S', 'T', 'C', 'P', 'X'}
            for number, suffix in re.findall(r'(?<![\d.])(\d+(?:\.\d+)?)\s*([A-Z]+)\b', expanded.upper()):
                if suffix not in non_identity_suffixes and not re.search(
                        r'(?<![\d.])' + re.escape(number) + r'\s*' + re.escape(suffix) + r'\b', recovered['output'].upper()):
                    return result
            return recovered

        numeric_orphan = any((any((ch.isdigit() for ch in o)) for o in orphan_quals))
        forte_retry_safe = False
        if forte_input:
            identities = {(item['variant'], item['sub_variant']) for item in brand_scope_items}
            forte_retry_safe = len(identities) == 1 and not forte_conflicts
        if not maxtra_input and not _lower_variant_retry and (not forte_input or forte_retry_safe) and (not numeric_orphan):
            stripped = input_name
            for o in orphan_quals:
                stripped = re.sub('\\b' + re.escape(o) + '\\b', ' ', stripped)
            stripped = re.sub('\\s+', ' ', stripped).strip()
            lower_vars = detect_variants_in_input(stripped, brand, available_variants, verbose=False,
                                                  _exclude_inferred_forte=_exclude_inferred_forte)
            if stripped and stripped != input_name and any((v != 'PLAIN' for v in lower_vars)):
                res = _process_product_core(stripped, brand_map, all_variants, all_sub_variants, client, verbose=verbose, _lower_variant_retry=True, forced_brand=forced_brand, auto_detect=auto_detect, _catalog_form_recovery=_catalog_form_recovery, _exclude_inferred_forte=_exclude_inferred_forte)
                if forte_input and res.get('product_code') and not any(
                        item['product_code'] == res['product_code'] and item['product'] == res['output']
                        for item in brand_scope_items):
                    return make_result('NO_CLEAR_MATCH', '', status='VARIANT_NOT_IN_MASTER', confidence='NONE')
                if res['status'] == 'MATCHED' and res['output'] != 'NO_CLEAR_MATCH':
                    if verbose:
                        print(f"\nRECOVERED (lower variant): '{brand}' lacks {orphan_quals}; mapped to lower variant '{res['output']}'")
                    res['status'] = 'RECOVERED_LOWER_VARIANT'
                    res['confidence'] = 'LOW'
                return _retry_nasal_abbreviation(res)
        suggestions = build_item_suggestions(brand_scope_items, input_name)
        if verbose:
            print(f"\nNO_CLEAR_MATCH: qualifier {orphan_quals} is not a known variant/sub-brand of '{brand}' in the master")
        return _retry_nasal_abbreviation(make_result('NO_CLEAR_MATCH', '', status='VARIANT_NOT_IN_MASTER', confidence='NONE', candidate_count=len(brand_scope_items), suggestions=suggestions))
    detected_variants = detect_variants_in_input(input_name, brand, available_variants, compound_token_letters=compound_token_letters, verbose=verbose, brand_items=brand_scope_items, _exclude_inferred_forte=_exclude_inferred_forte)
    _before_var = items
    items = filter_items_by_variant(items, detected_variants, reference_items=brand_scope_items, verbose=verbose)
    if _input_forms and items and (not _form_compatible(items, _input_forms, _catalog_form_recovery)):
        _pool = _form_compatible(_before_var, _input_forms, _catalog_form_recovery)
        if _pool:
            _retry = filter_items_by_variant(_pool, detected_variants, reference_items=_pool, verbose=verbose)
            if _retry:
                items = _retry
    # A newly eligible equivalent form may introduce a sibling variant (e.g. O
    # drops beside O DRY syrup). Do not parse its O 50 as this variant's strength.
    strength_variants = available_variants
    if items and any(_FORM_GROUPS.get(form) in _EQUIVALENT_FORM_GROUPS for form in _input_forms):
        strength_variants, _, _ = extract_available_attributes_from_items(items)
    detected_sub_variants = detect_sub_variants_in_input(input_name, brand, available_sub_variants, strength_variants, compound_token_digits=compound_token_digits, verbose=verbose)
    _before_sv = items
    items = filter_items_by_sub_variant(items, detected_sub_variants, input_name, reference_items=brand_scope_items, verbose=verbose, _catalog_form_recovery=_catalog_form_recovery)
    if _input_forms and items and (not _form_compatible(items, _input_forms, _catalog_form_recovery)):
        _pool = _form_compatible(_before_sv, _input_forms, _catalog_form_recovery)
        if _pool:
            _retry = filter_items_by_sub_variant(_pool, detected_sub_variants, input_name, reference_items=_pool, verbose=verbose, _catalog_form_recovery=_catalog_form_recovery)
            if _retry:
                items = _retry
    _before_ns = items
    pack_input = _pack_selection_input(input_name, items) if compact_brand_token(brand) != 'PRODIOF' else input_name
    items = filter_items_by_name_strength(items, pack_input, verbose=verbose)
    if _input_forms and items and (not _form_compatible(items, _input_forms, _catalog_form_recovery)):
        _pool = _form_compatible(_before_ns, _input_forms, _catalog_form_recovery)
        if _pool:
            _retry = filter_items_by_name_strength(_pool, pack_input, verbose=verbose)
            if _retry:
                items = _retry
    legacy_pack = compact_brand_token(brand) == 'PRODIOF'
    requested_pack = None if legacy_pack else _input_pack_measurement(pack_input, items, _prefer_volume=brand == 'MAXTRA')
    input_pack_size = extract_pack_size_from_input(input_name) if legacy_pack else str(requested_pack[0]) if requested_pack else ''
    if requested_pack:
        sources = _pack_source_brands(brand_map)
        identities = {(sources.get(id(item), brand), str(item.get('variant', 'PLAIN')).upper(),
                       str(item.get('sub_variant', 'PLAIN')).upper()) for item in items}
        if len(identities) > 1:
            items = _explicit_variant_items_before_pack(items, input_name, brand)
            identities = {(sources.get(id(item), brand), str(item.get('variant', 'PLAIN')).upper(),
                           str(item.get('sub_variant', 'PLAIN')).upper()) for item in items}
            if len(identities) > 1:
                return make_result('NO_CLEAR_MATCH', '', status='AMBIGUOUS_PACK_IDENTITY', confidence='NONE',
                                   candidate_count=len(items), suggestions=build_item_suggestions(items, input_name))
    items = filter_items_by_pack_size(items, input_pack_size, reference_items=brand_scope_items, verbose=verbose,
                                      _keep_form_ties=bool(_form_precedence(input_name)),
                                      _measurement=requested_pack, _legacy=legacy_pack)
    if requested_pack and not items:
        return make_result('NO_CLEAR_MATCH', '', status='NO_COMPATIBLE_PACK', confidence='NONE', candidate_count=0)
    if maxtra_input:
        valid = [item for item in items if str(item.get('product_code', '')).strip() not in {'', '0', '1', 'nan', 'None'}]
        if len(items) != 1 or len(valid) != 1:
            return make_result('NO_CLEAR_MATCH', '', status='AMBIGUOUS_MAXTRA_TARGET' if len(items) > 1 else 'MISSING_MAXTRA_TARGET',
                               confidence='NONE', candidate_count=len(items), suggestions=build_item_suggestions(items, input_name))
    if diof_input:
        valid = [item for item in items if str(item.get('product_code', '')).strip() not in {'', '0', '1', 'nan', 'None'}]
        if len(items) != 1 or len(valid) != 1:
            return make_result('NO_CLEAR_MATCH', status='AMBIGUOUS_DIOF_TARGET' if len(items) > 1 else 'MISSING_DIOF_TARGET',
                               confidence='NONE', candidate_count=len(items), suggestions=build_item_suggestions(items, input_name))
    if forte_input:
        valid = [item for item in items if str(item.get('product_code', '')).strip() not in {'', '0', '1', 'nan', 'None'}]
        if len(valid) != 1 or len(items) != 1:
            return make_result('NO_CLEAR_MATCH', '', status='AMBIGUOUS_FORTE_TARGET' if len(items) > 1 else 'MISSING_FORTE_TARGET',
                               confidence='NONE', candidate_count=len(items), suggestions=build_item_suggestions(items, input_name))
    if has_ambiguous_missing_strength(items, detected_sub_variants):
        suggestions = build_item_suggestions(items, input_name)
        if verbose:
            print('\nNO_CLEAR_MATCH: ambiguous strength (input does not resolve to one SKU)')
        return make_result('NO_CLEAR_MATCH', '', status='AMBIGUOUS_STRENGTH', confidence='NONE', candidate_count=len(items), suggestions=suggestions)
    if not items:
        suggestions = build_item_suggestions(brand_scope_items, input_name)
        recovered = choose_best_local_candidate(brand_scope_items, input_name, detected_variants, detected_sub_variants)
        if recovered:
            if verbose:
                print(f"\nRECOVERED (over-filtered) -> '{recovered['product']}'")
            return make_result(recovered['product'], recovered.get('product_code', ''), status='RECOVERED_NO_CANDIDATES', confidence='LOW', candidate_count=len(brand_scope_items), suggestions=suggestions)
        if verbose:
            print('\n[FAIL] No products left after filtering')
        return make_result('NO_CLEAR_MATCH', '', status='NO_CANDIDATES', confidence='NONE', candidate_count=0, suggestions=suggestions)
    local_ranked = rank_local_candidates(items, input_name, detected_variants, detected_sub_variants)
    if requested_pack and len(local_ranked) > 1:
        order = _form_precedence(input_name) or _form_precedence(normalize_ocr_numeric_noise(input_name))
        # Reuse the existing presentation tie-breaker, including OCR form aliases.
        local_ranked.sort(key=lambda pair: (pair[0], _final_form_preference(pair[1]['product'], order)), reverse=True)
    local_best_item = local_ranked[0][1] if local_ranked else None
    local_best_score = local_ranked[0][0] if local_ranked else 0
    local_second_score = local_ranked[1][0] if len(local_ranked) > 1 else None
    strong_local_match = local_best_score >= 90 and (local_second_score is None or local_best_score - local_second_score >= 8)
    if requested_pack and len(local_ranked) > 1 and not strong_local_match:
        tied = [item for score, item in local_ranked if score == local_best_score]
        preferences = [_final_form_preference(item['product'], order) for item in tied]
        if len(tied) < 2 or preferences.count(max(preferences)) != 1:
            return make_result('NO_CLEAR_MATCH', '', status='AMBIGUOUS_PACK_TARGET', confidence='NONE',
                               candidate_count=len(items), suggestions=_suggestions_from_ranked(local_ranked))
    if verbose and local_ranked:
        print('\n' + '=' * 80)
        print('LOCAL FINAL CANDIDATE RANKING')
        print('=' * 80)
        for score, item in local_ranked[:5]:
            print(f"  {score:5.1f} -> {item['product']}")
    local_suggestions = _suggestions_from_ranked(local_ranked)
    if local_best_item:
        if len(items) == 1:
            if verbose:
                print('\n[OK] Single candidate -- skipping reranker')
            formatted_output = format_output_with_pack_size(local_best_item['product'], input_name, local_best_item.get('pack_size', ''))
            return make_result(formatted_output, local_best_item.get('product_code', ''), status='MATCHED', confidence='HIGH', candidate_count=len(items), suggestions=local_suggestions)
        if strong_local_match:
            if verbose:
                print('\n[OK] Strong local match -- skipping reranker')
            formatted_output = format_output_with_pack_size(local_best_item['product'], input_name, local_best_item.get('pack_size', ''))
            return make_result(formatted_output, local_best_item.get('product_code', ''), status='MATCHED', confidence='HIGH', candidate_count=len(items), suggestions=local_suggestions)
    if local_best_item:
        return make_result(local_best_item['product'], local_best_item.get('product_code', ''), status='RECOVERED_LOCAL_ONLY', confidence='LOW', candidate_count=len(items), suggestions=local_suggestions)
    return make_result('NO_CLEAR_MATCH', '', status='LOCAL_ONLY_UNRESOLVED', confidence='NONE', candidate_count=len(items), suggestions=local_suggestions)

def suggest_local_match(input_name: str, brand_map: dict, verbose=False) -> tuple:
    result = process_product(input_name, brand_map, verbose=verbose, auto_detect=True)
    return (result['output'], result['product_code'])

def review_highlighted_output_file(output_xlsx_path: str, suggestions_xlsx_path: str='') -> str:
    print(f'Loading master product data...')
    df_master = load_master(MASTER_XLSX_PATH)
    brand_map = build_brand_product_map(df_master)
    print(f'Opening highlighted output file: {output_xlsx_path}')
    source_wb = load_workbook(output_xlsx_path)
    source_ws = source_wb[source_wb.sheetnames[0]]
    if not suggestions_xlsx_path:
        base = Path(output_xlsx_path)
        suggestions_xlsx_path = str(base.with_name(base.stem + '_suggestions.xlsx'))
    suggestions_wb = Workbook()
    suggestions_ws = suggestions_wb.active
    suggestions_ws.title = 'HighlightedSuggestions'
    suggestions_ws.append(['Row_No', 'Input_Column', 'Current_Output', 'Current_Product_Code', 'Suggested_Output', 'Suggested_Product_Code', 'Status'])
    highlighted_count = 0
    changed_count = 0
    for row_num in range(2, source_ws.max_row + 1):
        if not any((source_ws.cell(row_num, col).fill.fill_type for col in range(1, source_ws.max_column + 1))):
            continue
        highlighted_count += 1
        input_name = source_ws.cell(row_num, 1).value or ''
        current_output = source_ws.cell(row_num, 2).value or ''
        current_code = source_ws.cell(row_num, 3).value or ''
        suggested_output, suggested_code = suggest_local_match(input_name, brand_map, verbose=False)
        status = 'Needs manual review'
        if str(current_output).strip().upper() != str(suggested_output).strip().upper():
            status = 'Suggested change'
            changed_count += 1
        suggestions_ws.append([row_num, input_name, current_output, current_code, suggested_output, suggested_code, status])
    suggestions_wb.save(suggestions_xlsx_path)
    print(f'[OK] Reviewed {highlighted_count} rows, suggested {changed_count} changes')
    return suggestions_xlsx_path

class _MappingSheetWriter:
    """Maintain a workbook exclusively in its saver thread, retaining other sheets."""

    def __init__(self, workbook_path, columns, initial_rows):
        self.path = Path(workbook_path)
        self.columns = tuple(columns)
        self.initial_rows = tuple(initial_rows)
        self._workbook = None
        self._sheet = None
        self._applied = 0
        self._image_payloads = []

    def _write_row(self, row_idx, values):
        for col_idx, value in enumerate(values, 1):
            self._sheet.cell(row=row_idx, column=col_idx).value = (
                '' if isinstance(value, float) and value != value else value
            )

    def _open(self):
        try:
            wb = load_workbook(self.path)
        except FileNotFoundError:
            wb = Workbook()
            for name in wb.sheetnames:
                if name in ('Sheet', 'Sheet1'):
                    del wb[name]
        if 'mapping' in wb.sheetnames:
            del wb['mapping']
        self._workbook = wb
        self._sheet = wb.create_sheet('mapping')
        try:
            for sheet in wb:
                for image in sheet._images:
                    payload = image.ref.getvalue() if isinstance(image.ref, BytesIO) else image._data()
                    self._image_payloads.append((image, payload))
            self._write_row(1, self.columns)
            for row_idx, values in enumerate(self.initial_rows, 2):
                self._write_row(row_idx, values)
        except BaseException:
            self.close()
            raise
        # The sheet now owns these values; do not retain a second full row index.
        self.initial_rows = ()

    def __call__(self, completed_rows):
        if self._workbook is None:
            self._open()
        for row_idx, values in enumerate(islice(completed_rows, self._applied, None), self._applied + 2):
            self._write_row(row_idx, values)
        self._applied = len(completed_rows)
        # openpyxl closes image streams while serializing; refresh them for each
        # checkpoint so unrelated-sheet images survive reuse of this workbook.
        for image, payload in self._image_payloads:
            if isinstance(image.ref, BytesIO):
                image.ref.close()
            image.ref = BytesIO(payload)
        save_workbook_atomic(self._workbook, self.path)

    def close(self):
        if self._workbook is not None:
            try:
                for sheet in self._workbook:
                    for image in sheet._images:
                        if isinstance(image.ref, BytesIO):
                            image.ref.close()
            finally:
                self._workbook.close()
                self._workbook = self._sheet = None
                self._image_payloads = []


def _save_to_mapping_sheet(df, workbook_path):
    """Synchronous compatibility entry point with atomic workbook replacement."""
    frozen_rows = tuple(df.itertuples(index=False, name=None))
    writer = _MappingSheetWriter(workbook_path, tuple(df.columns), frozen_rows)
    try:
        writer(())
    finally:
        writer.close()

def format_row_result(position: int, total: int, input_name: str, res: dict) -> str:
    mapped = str(res.get('output', '') or '')
    code = str(res.get('product_code', '') or '')
    tail = f'{mapped} [{code}]' if code else mapped
    counter = f'[{position:>{len(str(total))}}/{total}]'
    return f"{counter} {input_name[:44]:<44} -> {tail}  ({res.get('status', '')})"

def process_excel_file(workbook_path=None, master_path=None, *, autosave_interval=50):
    workbook_path = Path(workbook_path or INPUT_OUTPUT_XLSX_PATH)
    master = load_master_data(master_path)
    brand_map = build_brand_product_map(master)
    all_variants = extract_all_variants_from_data(brand_map)
    all_sub_variants = extract_all_sub_variants_from_data(brand_map)
    source = pd.read_excel(workbook_path, sheet_name='garbage_check', engine='openpyxl', dtype=object)
    if len(source.columns) < 3:
        raise ValueError('The garbage_check sheet needs product, remark, and matched-brand columns.')
    product_column, garbage_column, brand_column = source.columns[:3]
    remarks = source[garbage_column].astype(str).str.strip()
    confirmed = source.loc[remarks.eq('0'), [product_column, brand_column]].copy()
    confirmed.rename(columns={product_column: INPUT_COLUMN, brand_column: '_brand_hint'}, inplace=True)
    confirmed['_auto_detect'] = False
    confirmed['_source_label'] = '0'
    maybe = source.loc[remarks.eq('MAYBE_PRODUCT'), [product_column]].copy()
    maybe.rename(columns={product_column: INPUT_COLUMN}, inplace=True)
    maybe['_brand_hint'] = None
    maybe['_auto_detect'] = True
    maybe['_source_label'] = 'MAYBE_PRODUCT'
    if PROCESS_CONFIRMED_ZERO_ONLY and PROCESS_MAYBE_PRODUCT_ONLY:
        raise ValueError('Confirmed-only and maybe-only processing are mutually exclusive.')
    if PROCESS_CONFIRMED_ZERO_ONLY:
        rows = confirmed
    elif PROCESS_MAYBE_PRODUCT_ONLY:
        rows = maybe
    else:
        rows = pd.concat([confirmed, maybe], ignore_index=True)
    rows = rows.reset_index(drop=True)
    if ROW_LIMIT is not None:
        rows = rows.iloc[:ROW_LIMIT].copy()
    output_columns = [OUTPUT_COLUMN, PRODUCT_CODE_COLUMN, STATUS_COLUMN, CONFIDENCE_COLUMN, CANDIDATES_COLUMN, SUGGESTIONS_COLUMN, SOURCE_COLUMN]
    for column in output_columns:
        rows[column] = pd.Series('', index=rows.index, dtype=object)
    public_columns = tuple(column for column in rows if not column.startswith('_'))
    writer = _MappingSheetWriter(
        workbook_path, public_columns,
        tuple(rows.loc[:, list(public_columns)].itertuples(index=False, name=None)),
    )
    print(f'Processing {len(rows)} product rows locally...')
    with AsyncWorkbookSaver(writer, interval=autosave_interval, on_progress=print) as saver:
        for idx, row in rows.iterrows():
            value = row[INPUT_COLUMN]
            input_name = '' if pd.isna(value) else str(value).strip()
            hint_value = row.get('_brand_hint')
            hint = None if pd.isna(hint_value) else str(hint_value).strip() or None
            if not input_name:
                result = make_result('SKIPPED_EMPTY_INPUT', status='SKIPPED_EMPTY', confidence='NONE')
            else:
                try:
                    result = process_product(input_name, brand_map, all_variants, all_sub_variants, forced_brand=hint, auto_detect=bool(row['_auto_detect']))
                except Exception as exc:
                    result = make_result(f'ERROR: {str(exc)[:100]}', status='ERROR', confidence='NONE')
            for column, key in ((OUTPUT_COLUMN, 'output'), (PRODUCT_CODE_COLUMN, 'product_code'), (STATUS_COLUMN, 'status'), (CONFIDENCE_COLUMN, 'confidence'), (CANDIDATES_COLUMN, 'candidate_count')):
                rows.at[idx, column] = result[key]
            rows.at[idx, SUGGESTIONS_COLUMN] = format_suggestions(result['suggestions'])
            rows.at[idx, SOURCE_COLUMN] = row['_source_label']
            print(format_row_result(idx + 1, len(rows), input_name, result))
            saver.append(tuple(rows.at[idx, column] for column in public_columns))
    results = rows.drop(columns=[c for c in rows if c.startswith('_')])
    print(f'Saved {len(results)} rows to {workbook_path} -> mapping')
    return results

def interactive_mode(master_path=None):
    master = load_master_data(master_path)
    brand_map = build_brand_product_map(master)
    print(f'Zuventus local product mapping: {len(brand_map)} brands')
    print('Commands: exit, brands, stats')
    while True:
        input_name = input('Enter INPUT PRODUCT NAME: ').strip()
        if input_name.lower() in {'exit', 'quit', 'q'}:
            break
        if not input_name:
            continue
        if input_name.lower() == 'brands':
            for brand, items in sorted(brand_map.items()):
                print(f'{brand}: {len(items)} products')
            continue
        if input_name.lower() == 'stats':
            print(f'Brands: {len(brand_map)}, products: {sum(map(len, brand_map.values()))}')
            continue
        result = process_product(input_name, brand_map, verbose=True, auto_detect=True)
        print(f"Result: {result['output']} [{result['product_code']}]")
        print(f"Status: {result['status']}; confidence: {result['confidence']}")
        if result['suggestions']:
            print('Suggestions: ' + format_suggestions(result['suggestions']))
_EXACT_CACHE_MAP = None
_EXACT_CACHE_INDEX = None

def _get_exact_catalog_index(brand_map: dict) -> dict:
    global _EXACT_CACHE_MAP, _EXACT_CACHE_INDEX
    if _EXACT_CACHE_MAP is brand_map:
        return _EXACT_CACHE_INDEX
    index = {}
    seen = set()
    for items in brand_map.values():
        for item in items:
            identity = (item['product'], item.get('product_code', ''), item.get('pack_size', ''))
            if identity not in seen:
                seen.add(identity)
                index.setdefault(norm(item['product']), []).append(item)
    _EXACT_CACHE_MAP, _EXACT_CACHE_INDEX = (brand_map, index)
    return index

def load_master_data(path=None) -> pd.DataFrame:
    return load_master(path or MASTER_XLSX_PATH)
_FORM_NAME_EQUIVALENTS = {'TABLET': 'TABLET', 'TABLETS': 'TABLET', 'TAB': 'TABLET', 'TABS': 'TABLET', 'TA': 'TABLET', 'CAPSULE': 'CAPSULE', 'CAPSULES': 'CAPSULE', 'CAP': 'CAPSULE', 'CAPS': 'CAPSULE', 'INJECTION': 'INJECTION', 'INJECTIONS': 'INJECTION', 'INJ': 'INJECTION', 'INJN': 'INJECTION', 'SYRUP': 'SYRUP', 'SYRUPS': 'SYRUP', 'SYP': 'SYRUP', 'SYR': 'SYRUP', 'SUSPENSION': 'SUSPENSION', 'SUSPENSIONS': 'SUSPENSION', 'SUSP': 'SUSPENSION', 'SUSPN': 'SUSPENSION', 'DROP': 'DROPS', 'DROPS': 'DROPS'}
# Reuse existing aliases without changing OCR or complete-name identity indexes.
_DOSAGE_FORM_ALIASES = {form: form for form, group in _FORM_GROUPS.items() if group in _EQUIVALENT_FORM_GROUPS}
_DOSAGE_FORM_ALIASES.update(_FORM_NAME_EQUIVALENTS)
_DOSAGE_FORM_ALIASES.update({alias.rstrip('.'): _FORM_NAME_EQUIVALENTS.get(form, form)
                            for alias, form in INPUT_TOKEN_EQUIVALENTS.items()
                            if _FORM_NAME_EQUIVALENTS.get(form, form) in _DOSAGE_FORM_ALIASES})
_DOSAGE_FORM_ALIASES.update({'PFS': 'INJECTION', 'SOL': 'SOLUTION', 'SOLUTIONS': 'SOLUTION',
                             'TONICS': 'TONIC', 'POW': 'POWDER', 'POWD': 'POWDER', 'POWDERS': 'POWDER',
                             'GRANULE': 'GRANULES', 'SACHETS': 'SACHET', 'LIQUIDS': 'LIQUID', 'VIALS': 'VIAL'})
_DOSAGE_FORM_ALIASES.update({alias: 'AMPULE' for alias in ('AMP', 'AMPS', 'AMPUL', 'AMPULE', 'AMPULES',
                                                         'AMPOULE', 'AMPOULES', 'AMPOUL', 'AMPOULS', 'AMPOULEP', 'AMPOULP')})
_DOSAGE_FORM_PATTERN = re.compile(r'(?<![A-Z])(?:' + '|'.join(re.escape(alias) for alias in sorted(_DOSAGE_FORM_ALIASES, key=len, reverse=True)) + r'|I\.?V\.?)(?![A-Z])', re.IGNORECASE)
FORM_HINT_TOKENS.update(set(_DOSAGE_FORM_ALIASES) - {'CAPSULS'})
GENERIC_VARIANT_MATCH_TOKENS.update(form for form, group in _FORM_GROUPS.items() if group in _EQUIVALENT_FORM_GROUPS)
_FORM_CACHE_MAP = None
_FORM_CACHE_INDEX = None

def _canonical_form_name(value: str) -> str:
    return ' '.join((_FORM_NAME_EQUIVALENTS.get(token, token) for token in norm(value).split()))

def _get_form_catalog_index(brand_map: dict) -> dict:
    global _FORM_CACHE_MAP, _FORM_CACHE_INDEX
    if _FORM_CACHE_MAP is brand_map:
        return _FORM_CACHE_INDEX
    index = {}
    for items in _get_exact_catalog_index(brand_map).values():
        for item in items:
            index.setdefault(_canonical_form_name(item['product']), []).append(item)
    _FORM_CACHE_MAP, _FORM_CACHE_INDEX = (brand_map, index)
    return index
if __name__ == '__main__':
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(errors='backslashreplace')
    import argparse
    parser = argparse.ArgumentParser(description='Zuventus local product mapping')
    parser.add_argument('--master', type=Path, default=MASTER_XLSX_PATH)
    parser.add_argument('--workbook', type=Path, default=INPUT_OUTPUT_XLSX_PATH)
    parser.add_argument('--product', help='Map one product locally')
    parser.add_argument('--interactive', action='store_true')
    parser.add_argument('--autosave-every', type=int, default=50, metavar='ROWS', help='Background checkpoint interval; 0 saves only the final workbook (default: 50)')
    args = parser.parse_args()
    if args.autosave_every < 0:
        parser.error('--autosave-every must be zero or greater')
    if args.product is not None:
        catalog = build_brand_product_map(load_master_data(args.master))
        result = process_product(args.product, catalog, auto_detect=True)
        for key in ('output', 'product_code', 'status', 'confidence'):
            print(f'{key}: {result[key]}')
    elif args.interactive:
        interactive_mode(args.master)
    else:
        process_excel_file(args.workbook, args.master, autosave_interval=args.autosave_every)
