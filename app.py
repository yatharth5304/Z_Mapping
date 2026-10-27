#!/usr/bin/env python3
"""Run the local Zuventus classification and product-mapping Excel workflow."""

from __future__ import annotations

import argparse
from contextlib import suppress
from datetime import datetime, timedelta, timezone
import os
from pathlib import Path
import shutil
import sys
import time

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ERROR_CODES
from openpyxl.utils import get_column_letter

from excel_autosave import AsyncWorkbookSaver, save_workbook_atomic

import garbage_check as gc
import mapping as mp



PROJECT_DIR = Path(__file__).resolve().parent
INPUT_XLSX = Path(os.environ.get("MAPPING_INPUT_XLSX", PROJECT_DIR / "mapping.xlsx"))
MASTER_XLSX = Path(os.environ.get("MASTER_XLSX", PROJECT_DIR / "zuventus_product.xlsx"))
GC_COLS = [
    "PRODUCT_NAME", "Garbage check", "Matched brand", "Matched text",
    "Match type", "Matched garbage phrase", "Garbage match type",
]
FINAL_COLS = ["PRODUCT_NAME", "PRODUCT_CODE", "MASTER_PRODUCT_NAME"]


def validate_input_file(xlsx_path: Path) -> list[str]:
    """Read the first sheet's PRODUCT_NAME column, preserving order and repeats."""
    if not xlsx_path.is_file():
        raise FileNotFoundError(f"Input file not found: {xlsx_path}")
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        rows = iter(wb.worksheets[0].iter_rows(values_only=True))
        header = next(rows, ())
        matches = [i for i, value in enumerate(header)
                   if str(value).strip().upper() == "PRODUCT_NAME"]
        if len(matches) != 1:
            raise ValueError("The first sheet must contain exactly one PRODUCT_NAME column")
        column = matches[0]
        names = ["" if column >= len(row) or row[column] is None else str(row[column])
                 for row in rows]
        # Excel formatting can extend the used range past the last actual input.
        while names and not names[-1].strip():
            names.pop()
        if not names:
            raise ValueError("Input file has no product rows")
        return names
    finally:
        wb.close()


def create_backup(xlsx_path: Path) -> Path:
    """Keep a recoverable copy before replacing an existing output workbook."""
    stamp = datetime.now(timezone(timedelta(hours=5, minutes=30))).strftime("%Y%m%d_%H%M%S_%f")
    target = xlsx_path.with_name(f"{xlsx_path.stem}_backup_{stamp}{xlsx_path.suffix}")
    shutil.copy2(xlsx_path, target)
    return target


def build_garbage_check_resources(master_xlsx: Path):
    brands = gc.load_brands(str(master_xlsx))
    brand_index = gc.build_brand_index(brands)
    evidence = gc.build_master_product_evidence(brand_index, str(master_xlsx))
    return brand_index, evidence, gc.build_garbage_set()


def run_garbage_check_phase(product_names, brand_index, master_evidence, garbage_rows):
    results = []
    for name in product_names:
        values = gc.classify_single_product(name, brand_index, master_evidence, garbage_rows)
        results.append(dict(zip(GC_COLS, [name, *values])))
    return results


def run_mapping_phase(gc_results, master_xlsx: Path, *, verbose=False, on_row=None):
    """Map confirmed rows only; retain diagnostics separately from public output."""
    df = mp.load_master_data(str(master_xlsx))
    brand_map = mp.build_brand_product_map(df)
    variants = mp.extract_all_variants_from_data(brand_map)
    subvariants = mp.extract_all_sub_variants_from_data(brand_map)
    final_results, diagnostics = [], []
    for index, row in enumerate(gc_results, start=1):
        name, classification = row["PRODUCT_NAME"], row["Garbage check"]
        result = None
        if classification == "1":
            code, master_name = "1", "GARBAGE"
        elif classification == "0":
            result = mp.process_product(
                name, brand_map, variants, subvariants, None,
                forced_brand=row["Matched brand"], auto_detect=False, verbose=False,
            )
            raw_code = result.get("product_code")
            if raw_code is None or str(raw_code).strip() in {"", "0", "1", "None", "nan"}:
                code, master_name = "0", "NO SUGGESTION"
            else:
                code = str(raw_code).strip()
                master_name = str(result["output"])
        else:
            code, master_name = "0", "NO SUGGESTION"
        final_results.append(dict(zip(FINAL_COLS, [name, code, master_name])))
        diagnostics.append(result)
        if verbose:
            status = result.get("status", "") if result else classification
            print(f"[{index}/{len(gc_results)}] {name} -> {master_name} [{code}] ({status})")
        if on_row is not None:
            on_row(final_results[-1])
    return final_results, diagnostics


def write_output_workbook(output_xlsx: Path, gc_results, final_results, *, progress=None):
    """Stream independent rows in linear time and atomically commit both sheets."""
    wb = Workbook(write_only=True)
    try:
        for title, columns, rows in (("garbage_check", GC_COLS, gc_results), ("mapping", FINAL_COLS, final_results)):
            ws = wb.create_sheet(title)
            ws.freeze_panes = "A2"
            ws.column_dimensions["A"].width = 55
            if title == "mapping":
                ws.column_dimensions["B"].width = 18
                ws.column_dimensions["C"].width = 65
            ws.append(columns)
            count = 0
            if progress:
                progress(f"[output] Writing {title} worksheet...")
            for count, row in enumerate(rows, start=1):
                values = [row.get(key, "") for key in columns] if isinstance(row, dict) else row
                cells = []
                for value in values:
                    # Ordinary scalars use openpyxl's reusable streaming cell.
                    # Only formula/error-looking text needs an explicit literal
                    # cell; wrapping every value adds an exception per cell in
                    # openpyxl's streaming adapter.
                    if isinstance(value, str) and (value.startswith("=") or value in ERROR_CODES):
                        cell = WriteOnlyCell(ws, value=value)
                        cell.data_type = "s"
                        cells.append(cell)
                    else:
                        cells.append(value)
                ws.append(cells)
                if progress and count % 10000 == 0:
                    progress(f"[output] {title}: {count} rows written")
            ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{count + 1}"
        if progress:
            progress("[output] Saving Excel workbook...")
        save_workbook_atomic(wb, output_xlsx)
    finally:
        wb.close()
        # openpyxl normally removes its streaming XML files after save. Also
        # clean them after an interrupted/failed serialization, before retrying.
        for ws in wb.worksheets:
            if not ws.closed:
                with suppress(Exception):
                    ws.close()
            writer = getattr(ws, "_writer", None)
            if writer is not None and Path(writer.out).exists():
                with suppress(OSError, ValueError):
                    writer.cleanup()


def run_pipeline(input_xlsx=INPUT_XLSX, master_xlsx=MASTER_XLSX, output_xlsx=None, *, backup=True, verbose=True, autosave_interval=1000):
    """Execute both local stages and return rows, diagnostics, counts, and paths."""
    input_path = Path(input_xlsx).expanduser().resolve()
    master_path = Path(master_xlsx).expanduser().resolve()
    output_path = Path(output_xlsx).expanduser().resolve() if output_xlsx else input_path
    if input_path == master_path or output_path == master_path:
        raise ValueError("The Product Master cannot be used as the input or output workbook")
    if not master_path.is_file():
        raise FileNotFoundError(f"Product Master not found: {master_path}")
    names = validate_input_file(input_path)
    if verbose:
        print(f"Zuventus local mapping: {len(names)} rows")
        print(f"Master: {master_path}")
    resources = build_garbage_check_resources(master_path)
    gc_results = run_garbage_check_phase(names, *resources)
    # Every checkpoint retains all source descriptions on the first sheet,
    # including rows whose mapping has not yet been processed.
    frozen_gc = tuple(tuple(row.get(column, "") for column in GC_COLS) for row in gc_results)
    backup_path = create_backup(output_path) if backup and output_path.exists() else None
    progress = (lambda message: print(message, flush=True)) if verbose else None

    def write_checkpoint(rows):
        write_output_workbook(output_path, frozen_gc, rows, progress=progress)

    with AsyncWorkbookSaver(write_checkpoint, interval=autosave_interval, on_progress=progress) as saver:
        final_results, diagnostics = run_mapping_phase(
            gc_results, master_path, verbose=verbose,
            on_row=lambda row: saver.append(tuple(row[column] for column in FINAL_COLS)),
        )
        if progress:
            progress("[output] Row processing complete; waiting for the final Excel checkpoint...")
    counts = {
        "total": len(final_results),
        "mapped": sum(row["PRODUCT_CODE"] not in {"0", "1"} for row in final_results),
        "garbage": sum(row["PRODUCT_CODE"] == "1" for row in final_results),
        "unresolved": sum(row["PRODUCT_CODE"] == "0" for row in final_results),
    }
    if verbose:
        print("Total: {total}; mapped: {mapped}; garbage: {garbage}; unresolved: {unresolved}".format(**counts))
        print(f"Saved: {output_path}")
        if backup_path:
            print(f"Backup: {backup_path}")
    return {"output": output_path, "backup": backup_path, "counts": counts,
            "garbage_check": gc_results, "mapping": final_results, "diagnostics": diagnostics}


def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="backslashreplace")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", "-i", type=Path, default=INPUT_XLSX, help="Excel file containing PRODUCT_NAME (default: mapping.xlsx)")
    parser.add_argument("--master", type=Path, default=MASTER_XLSX, help="Zuventus Product Master")
    parser.add_argument("--output", "-o", type=Path, help="Separate result workbook (default: update input in place)")
    parser.add_argument("--autosave-every", type=int, default=1000, help="Checkpoint interval in processed rows (0: final save only)")
    args = parser.parse_args(argv)
    if args.autosave_every < 0:
        parser.error("--autosave-every must be nonnegative")
    started = time.perf_counter()
    run_pipeline(args.input, args.master, args.output, autosave_interval=args.autosave_every)
    print(f"Completed in {time.perf_counter() - started:.1f} seconds")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        sys.exit(130)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)
